"""The generated deck read back as an editable scene (RF-46 … RF-48).

The editor and the mapping preview both draw from this, and only from this. Every number is taken
out of the .pptx the user is about to download — position, size, rotation, insets, font size,
colour — in EMU and points, with no rounding to CSS pixels on the way.

What changed in RF-48, after measuring the project's own template:

* geometry comes from :mod:`cadencia.report.core.geometry`, so the 49 shapes nested in groups land where
  the slide puts them rather than where their group's author was working, and the connectors and
  freeforms carry a real outline instead of a rectangle;
* fill, line and text come from :mod:`cadencia.report.core.properties`, so the 73 shapes that state no
  fill of their own and the 16 paragraphs that state no size get the answer the theme, the layout
  and the master actually give;
* text is kept as **runs**, not flattened to the first one, so a bold word inside a sentence
  survives being edited.

Anything still not understood is marked ``renderable: false``. The canvas then shows that shape as
a crop of the rendered slide instead of guessing — a faithful picture that simply is not editable,
which is always better than a wrong drawing.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from cadencia.report.core import geometry as geo
from cadencia.report.core import pptx_utils as px
from cadencia.report.core.edits import DeckEdits, shape_key, slide_key
from cadencia.report.core.manifest import parse_target
from cadencia.report.core.mapping import DeckMapping
from cadencia.report.core.properties import PropertyResolver

EMU_PER_PT = 12700
_ANCHOR = {"t": "top", "ctr": "middle", "b": "bottom"}
_ALIGN = {"l": "left", "ctr": "center", "r": "right", "just": "justify"}
#: Shapes the vector canvas cannot draw honestly; it crops the rendered slide for these instead.
_RASTER_KINDS = {"chart", "smartart", "ole", "media"}


def _bullet_of(paragraph: Any) -> dict[str, Any]:
    ppr = paragraph._p.pPr
    if ppr is None:
        return {"kind": "inherit"}
    if ppr.find(qn("a:buNone")) is not None:
        return {"kind": "none"}
    char = ppr.find(qn("a:buChar"))
    if char is not None:
        return {"kind": "char", "char": char.get("char") or "•"}
    if ppr.find(qn("a:buAutoNum")) is not None:
        return {"kind": "number"}
    return {"kind": "inherit"}


def _paragraph_of(
    paragraph: Any, shape: Any, slide: Any, resolver: PropertyResolver, font_scale: float
) -> dict[str, Any]:
    """One paragraph with its resolved defaults and every run kept separate."""
    defaults = resolver.paragraph_defaults(shape, slide, paragraph.level)
    ppr = paragraph._p.pPr
    align = _ALIGN.get((ppr.get("algn") if ppr is not None else None) or "") or _ALIGN.get(
        defaults.get("align") or ""
    )
    runs = []
    for run in paragraph.runs:
        style = resolver.run_style(run, defaults)
        if font_scale != 1.0 and style.get("size_pt"):
            style["size_pt"] = round(style["size_pt"] * font_scale, 2)
        runs.append({"text": run.text, **style})
    base = dict(defaults)
    if font_scale != 1.0 and base.get("size_pt"):
        base["size_pt"] = round(base["size_pt"] * font_scale, 2)
    base.pop("align", None)
    return {
        "text": paragraph.text,
        "level": paragraph.level,
        "align": align,
        "line_spacing": paragraph.line_spacing if isinstance(paragraph.line_spacing, float) else None,
        "space_before_pt": round(paragraph.space_before.pt, 2)
        if paragraph.space_before is not None
        else None,
        "space_after_pt": round(paragraph.space_after.pt, 2) if paragraph.space_after is not None else None,
        "bullet": _bullet_of(paragraph),
        # The paragraph's resolved style is what an empty paragraph and the editor's defaults use.
        "style": runs[0] if runs else base,
        "runs": runs or [{"text": "", **base}],
    }


def _text_of(shape: Any, slide: Any, resolver: PropertyResolver, font_scale: float) -> dict[str, Any]:
    frame = shape.text_frame
    body = frame._txBody.bodyPr
    autofit = "none"
    if body.find(qn("a:normAutofit")) is not None:
        autofit = "shrink"
    elif body.find(qn("a:spAutoFit")) is not None:
        autofit = "resize"
    return {
        "paragraphs": [_paragraph_of(p, shape, slide, resolver, font_scale) for p in frame.paragraphs],
        "anchor": _ANCHOR.get(body.get("anchor") or "t", "top"),
        "word_wrap": (body.get("wrap") or "square") != "none",
        "autofit": autofit,
        "columns": int(body.get("numCol") or 1),
        "insets_pt": {
            "l": round((frame.margin_left or 0) / EMU_PER_PT, 2),
            "r": round((frame.margin_right or 0) / EMU_PER_PT, 2),
            "t": round((frame.margin_top or 0) / EMU_PER_PT, 2),
            "b": round((frame.margin_bottom or 0) / EMU_PER_PT, 2),
        },
    }


def _kind_of(shape: Any) -> str:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if getattr(shape, "has_table", False) and shape.has_table:
        return "table"
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return "chart"
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "image"
    if shape._element.tag.endswith("graphicFrame"):
        return "smartart"
    if geo.line_shape(shape):
        return "line"
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        return "textbox"
    return "shape"


def _shape_scene(placed: geo.Placed, slide: Any, resolver: PropertyResolver, order: int) -> dict[str, Any]:
    shape = placed.shape
    kind = _kind_of(shape)
    shape_outline = geo.outline(shape, placed.cx, placed.cy)
    flip_h, flip_v = geo.flipped(shape)
    entry: dict[str, Any] = {
        "key": shape_key(shape),
        "shape_id": shape.shape_id,
        "name": shape.name,
        "kind": kind,
        "z": order,
        "depth": placed.depth,
        "parent": f"sh{placed.parent_id}" if placed.parent_id else None,
        # Absolute slide geometry: the group transform is already folded in.
        "frame": {
            "x": placed.x,
            "y": placed.y,
            "cx": placed.cx,
            "cy": placed.cy,
            "rotation": round(placed.rotation, 3),
        },
        "flip": {"h": flip_h, "v": flip_v},
        "geometry": shape_outline["preset"],
        "path": shape_outline["path"],
        "stroke_only": shape_outline["stroke_only"] or kind == "line",
        "fill": resolver.fill(shape, slide),
        "line": resolver.line(shape, slide),
        "placeholder": None,
        # A group is a container: selectable, never painted.
        "shadow": resolver.effects(shape),
        "renderable": kind not in _RASTER_KINDS,
        "container": kind == "group",
        "alt": px.alt_text(shape),
    }
    if shape.is_placeholder:
        fmt = shape.placeholder_format
        entry["placeholder"] = {"idx": fmt.idx, "type": str(fmt.type)}
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame and kind != "table":
        entry["text"] = _text_of(shape, slide, resolver, placed.transform.font_scale)
    if kind == "image":
        try:
            entry["image"] = {
                "content_type": shape.image.content_type,
                "size": list(shape.image.size),
                "crop": {
                    "left": round(shape.crop_left, 4),
                    "right": round(shape.crop_right, 4),
                    "top": round(shape.crop_top, 4),
                    "bottom": round(shape.crop_bottom, 4),
                },
            }
        except (AttributeError, ValueError, NotImplementedError):
            entry["image"] = None
    if kind == "table":
        table = shape.table
        defaults = resolver.paragraph_defaults(shape, slide, 0)
        entry["table"] = {
            "rows": len(table.rows),
            "cols": len(table.columns),
            "col_widths": [int(c.width or 0) for c in table.columns],
            "row_heights": [int(r.height or 0) for r in table.rows],
            "cells": [[cell.text for cell in row.cells] for row in table.rows],
            "fills": [[resolver.fill(cell, slide).get("color") for cell in row.cells] for row in table.rows],
            "styles": [
                [
                    resolver.run_style(cell.text_frame.paragraphs[0].runs[0], defaults)
                    if cell.text_frame.paragraphs and cell.text_frame.paragraphs[0].runs
                    else defaults
                    for cell in row.cells
                ]
                for row in table.rows
            ],
        }
    if kind == "chart":
        entry["chart"] = {"chart_type": str(shape.chart.chart_type)}
    return entry


def _background_of(slide: Any) -> dict[str, str] | None:
    """The slide's *own* background, if someone set one — as opposed to a plain solid fill
    inherited from the layout or master, which `type` reports as `MSO_FILL_TYPE.BACKGROUND`
    rather than `None` and is not this slide's own override."""
    try:
        fill = slide.background.fill
        if fill.type == MSO_FILL_TYPE.SOLID:
            return {"type": "solid", "color": str(fill.fore_color.rgb)}
    except (AttributeError, TypeError):
        pass
    return None


def bound_targets(mapping: DeckMapping, origin_slide_id: int) -> dict[int, str]:
    """Shape id -> the field bound to it, so the editor can mark what a re-run will overwrite."""
    out: dict[int, str] = {}
    for binding in mapping.for_slide(origin_slide_id).bindings:
        try:
            target = parse_target(binding.target)
        except ValueError:
            continue
        out.setdefault(target.shape_id, binding.key or "texto fixo")
    return out


def deck_scene(
    path: str | Path,
    produced: list[dict[str, Any]],
    mapping: DeckMapping,
    edits: DeckEdits | None = None,
    *,
    theme: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Read a generated .pptx into the scene the editor draws and manipulates."""
    prs = Presentation(str(path))
    resolver = PropertyResolver(prs)
    origin = {int(entry["index"]): entry for entry in produced}

    slides = []
    for number, slide in enumerate(prs.slides, start=1):
        meta = origin.get(number, {})
        origin_id = int(meta.get("origin_slide_id") or slide.slide_id)
        key = slide_key(origin_id, meta.get("board_id"), int(meta.get("page") or 1))
        bound = bound_targets(mapping, origin_id)
        shapes = []
        for order, placed in enumerate(geo.iter_placed(slide.shapes)):
            entry = _shape_scene(placed, slide, resolver, order)
            field = bound.get(placed.shape.shape_id)
            entry["bound_field"] = field
            entry["bound"] = field is not None
            shapes.append(entry)
        slides.append(
            {
                "key": key,
                "index": number,
                "origin_slide_id": origin_id,
                "board_id": meta.get("board_id"),
                "page": int(meta.get("page") or 1),
                "layout": slide.slide_layout.name,
                "notes": geo.notes_of(slide)[:600],
                "background": _background_of(slide),
                "shapes": shapes,
                "edited": bool(edits and edits.for_slide(key).shapes),
            }
        )
    colors = dict(resolver.theme.colors)
    return {
        "slide_width": int(prs.slide_width),
        "slide_height": int(prs.slide_height),
        "theme": {
            "colors": colors,
            "fonts": {"major": resolver.theme.major_font, "minor": resolver.theme.minor_font},
            "background": colors.get("lt1", "FFFFFF"),
        },
        "slides": slides,
    }


def picture_blob(path: str | Path, key: str) -> tuple[bytes, str] | None:
    """The bytes of one picture already embedded in a generated deck, found by its scene key.

    The editor draws images from the deck itself rather than from the asset store, because most of
    them came with the template and were never uploaded here.
    """
    prs = Presentation(str(path))
    for slide in prs.slides:
        for shape, _parent in px.iter_shapes(slide.shapes):
            if shape_key(shape) != key or shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            try:
                return shape.image.blob, shape.image.content_type
            except (AttributeError, ValueError, NotImplementedError):
                return None
    return None
