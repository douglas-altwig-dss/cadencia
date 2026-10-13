"""A stable fingerprint of a generated deck, for the refactor in progress (stages 0–8).

The presentation module is about to be taken apart — typed variables, a data-source port, one
renderer, the status palette moving house. Unit tests say each piece still works; they do not say
that the *same .pptx* comes out the other end. This does: it walks a produced file and writes down
everything a reader would notice — the slides, the shapes, their words, their table cells, the
fills that carry a status colour — in a form that diffs cleanly.

Deliberately not captured: picture bytes. Charts are rendered by matplotlib and a byte comparison
would fail on a font hint or a library patch, which says nothing about whether the deck is right.
That a picture *is there*, and how big its box is, is captured.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from pptx import Presentation

from cadencia.report.core import pptx_utils as px


def _fill(shape: Any) -> str | None:
    """The solid colour a shape states for itself, which is what `status_fill` writes."""
    try:
        fill = shape.fill
        if fill.type is not None and "SOLID" in str(fill.type):
            return str(fill.fore_color.rgb)
    except (AttributeError, TypeError, ValueError, NotImplementedError):
        return None
    return None


def _shape_digest(shape: Any) -> dict[str, Any]:
    entry: dict[str, Any] = {"id": shape.shape_id, "name": shape.name}
    fill = _fill(shape)
    if fill:
        entry["fill"] = fill
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text:
            entry["text"] = text.splitlines()
    if getattr(shape, "has_table", False) and shape.has_table:
        entry["table"] = [[cell.text_frame.text.strip() for cell in row.cells] for row in shape.table.rows]
    if shape.shape_type is not None and "PICTURE" in str(shape.shape_type):
        # The bytes are matplotlib's and may move; the box is ours and must not.
        entry["picture"] = {"cx": int(shape.width or 0), "cy": int(shape.height or 0)}
    return entry


def deck_digest(path: str | Path) -> dict[str, Any]:
    """Everything a reader would notice about a generated deck, in diffable form."""
    prs = Presentation(str(path))
    slides = []
    for number, slide in enumerate(prs.slides, start=1):
        shapes = [_shape_digest(shape) for shape, _parent in px.iter_shapes(slide.shapes)]
        slides.append(
            {
                "number": number,
                "layout": slide.slide_layout.name,
                # Stage 5 will start writing this; capturing it now means the snapshot shows
                # the day it changes instead of the day somebody notices.
                "name": slide.element.cSld.get("name") or "",
                "shapes": sorted((s for s in shapes if len(s) > 2), key=lambda s: s["id"]),
            }
        )
    return {"slide_count": len(prs.slides), "slides": slides}


def run_digest(run: Any, validation: dict[str, Any], path: str | Path) -> dict[str, Any]:
    """The digest plus what the pipeline said about the run, which is half of what can regress."""
    return {
        "produced": [
            {k: entry[k] for k in ("index", "origin_slide_id", "board_id", "page") if k in entry}
            for entry in sorted(run, key=lambda e: e["index"])
        ],
        "validation": {
            "ok": validation["ok"],
            "errors": sorted(e.get("kind", "") for e in validation["errors"]),
            "warnings": sorted(w.get("kind", "") for w in validation["warnings"]),
            "package_ok": validation["package"]["ok"],
        },
        "deck": deck_digest(path),
    }


def compare(name: str, actual: dict[str, Any], *, update: bool = False) -> None:
    """Assert against the committed snapshot, or write it when there is none.

    A missing snapshot writes and passes — that is how the first run of a new case records itself.
    A *different* snapshot always fails: the whole point is that a refactor cannot quietly change
    the file. Set CADENCIA_UPDATE_GOLDEN=1 to re-record on purpose.
    """
    target = Path(__file__).parent / "golden" / f"{name}.json"
    target.parent.mkdir(parents=True, exist_ok=True)
    rendered = json.dumps(actual, indent=2, ensure_ascii=False, sort_keys=True)
    if update or not target.exists():
        target.write_text(rendered + "\n", encoding="utf-8")
        return
    expected = target.read_text(encoding="utf-8").strip()
    assert rendered == expected, (
        f"O deck gerado mudou em relação a tests/golden/{name}.json.\n"
        "Se a mudança é intencional, confira o diff e regrave com CADENCIA_UPDATE_GOLDEN=1."
    )
