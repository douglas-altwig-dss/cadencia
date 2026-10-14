"""Deck data added for the Apresentação roadmap: forecast, flow, trends, changes, data health and
the epic plan (#1–#6), and the chart styles that draw them."""

from __future__ import annotations

import asyncio

import pytest
from pptx import Presentation

from cadencia.demo import DemoJira
from cadencia.demo_seed import save_demo_team
from cadencia.report.adapters.jira_source import JiraSource
from cadencia.report.core.charts import render_chart_png
from cadencia.report.core.context import FIELD_INDEX, Series
from cadencia.report.core.generate import generate_deck
from cadencia.report.core.manifest import extract_manifest
from cadencia.report.core.mapping import DeckMapping
from cadencia.report.core.sample_template import build_sample_template
from cadencia.store.db import session_scope
from tests.support import NOW, collect_demo


@pytest.fixture
def deck(sessions, settings):
    asyncio.run(collect_demo(DemoJira(NOW, epic_plan=True), sessions, settings, select=(1, 2, 3, 4)))
    with session_scope(sessions) as session:
        save_demo_team(session)
    with session_scope(sessions) as session:
        return JiraSource(session, settings).deck_data([1, 4], reference=NOW, now=NOW, registry={})


def test_every_new_field_is_in_the_catalogue():
    for key in (
        "grafico.previsao",
        "grafico.cfd",
        "grafico.aging",
        "tendencia.vazao",
        "mudancas.novos_riscos",
        "plano.epicos",
        "grafico.gantt",
        "saude.problemas",
        "board.status_dados",
    ):
        assert key in FIELD_INDEX


def test_forecast_flow_and_trends_for_a_scrum_board(deck):
    board = deck.board_by_id(1)
    assert (
        board.texts["previsao.p85_data"]
        and board.numbers["num.previsao_p85_dias"] >= board.numbers["num.previsao_p50_dias"]
    )
    assert 0 <= board.numbers["num.previsao_chance_sprint"] <= 1
    forecast = board.series["grafico.previsao"]
    assert forecast.categories == ["50% de chance", "85% de chance", "95% de chance"]
    cfd = board.series["grafico.cfd"]
    assert cfd.style == "stacked_area" and next(iter(cfd.series)) == "Concluído"  # done is the bottom band
    assert all(len(values) == len(cfd.categories) for values in cfd.series.values())
    aging = board.series["grafico.aging"]
    assert aging.style == "scatter" and {m["label"] for m in aging.marks} >= {"CAD-11"}
    assert "num.vazao_janela_delta" in board.numbers and board.texts["tendencia.vazao"]


def test_changes_since_the_last_checkpoint_and_data_health(deck):
    board = deck.board_by_id(1)
    assert board.texts["mudancas.periodo"].startswith("desde ")
    assert "num.em_risco_delta" in board.numbers and board.texts["tendencia.em_risco"]
    assert all(
        set(r.values) == set(FIELD_INDEX["mudancas.novos_riscos"].columns)
        for r in board.rows["mudancas.novos_riscos"]
    )
    assert board.statuses["board.status_dados"] in {"em_dia", "risco", "atrasado"}
    assert board.numbers["num.saude_problemas"] == len(board.rows["saude.problemas"])


def test_the_epic_plan_reaches_the_deck(deck):
    board = deck.board_by_id(4)
    [row] = board.rows["plano.epicos"]
    assert row.values["chave"] == "ANL-1" and row.values["situacao_plano"] == "pronto"
    assert int(row.values["plano_dias"]) <= int(row.values["base_dias"])
    assert board.texts["plano.epico"].startswith("ANL-1") and board.numbers["num.plano_dias"] > 0
    gantt = board.series["grafico.gantt"]
    assert gantt.style == "gantt" and gantt.marks and all(m["lane"] in gantt.categories for m in gantt.marks)
    assert board.series["grafico.horas_por_papel"].categories


@pytest.mark.parametrize("style", ["bar", "stacked_area", "line", "scatter", "gantt"])
def test_every_chart_style_renders(tmp_path, style):
    series = Series(
        categories=["A", "B", "C"],
        series={"um": [1.0, 2.0, 3.0], "dois": [2.0, 1.0, 0.0]} if style != "gantt" else {"jornada": [6.0]},
        style=style,
        marks=[
            {"x": 0, "y": 3.0, "label": "X-1"},
            {"lane": "A", "start": 0, "end": 4, "color": "#4dabf7", "label": "S1"},
            {"lane": "B", "start": 4, "end": 10, "label": "S2"},
        ]
        if style in {"scatter", "gantt"}
        else [],
    )
    if style == "scatter":
        series.marks = series.marks[:1]
        series.series = {"P85": [4.0]}
    if style == "gantt":
        series.marks = series.marks[1:]
    out = render_chart_png(series, tmp_path / f"{style}.png", 5_000_000, 2_500_000)
    assert out.stat().st_size > 1000


def test_a_gantt_bound_to_a_native_chart_is_rendered_over_it(deck, tmp_path):
    template = build_sample_template(tmp_path / "modelo.pptx")
    manifest = extract_manifest(template)
    slide = manifest["slides"][1]
    chart = next(a for a in slide["areas"] if a["kind"] == "chart")
    mapping = DeckMapping.model_validate(
        {
            "slides": {
                str(slide["slide_id"]): {
                    "scope": "board",
                    "board_id": 4,
                    "bindings": [{"target": chart["id"], "mode": "chart", "field": "grafico.gantt"}],
                }
            }
        }
    )
    result = generate_deck(template, mapping, deck, tmp_path / "saida.pptx")
    assert not [p for p in result.problems if p.target == chart["id"]]
    produced = Presentation(str(tmp_path / "saida.pptx"))
    board_slide = produced.slides[1]
    assert not any(getattr(s, "has_chart", False) and s.has_chart for s in board_slide.shapes)
    assert any(s.shape_type == 13 for s in board_slide.shapes)  # MSO_SHAPE_TYPE.PICTURE


# ------------------------------------------------------------------ narrative suggested by AI (#7)


class _Block:
    def __init__(self, text):
        self.type, self.text = "text", text


class _FakeClient:
    def __init__(self, text="O time entregou 5 itens.", stop="end_turn"):
        self.calls = []
        client = self

        class _Messages:
            def create(self, **kwargs):
                client.calls.append(kwargs)
                return type("R", (), {"stop_reason": stop, "content": [_Block(text)]})()

        self.beta = type("B", (), {"messages": _Messages()})()


def test_the_narrative_sees_numbers_never_names(deck, settings):
    from cadencia.report.narrative import NarrativeUnavailable, draft_narrative, narrative_facts

    facts = narrative_facts(deck.board_by_id(1))
    assert facts["board"] == "Plataforma" and "em_risco" in facts
    flat = str(facts)
    assert not any(name in flat for name in ("Ana", "Bruno", "Carla", "Diego", "Elisa"))
    with pytest.raises(NarrativeUnavailable):
        draft_narrative(settings, facts, client=_FakeClient())  # off unless enabled on purpose
    enabled = settings.model_copy(update={"ai_enabled": True})
    client = _FakeClient()
    assert draft_narrative(enabled, facts, client=client) == "O time entregou 5 itens."
    [call] = client.calls
    assert call["model"] == "claude-opus-5" and call["fallbacks"] == "default"
    assert "Plataforma" in call["messages"][0]["content"]
    from cadencia.errors import Invalid

    with pytest.raises(Invalid):
        draft_narrative(enabled, facts, client=_FakeClient(stop="refusal"))
