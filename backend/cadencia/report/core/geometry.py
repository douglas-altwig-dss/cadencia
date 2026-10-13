"""Where a shape really is, and what outline it really has (RF-48).

Two things the old reader got wrong on any non-trivial deck:

**Group transforms.** A shape inside a group stores its position in the *group's* coordinate space,
not the slide's. The group carries ``off``/``ext`` (where it sits on the slide) and
``chOff``/``chExt`` (the space its children are written in); the child's slide position is the
child's own coordinates mapped through that pair, and nested groups compose. Reading ``shape.left``
for a grouped shape — which is what happened before — puts it wherever the group's author happened
to be working. The project's own template has 22 groups holding 49 shapes.

**Outlines.** ``prstGeom`` names a shape, it does not describe one, and ``custGeom`` describes one
in its own coordinate space. Drawing every shape as a rectangle turns the template's 21 connectors
into filled bars. Both are translated here into SVG path data, which is what the canvas draws and
what keeps a line a line.
"""

from __future__ import annotations

import math
from collections.abc import Iterator
from dataclasses import dataclass
from typing import Any

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


@dataclass(frozen=True)
class Transform:
    """Group space -> slide space. Composed as groups nest."""

    dx: int = 0
    dy: int = 0
    sx: float = 1.0
    sy: float = 1.0
    rotation: float = 0.0

    def apply(self, x: int, y: int, cx: int, cy: int) -> tuple[int, int, int, int]:
        return (
            round(self.dx + x * self.sx),
            round(self.dy + y * self.sy),
            max(round(cx * self.sx), 0),
            max(round(cy * self.sy), 0),
        )

    @property
    def font_scale(self) -> float:
        """Text inside a scaled group scales with it; the geometric mean handles a squashed group."""
        return math.sqrt(max(self.sx, 1e-6) * max(self.sy, 1e-6))

    def compose(self, inner: Transform) -> Transform:
        return Transform(
            dx=round(self.dx + inner.dx * self.sx),
            dy=round(self.dy + inner.dy * self.sy),
            sx=self.sx * inner.sx,
            sy=self.sy * inner.sy,
            rotation=self.rotation + inner.rotation,
        )


IDENTITY = Transform()


def group_transform(group: Any) -> Transform:
    """The child-space -> parent-space mapping a group declares."""
    xfrm = group._element.find(f"{A}xfrm")
    if xfrm is None:
        xfrm = group._element.find(f".//{A}xfrm")
    if xfrm is None:
        return IDENTITY
    off, ext = xfrm.find(f"{A}off"), xfrm.find(f"{A}ext")
    ch_off, ch_ext = xfrm.find(f"{A}chOff"), xfrm.find(f"{A}chExt")
    if off is None or ext is None or ch_off is None or ch_ext is None:
        return IDENTITY

    def number(element: Any, *names: str) -> int:
        for name in names:
            raw = element.get(name)
            if raw is not None:
                try:
                    return int(raw)
                except ValueError:
                    return 0
        return 0

    cw, ch = number(ch_ext, "cx"), number(ch_ext, "cy")
    sx = number(ext, "cx") / cw if cw else 1.0
    sy = number(ext, "cy") / ch if ch else 1.0
    rotation = int(xfrm.get("rot") or 0) / 60000
    return Transform(
        dx=round(number(off, "x") - number(ch_off, "x") * sx),
        dy=round(number(off, "y") - number(ch_off, "y") * sy),
        sx=sx,
        sy=sy,
        rotation=rotation,
    )


@dataclass
class Placed:
    """A shape with the geometry it has *on the slide*, whatever it is nested inside."""

    shape: Any
    x: int
    y: int
    cx: int
    cy: int
    rotation: float
    transform: Transform
    depth: int
    parent_id: int | None
    z: int


def iter_placed(
    shapes: Any, transform: Transform = IDENTITY, depth: int = 0, parent_id: int | None = None
) -> Iterator[Placed]:
    """Every shape in painting order, with its absolute slide geometry.

    Groups are yielded before their children — the canvas needs them to draw a selection box, and
    the order matches the ``spTree``, which is the painting order PowerPoint uses.
    """
    for counter, shape in enumerate(shapes):
        local_rot = float(shape.rotation or 0) if hasattr(shape, "rotation") else 0.0
        x, y, cx, cy = transform.apply(
            int(shape.left or 0), int(shape.top or 0), int(shape.width or 0), int(shape.height or 0)
        )
        yield Placed(
            shape=shape,
            x=x,
            y=y,
            cx=cx,
            cy=cy,
            rotation=local_rot + transform.rotation,
            transform=transform,
            depth=depth,
            parent_id=parent_id,
            z=counter,
        )
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_placed(
                shape.shapes, transform.compose(group_transform(shape)), depth + 1, shape.shape_id
            )


# ---------------------------------------------------------------- outlines

#: Presets drawn from their own formula, in a 0..1 box the canvas scales. Anything missing falls
#: back to a rectangle, which is what PowerPoint's own "unknown geometry" looks like anyway.
_PRESETS: dict[str, str] = {
    "rect": "M0,0 L1,0 L1,1 L0,1 Z",
    "roundRect": "",  # handled with the adjustment value
    "ellipse": "M0,0.5 A0.5,0.5 0 1,1 1,0.5 A0.5,0.5 0 1,1 0,0.5 Z",
    "triangle": "M0.5,0 L1,1 L0,1 Z",
    "rtTriangle": "M0,0 L0,1 L1,1 Z",
    "diamond": "M0.5,0 L1,0.5 L0.5,1 L0,0.5 Z",
    "parallelogram": "M0.25,0 L1,0 L0.75,1 L0,1 Z",
    "trapezoid": "M0.25,0 L0.75,0 L1,1 L0,1 Z",
    "pentagon": "M0.5,0 L1,0.38 L0.81,1 L0.19,1 L0,0.38 Z",
    "hexagon": "M0.25,0 L0.75,0 L1,0.5 L0.75,1 L0.25,1 L0,0.5 Z",
    "line": "M0,0 L1,1",
    "straightConnector1": "M0,0 L1,1",
    "bentConnector3": "M0,0 L0.5,0 L0.5,1 L1,1",
    "rightArrow": "M0,0.3 L0.6,0.3 L0.6,0 L1,0.5 L0.6,1 L0.6,0.7 L0,0.7 Z",
    "leftArrow": "M1,0.3 L0.4,0.3 L0.4,0 L0,0.5 L0.4,1 L0.4,0.7 L1,0.7 Z",
    "upArrow": "M0.3,1 L0.3,0.4 L0,0.4 L0.5,0 L1,0.4 L0.7,0.4 L0.7,1 Z",
    "downArrow": "M0.3,0 L0.3,0.6 L0,0.6 L0.5,1 L1,0.6 L0.7,0.6 L0.7,0 Z",
    "star5": (
        "M0.5,0 L0.62,0.35 L1,0.35 L0.69,0.57 L0.81,0.93 L0.5,0.71 L0.19,0.93 L0.31,0.57 L0,0.35 L0.38,0.35 Z"
    ),
    "chevron": "M0,0 L0.75,0 L1,0.5 L0.75,1 L0,1 L0.25,0.5 Z",
    "plus": (
        "M0.35,0 L0.65,0 L0.65,0.35 L1,0.35 L1,0.65 L0.65,0.65 "
        "L0.65,1 L0.35,1 L0.35,0.65 L0,0.65 L0,0.35 L0.35,0.35 Z"
    ),
}
#: Presets whose outline is a stroke, not an area: never filled, however the file states a fill.
STROKE_ONLY = {"line", "straightConnector1", "bentConnector3", "curvedConnector3", "bentConnector2"}


def _round_rect_path(radius: float) -> str:
    r = max(0.0, min(radius, 0.5))
    if r <= 0:
        return _PRESETS["rect"]
    return f"M{r},0 L{1 - r},0 Q1,0 1,{r} L1,{1 - r} Q1,1 {1 - r},1 L{r},1 Q0,1 0,{1 - r} L0,{r} Q0,0 {r},0 Z"


def _adjustments(shape: Any) -> list[float]:
    try:
        return [float(value) for value in shape.adjustments]
    except (AttributeError, NotImplementedError, TypeError, ValueError):
        return []


def _custom_path(shape: Any, cx: int, cy: int) -> str | None:
    """``custGeom`` translated to SVG, normalised into the shape's own 0..1 box."""
    try:
        cust = shape._element.spPr.find(f"{A}custGeom")
    except AttributeError:
        return None
    if cust is None:
        return None
    path = cust.find(f"{A}pathLst/{A}path")
    if path is None:
        return None
    width = int(path.get("w") or 0) or cx or 1
    height = int(path.get("h") or 0) or cy or 1

    def point(element: Any) -> tuple[float, float]:
        return (
            round(int(element.get("x") or 0) / width, 5),
            round(int(element.get("y") or 0) / height, 5),
        )

    out: list[str] = []
    for command in path:
        tag = command.tag[len(A) :]
        if tag == "moveTo":
            x, y = point(command.find(f"{A}pt"))
            out.append(f"M{x},{y}")
        elif tag == "lnTo":
            x, y = point(command.find(f"{A}pt"))
            out.append(f"L{x},{y}")
        elif tag == "cubicBezTo":
            points = [point(p) for p in command.findall(f"{A}pt")]
            if len(points) == 3:
                out.append("C" + " ".join(f"{x},{y}" for x, y in points))
        elif tag == "quadBezTo":
            points = [point(p) for p in command.findall(f"{A}pt")]
            if len(points) == 2:
                out.append("Q" + " ".join(f"{x},{y}" for x, y in points))
        elif tag == "close":
            out.append("Z")
    return " ".join(out) or None


def outline(shape: Any, cx: int, cy: int) -> dict[str, Any]:
    """The shape's outline as SVG path data in a 0..1 box, plus how it should be treated."""
    try:
        spPr = shape._element.spPr
    except AttributeError:
        spPr = None
    preset = None
    if spPr is not None:
        element = spPr.find(f"{A}prstGeom")
        preset = element.get("prst") if element is not None else None

    custom = _custom_path(shape, cx, cy)
    if custom:
        return {"preset": "custGeom", "path": custom, "stroke_only": False}
    if preset == "roundRect":
        adjust = _adjustments(shape)
        return {
            "preset": preset,
            "path": _round_rect_path(adjust[0] if adjust else 0.16),
            "stroke_only": False,
        }
    if preset in _PRESETS:
        return {"preset": preset, "path": _PRESETS[preset], "stroke_only": preset in STROKE_ONLY}
    if shape.shape_type == MSO_SHAPE_TYPE.LINE or (preset or "").endswith("Connector1"):
        return {"preset": preset or "line", "path": _PRESETS["line"], "stroke_only": True}
    return {"preset": preset, "path": _PRESETS["rect"], "stroke_only": False}


def line_shape(shape: Any) -> bool:
    """A connector: drawn as a stroke from corner to corner, never as a filled box."""
    if shape.shape_type == MSO_SHAPE_TYPE.LINE:
        return True
    try:
        element = shape._element.spPr.find(f"{A}prstGeom")
    except AttributeError:
        return False
    return element is not None and (element.get("prst") or "") in STROKE_ONLY


def flipped(shape: Any) -> tuple[bool, bool]:
    """``flipH``/``flipV``, which is how a connector points the other way."""
    try:
        xfrm = shape._element.spPr.find(f"{A}xfrm")
    except AttributeError:
        return False, False
    if xfrm is None:
        return False, False
    return xfrm.get("flipH") in ("1", "true"), xfrm.get("flipV") in ("1", "true")


def notes_of(slide: Any) -> str:
    return slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""


def shape_element_tag(shape: Any) -> str:
    return shape._element.tag.split("}")[-1]


def has_qn(shape: Any, tag: str) -> bool:
    return shape._element.find(qn(tag)) is not None
