"""Manual edits to a generated deck, and how they are replayed onto the next one (RF-46).

The module before this one answers "what does the data say"; this one answers "and then I moved
that box". The two are deliberately separate files: a generation always starts from the pristine
template, fills it with fresh data, and *then* replays these edits, so a re-run updates the numbers
without losing the layout work.

Three rules make the replay safe:

* **Absolute, never relative.** An edit says "this shape is at x=1143000", never "move it right by
  20". Applying the same edit twice lands in the same place, so the editor can show the result
  locally and the server can reach it again from scratch.
* **Keyed by origin, not by output position.** A slide is addressed by the template slide it came
  from, the board it was cloned for and its pagination page — the identity
  :func:`cadencia.report.core.generate.generate_deck` already records in ``produced``. Adding a board or
  a sprint shuffles the output order without moving anybody's edits.
* **Inserted shapes carry their key in their name.** ``cadencia:<key>`` survives a save in
  PowerPoint, so a box the user created can be found again on the next run.
"""

from __future__ import annotations

import copy
from collections.abc import Mapping
from typing import Any, Literal

from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE_TYPE
from pydantic import BaseModel, Field

from cadencia.report.core import pptx_utils as px

GROUP = MSO_SHAPE_TYPE.GROUP

EDITS_VERSION = 1
NAME_PREFIX = "cadencia:"
EMU_PER_PT = 12700

ShapeKind = Literal["textbox", "shape", "image"]
Preset = Literal["rect", "roundRect", "ellipse", "triangle", "diamond", "arrow", "line", "star"]


#: RF-64: the board and page of an edit made on the **template** rather than on one produced
#: slide. A template has neither — it is the design every copy starts from — so a template edit
#: is addressed to all of them and merged underneath whatever a particular copy says.
WILDCARD = "*"


def slide_key(origin_slide_id: int, board_id: int | None, page: int) -> str:
    """Stable identity of a produced slide: template slide, board, page."""
    return f"{origin_slide_id}:{board_id or 0}:{page}"


def template_slide_key(origin_slide_id: int) -> str:
    """Where an edit made on the template itself is stored: every board, every page."""
    return f"{origin_slide_id}:{WILDCARD}:{WILDCARD}"


def shape_key(shape: Any) -> str:
    """How an edit addresses a shape: its own id, or the key we stamped into an inserted one."""
    name = str(getattr(shape, "name", "") or "")
    if name.startswith(NAME_PREFIX):
        return name[len(NAME_PREFIX) :]
    return f"sh{shape.shape_id}"


# ---------------------------------------------------------------- the document


class Frame(BaseModel):
    """Geometry in EMU, straight from the OOXML, so the canvas and the file cannot disagree."""

    x: int | None = None
    y: int | None = None
    cx: int | None = None
    cy: int | None = None
    rotation: float | None = None
    #: Mirrors `<a:xfrm flipH/flipV>` — `None` leaves it as the template has it, same as every
    #: other field here.
    flip_h: bool | None = None
    flip_v: bool | None = None


class Shadow(BaseModel):
    """The one effect a deck actually uses (`properties.effects` reads the same shape back).
    `enabled=False` removes an existing shadow; the field is only present on the edit at all when
    someone has touched it, same as `fill`/`line`'s own "none"."""

    enabled: bool = True
    color: str = "000000"
    alpha: float = 0.6
    blur_pt: float = 4.0
    distance_pt: float = 3.0
    angle: float = 45.0


class Fill(BaseModel):
    type: Literal["none", "solid"] = "solid"
    color: str | None = None  # RRGGBB
    alpha: float | None = None  # 0..1, 1 is opaque


class Line(BaseModel):
    type: Literal["none", "solid"] = "solid"
    color: str | None = None
    width_pt: float | None = None
    dash: Literal["solid", "dash", "dot", "dashDot", "lgDash"] | None = None


class RunStyle(BaseModel):
    font: str | None = None
    size_pt: float | None = None
    bold: bool | None = None
    italic: bool | None = None
    underline: bool | None = None
    color: str | None = None
    #: `<a:highlight>` — the text-marker colour, RRGGBB. Empty string clears it; `None` leaves it.
    highlight: str | None = None
    #: A URL this run jumps to when clicked — Slides' "Insert > Link" on selected text. Empty
    #: string clears an existing link; ``None`` leaves it as it was.
    link: str | None = None


class RunEdit(BaseModel):
    """One stretch of text with its own look, so a bold word inside a sentence survives."""

    text: str = ""
    style: RunStyle = Field(default_factory=RunStyle)


class ParagraphEdit(BaseModel):
    text: str = ""
    #: When present, the paragraph is written run by run and ``text`` is only a summary.
    runs: list[RunEdit] | None = None
    align: Literal["left", "center", "right", "justify"] | None = None
    level: int = 0
    line_spacing: float | None = None  # multiple of the line height
    space_before_pt: float | None = None
    space_after_pt: float | None = None
    bullet: Literal["none", "char", "number"] | None = None
    bullet_char: str = "•"
    style: RunStyle = Field(default_factory=RunStyle)


class TextEdit(BaseModel):
    """``paragraphs is None`` keeps whatever the data put there; a list replaces it wholesale."""

    paragraphs: list[ParagraphEdit] | None = None
    word_wrap: bool | None = None
    autofit: Literal["none", "shrink", "resize"] | None = None
    anchor: Literal["top", "middle", "bottom"] | None = None
    insets_pt: dict[str, float] | None = None  # l, r, t, b
    columns: int | None = None


class ImageEdit(BaseModel):
    asset_id: str | None = None
    crop: dict[str, float] | None = None  # left, right, top, bottom as fractions


class ShapeEdit(BaseModel):
    op: Literal["update", "insert", "delete"] = "update"
    kind: ShapeKind | None = None  # insert only
    preset: Preset | None = None  # insert of kind "shape"
    frame: Frame | None = None
    z: int | None = None  # absolute index in the slide's shape tree
    fill: Fill | None = None
    line: Line | None = None
    shadow: Shadow | None = None
    corner_radius: float | None = None  # 0..0.5 of the shorter side
    text: TextEdit | None = None
    image: ImageEdit | None = None
    #: #15: what a screen reader says for this shape. "" clears it; None leaves the file's own.
    alt_text: str | None = None


class CellStyle(BaseModel):
    """RF-91: one table cell's look. Absolute like everything here, so a replay lands the same."""

    fill: str | None = None  # RRGGBB; "none" clears the fill
    bold: bool | None = None
    italic: bool | None = None
    color: str | None = None
    size_pt: float | None = None
    align: Literal["left", "center", "right", "justify"] | None = None


class TableEdit(BaseModel):
    """RF-91: column widths in EMU and per-cell styles keyed ``"<row>:<col>"``.

    Cell *text* stays in `SlideEdits.cells`, where it already was; this is the rest of what a table
    needs, kept per table so a data table re-filled with a different number of rows keeps its
    widths and header styling, and a style addressed to a row that no longer exists is reported.
    """

    col_widths: list[int] | None = None
    cells: dict[str, CellStyle] = Field(default_factory=dict)


class SlideEdits(BaseModel):
    shapes: dict[str, ShapeEdit] = Field(default_factory=dict)
    #: Speaker notes, replaced wholesale when set.
    notes: str | None = None
    #: Table cell text overrides, keyed ``"<shape key>:<row>:<col>"``.
    cells: dict[str, str] = Field(default_factory=dict)
    #: Groups to dissolve, so their children can be edited one by one.
    ungrouped: list[str] = Field(default_factory=list)
    #: Selection groups: a name the editor invented -> the shape keys that move and style together.
    #:
    #: They carry **no geometry**, which is the point. A real ``<p:grpSp>`` would give its children
    #: coordinates relative to the group, and this document is absolute overrides on purpose (a
    #: relative one drifts on the second replay). Members keep their own absolute frames, so this
    #: never reaches the .pptx: ``apply_shape_edits`` ignores it, and a deck generated from a
    #: document that holds nothing else is byte-for-byte the deck generated from an empty one.
    groups: dict[str, list[str]] = Field(default_factory=dict)
    #: RF-89: shapes the canvas will not let anyone drag, resize or delete, and shapes the editor
    #: hides while you work behind them. Both are about *editing* and, like `groups`, never reach
    #: the .pptx: a locked logo is still exactly where the template put it.
    locked: list[str] = Field(default_factory=list)
    hidden: list[str] = Field(default_factory=list)
    #: RF-91: table layout and cell styling, by table shape key.
    tables: dict[str, TableEdit] = Field(default_factory=dict)
    #: The slide's own background — Slides' "Background" button. ``None`` means "untouched, use
    #: whatever the layout/master already show"; a real `Fill` (including ``type="none"``, which
    #: writes an explicit "inherit") is a genuine override and therefore replays into the file.
    background: Fill | None = None

    def replays(self) -> bool:
        """Whether anything here changes the file (locks, hiding and groups do not)."""
        return bool(
            self.shapes
            or self.cells
            or self.ungrouped
            or self.tables
            or self.notes is not None
            or self.background is not None
        )


class NamedStyle(BaseModel):
    """#22: a look somebody named ("Título de seção", "Destaque") and applies with one click.

    Editor-only, like groups: applying it writes ordinary absolute run/paragraph/fill overrides
    into the shapes, so the file never needs to know a style existed."""

    name: str
    run: RunStyle = Field(default_factory=RunStyle)
    align: Literal["left", "center", "right", "justify"] | None = None
    fill: str | None = None  # RRGGBB for the shape; None leaves it alone


class DeckEdits(BaseModel):
    version: int = EDITS_VERSION
    slides: dict[str, SlideEdits] = Field(default_factory=dict)
    #: Final order of slide keys. Empty means "whatever the generator produced".
    order: list[str] = Field(default_factory=list)
    deleted: list[str] = Field(default_factory=list)
    #: New slide key -> the key it was duplicated from, in creation order.
    duplicates: dict[str, str] = Field(default_factory=dict)
    #: #22: the template's named styles, by an id the editor invented. Never replayed.
    styles: dict[str, NamedStyle] = Field(default_factory=dict)

    def for_slide(self, key: str) -> SlideEdits:
        """This slide's edits, with the template's underneath them (RF-64).

        A template edit belongs to the slide's *design* and a produced slide's edit belongs to one
        board's copy of it, so the copy wins shape by shape. Merging here rather than at every call
        site is what keeps `apply_shape_edits`, the validator and the canvas from each having to
        remember the rule.
        """
        own = self.slides.get(key)
        origin = key.split(":", 1)[0]
        shared = self.slides.get(f"{origin}:{WILDCARD}:{WILDCARD}") if ":" in key else None
        if shared is None:
            return own or SlideEdits()
        if own is None:
            return shared
        tables = {**shared.tables}
        for key, table in own.tables.items():
            base = tables.get(key)
            tables[key] = (
                table
                if base is None
                else TableEdit(
                    col_widths=table.col_widths if table.col_widths is not None else base.col_widths,
                    cells={**base.cells, **table.cells},
                )
            )
        return SlideEdits(
            shapes={**shared.shapes, **own.shapes},
            notes=own.notes if own.notes is not None else shared.notes,
            cells={**shared.cells, **own.cells},
            ungrouped=[*shared.ungrouped, *(u for u in own.ungrouped if u not in shared.ungrouped)],
            groups={**shared.groups, **own.groups},
            locked=[*shared.locked, *(k for k in own.locked if k not in shared.locked)],
            hidden=[*shared.hidden, *(k for k in own.hidden if k not in shared.hidden)],
            tables=tables,
            background=own.background if own.background is not None else shared.background,
        )

    def touched(self) -> set[str]:
        """Slides carrying any manual change — the trigger for replaying at all."""
        # `groups` is deliberately absent: a selection group changes nothing in the file, so a
        # document carrying only groups has nothing to replay.
        # So are `locked` and `hidden`, for the same reason.
        marked = {key for key, slide in self.slides.items() if slide.replays()}
        return marked | set(self.deleted) | set(self.duplicates)


# ---------------------------------------------------------------- applying


def _set_fill(target: Any, fill: Fill) -> None:
    if fill.type == "none":
        target.fill.background()
        return
    if fill.color:
        px.set_solid_fill(target, fill.color)
    if fill.alpha is not None and fill.alpha < 1:
        px.set_fill_alpha(target, fill.alpha)


def _set_line(shape: Any, line: Line) -> None:
    if line.type == "none":
        shape.line.fill.background()
        return
    if line.color:
        px.set_line_color(shape, line.color)
    if line.width_pt is not None:
        shape.line.width = round(line.width_pt * EMU_PER_PT)
    if line.dash:
        px.set_line_dash(shape, line.dash)


def _apply_frame(shape: Any, frame: Frame) -> None:
    if frame.x is not None:
        shape.left = int(frame.x)
    if frame.y is not None:
        shape.top = int(frame.y)
    if frame.cx is not None:
        shape.width = max(int(frame.cx), 1)
    if frame.cy is not None:
        shape.height = max(int(frame.cy), 1)
    if frame.rotation is not None:
        shape.rotation = float(frame.rotation)
    if frame.flip_h is not None:
        shape._element.get_or_add_xfrm().flipH = bool(frame.flip_h)
    if frame.flip_v is not None:
        shape._element.get_or_add_xfrm().flipV = bool(frame.flip_v)


def _apply_text(shape: Any, edit: TextEdit) -> None:
    if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
        return
    frame = shape.text_frame
    if edit.paragraphs is not None:
        px.write_paragraphs(frame, [p.model_dump() for p in edit.paragraphs])
    if edit.word_wrap is not None:
        frame.word_wrap = edit.word_wrap
    if edit.anchor is not None:
        px.set_anchor(frame, edit.anchor)
    if edit.autofit is not None:
        px.set_autofit(frame, edit.autofit)
    if edit.insets_pt is not None:
        px.set_insets(frame, edit.insets_pt)
    if edit.columns is not None:
        px.set_columns(frame, edit.columns)


def _insert_shape(slide: Any, key: str, edit: ShapeEdit, assets: Mapping[str, str]) -> Any | None:
    frame = edit.frame or Frame()
    box = (
        int(frame.x or 0),
        int(frame.y or 0),
        max(int(frame.cx or 914400), 1),
        max(int(frame.cy or 457200), 1),
    )
    if edit.kind == "image":
        path = assets.get((edit.image.asset_id if edit.image else None) or "")
        if not path:
            return None
        shape = slide.shapes.add_picture(path, *box)
    elif edit.kind == "shape" and edit.preset == "line":
        # A line is a connector (`<p:cxnSp>`), not an autoshape — corner to corner of its box,
        # the same "line" preset geometry `geometry.py` already reads back for `iter_placed`.
        x, y, cx, cy = box
        shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, y, x + cx, y + cy)
    elif edit.kind == "shape":
        shape = slide.shapes.add_shape(px.preset_shape(edit.preset or "rect"), *box)
        shape.text_frame.text = ""
    else:
        shape = slide.shapes.add_textbox(*box)
    shape.name = f"{NAME_PREFIX}{key}"
    return shape


def _apply_cells(slide: Any, edits: SlideEdits, existing: Mapping[str, Any]) -> list[str]:
    """Table cell text, addressed by shape and coordinates."""
    problems: list[str] = []
    for address, text in edits.cells.items():
        key, _, position = address.rpartition(":")
        row_key, _, column = key.rpartition(":")
        try:
            row_index, column_index = int(column), int(position)
        except ValueError:
            continue
        shape = existing.get(row_key)
        if shape is None or not getattr(shape, "has_table", False) or not shape.has_table:
            problems.append(f"A tabela {row_key} não existe mais neste slide.")
            continue
        table = shape.table
        if row_index >= len(table.rows) or column_index >= len(table.columns):
            problems.append(f"A célula {row_index + 1},{column_index + 1} saiu da tabela {row_key}.")
            continue
        px.set_frame_text(table.cell(row_index, column_index).text_frame, text)
    return problems


def _style_cell(cell: Any, style: CellStyle) -> None:
    if style.fill == "none":
        cell.fill.background()
    elif style.fill:
        px.set_solid_fill(cell, style.fill)
    for paragraph in cell.text_frame.paragraphs:
        if style.align:
            px.set_paragraph_align(paragraph, style.align)
        for run in paragraph.runs:
            px.set_run_style(
                run, bold=style.bold, italic=style.italic, color=style.color, size_pt=style.size_pt
            )


def _apply_tables(edits: SlideEdits, existing: Mapping[str, Any]) -> list[str]:
    """RF-91: column widths and cell styles, after the data (and any typed cell text) is in."""
    problems: list[str] = []
    for key, table_edit in edits.tables.items():
        shape = existing.get(key)
        if shape is None or not getattr(shape, "has_table", False) or not shape.has_table:
            problems.append(f"A tabela {key} não existe mais neste slide.")
            continue
        table = shape.table
        if table_edit.col_widths:
            for index, width in enumerate(table_edit.col_widths[: len(table.columns)]):
                if width and width > 0:
                    table.columns[index].width = int(width)
            shape.width = sum(int(c.width) for c in table.columns)
        for address, style in table_edit.cells.items():
            row, _, column = address.partition(":")
            try:
                row_index, column_index = int(row), int(column)
            except ValueError:
                continue
            if row_index >= len(table.rows) or column_index >= len(table.columns):
                problems.append(f"A célula {row_index + 1},{column_index + 1} saiu da tabela {key}.")
                continue
            _style_cell(table.cell(row_index, column_index), style)
    return problems


def apply_shape_edits(slide: Any, edits: SlideEdits, assets: Mapping[str, str]) -> list[str]:
    """Apply one slide's edits in place. Returns the problems worth telling the user about."""
    problems: list[str] = []
    # Ungrouping first: it changes which shapes exist, and every later edit addresses the result.
    for key in edits.ungrouped:
        group = next(
            (s for s, _p in px.iter_shapes(slide.shapes) if shape_key(s) == key and s.shape_type == GROUP),
            None,
        )
        if group is None:
            problems.append(f"O grupo {key} não existe mais neste slide.")
            continue
        px.ungroup(slide, group)

    existing = {shape_key(shape): shape for shape, _parent in px.iter_shapes(slide.shapes)}

    for key, edit in edits.shapes.items():
        if edit.op == "delete":
            shape = existing.get(key)
            if shape is not None:
                shape._element.getparent().remove(shape._element)
            continue
        shape = existing.get(key)
        if shape is None:
            if edit.op != "insert":
                problems.append(f"A forma {key} não existe mais neste slide; a edição foi ignorada.")
                continue
            shape = _insert_shape(slide, key, edit, assets)
            if shape is None:
                problems.append(f"A imagem da forma {key} não está mais disponível.")
                continue
        if edit.frame is not None:
            _apply_frame(shape, edit.frame)
        if edit.fill is not None:
            _set_fill(shape, edit.fill)
        if edit.line is not None:
            _set_line(shape, edit.line)
        if edit.shadow is not None:
            px.set_shadow(shape, edit.shadow.model_dump() if edit.shadow.enabled else None)
        if edit.corner_radius is not None:
            px.set_corner_radius(shape, edit.corner_radius)
        if edit.image is not None and edit.image.crop is not None:
            px.set_crop(shape, edit.image.crop)
        if edit.text is not None:
            _apply_text(shape, edit.text)
        if edit.alt_text is not None:
            px.set_alt_text(shape, edit.alt_text)

    problems += _apply_cells(slide, edits, existing)
    problems += _apply_tables(edits, existing)
    if edits.notes is not None:
        px.set_notes(slide, edits.notes)
    if edits.background is not None:
        _set_fill(slide.background, edits.background)

    # Z-order last: every shape exists by now, so the indices mean what the editor meant.
    ordered = [(k, e.z) for k, e in edits.shapes.items() if e.z is not None and e.op != "delete"]
    for key, z in sorted(ordered, key=lambda kv: kv[1] or 0):
        shape = next((s for s, _ in px.iter_shapes(slide.shapes) if shape_key(s) == key), None)
        if shape is not None:
            px.move_shape_to(slide, shape, int(z))
    return problems


def apply_deck_edits(
    prs: Any, produced: list[dict[str, Any]], edits: DeckEdits, assets: Mapping[str, str]
) -> tuple[list[str], dict[str, int]]:
    """Replay a whole edit document onto a freshly generated deck.

    Returns the problems and the slide key -> 1-based index map of the result, which is what the
    editor needs to line its scene up with the file.
    """
    problems: list[str] = []
    slides = list(prs.slides)
    by_key: dict[str, Any] = {}
    for entry in produced:
        key = slide_key(entry["origin_slide_id"], entry.get("board_id"), entry.get("page", 1))
        index = int(entry["index"]) - 1
        if 0 <= index < len(slides):
            by_key.setdefault(key, slides[index])

    # 1. Duplicates first, so a copy can be edited like any other slide.
    for new_key, source_key in edits.duplicates.items():
        source = by_key.get(source_key)
        if source is None:
            problems.append(f"O slide duplicado {new_key} perdeu a origem; não foi recriado.")
            continue
        copy_at = px.slide_index(prs, source) + 1
        by_key[new_key] = px.duplicate_slide(prs, source, index=copy_at)

    # 2. Shape edits, on originals and copies alike.
    for key, slide in list(by_key.items()):
        slide_edits = edits.for_slide(key)
        if slide_edits.replays():
            problems += apply_shape_edits(slide, slide_edits, assets)

    # 3. Deletions.
    for key in edits.deleted:
        slide = by_key.pop(key, None)
        if slide is not None:
            px.delete_slide(prs, slide)

    # 4. Explicit ordering. Keys the order does not mention keep their relative place at the end.
    if edits.order:
        wanted = [k for k in edits.order if k in by_key]
        for position, key in enumerate(wanted):
            px.move_slide(prs, by_key[key], position)

    return problems, {key: px.slide_index(prs, slide) + 1 for key, slide in by_key.items()}


def merge_edits(base: DeckEdits, patch: DeckEdits) -> DeckEdits:
    """Later wins, shape by shape — used when the editor saves only what it changed."""
    out = base.model_copy(deep=True)
    for key, slide in patch.slides.items():
        target = out.slides.setdefault(key, SlideEdits())
        target.shapes.update(copy.deepcopy(slide.shapes))
    if patch.order:
        out.order = list(patch.order)
    out.deleted = sorted(set(out.deleted) | set(patch.deleted))
    out.duplicates.update(patch.duplicates)
    return out
