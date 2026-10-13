"""Deck generation (RF-39 … RF-42).

The generator starts from the original template file and replaces content inside it; it never
draws slides. Order of work:

1. structure — drop slides marked "remove", clone "repeat" blocks once per board (consecutive
   repeat slides are cloned together, keeping template order);
2. pagination — tables whose rows do not fit get continuation copies of their slide;
3. bindings — every mapped area receives its field;
4. inline ``{{marcadores}}`` anywhere in the deck are replaced.
"""

from __future__ import annotations

import copy
import math
import tempfile
from collections.abc import Collection, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from pptx import Presentation

from cadencia.report.core import layout as lay
from cadencia.report.core import pptx_utils as px
from cadencia.report.core.charts import render_chart_png, replace_chart_data
from cadencia.report.core.context import (
    MISSING,
    RENDER_ONLY_STYLES,
    BoardDeck,
    DeckData,
    Group,
    Row,
    RowGroup,
    Series,
    resolve,
    spec_for,
    variable_name,
)
from cadencia.report.core.images import render_qr_png
from cadencia.report.core.manifest import parse_target
from cadencia.report.core.mapping import Binding, DeckMapping, SlideMapping
from cadencia.report.core.text import MARKER, MODIFIERS, replace_in_paragraph, replace_in_slide, split_marker

BOTTOM_MARGIN = 0.04
TABLE_MODES = {"table_rows", "grouped_rows"}
CONTINUED = " (cont.)"


@dataclass
class Problem:
    kind: str  # vazio | area_ausente | campo_desconhecido | marcador_desconhecido | sem_dados | truncado
    message: str
    slide: int | None = None
    target: str | None = None
    field: str | None = None

    def as_dict(self) -> dict[str, Any]:
        return self.__dict__.copy()


@dataclass
class GenerationResult:
    path: Path
    slide_count: int = 0
    problems: list[Problem] = field(default_factory=list)
    produced: list[dict[str, Any]] = field(default_factory=list)
    #: #30: template slides a ``hide_when_empty`` dropped, once per copy — what the usage view counts.
    dropped: list[int] = field(default_factory=list)


@dataclass
class _Placed:
    slide: Any
    origin_id: int
    mapping: SlideMapping
    board: BoardDeck | None
    page: int = 0


@dataclass
class _TableRow:
    kind: str  # "group" | "item"
    values: list[str | None]
    fills: dict[int, str]
    colors: dict[int, str] = field(default_factory=dict)  # RF-92: text colour per column


def generate_deck(
    template_path: str | Path,
    mapping: DeckMapping,
    deck: DeckData,
    out_path: str | Path,
    *,
    theme: dict[str, Any] | None = None,
    audience: str = "",
    only_slides: Collection[int] | None = None,
) -> GenerationResult:
    """Fill ``template_path``. ``audience`` (#25) keeps the slides meant for that variant;
    ``only_slides`` (#10) keeps just those template slides — a partial deck for a quick fix."""
    prs = Presentation(str(template_path))
    if audience or only_slides is not None:
        mapping = narrowed(mapping, [s.slide_id for s in prs.slides], audience, only_slides)
    result = GenerationResult(path=Path(out_path))
    placed = _structure(prs, mapping, deck, result)
    with tempfile.TemporaryDirectory(prefix="cadencia-deck-") as scratch:
        final: list[_Placed] = []
        for item in placed:
            final += _paginate_and_bind(prs, item, mapping, deck, result, theme or {}, Path(scratch))
        by_slide = {p.slide.slide_id: p for p in final}
        for number, slide in enumerate(prs.slides, start=1):
            context = by_slide.get(slide.slide_id)
            board = context.board if context else None
            origin = context.origin_id if context else None
            replace_in_slide(slide, _marker_resolver(deck, board, mapping, result, number, origin))
            # RF-65: the name goes into the file, so it survives the deck leaving Cadência. The
            # board's own name is appended on a repeated slide, which is what makes thirty copies
            # of "Resumo" findable in PowerPoint's selection pane.
            if origin is not None:
                given = mapping.for_slide(origin).name
                if given:
                    suffix = (
                        f" · {board.name}" if board and mapping.for_slide(origin).scope == "repeat" else ""
                    )
                    px.set_slide_name(slide, f"{given}{suffix}")
        result.produced = [
            {
                "index": px.slide_index(prs, p.slide) + 1,
                "origin_slide_id": p.origin_id,
                "board_id": p.board.board_id if p.board else None,
                "page": p.page + 1,
            }
            for p in final
        ]
        result.produced.sort(key=lambda r: r["index"])
        expand_notes(prs, result.produced, deck, mapping, result)
        result.slide_count = len(prs.slides)
        Path(out_path).parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out_path))
    return result


# ---------------------------------------------------------------- structure


def narrowed(
    mapping: DeckMapping, slide_ids: Sequence[int], audience: str, only_slides: Collection[int] | None
) -> DeckMapping:
    """A copy of ``mapping`` where every slide outside the audience or the chosen set is removed.

    Removal is the motor's own vocabulary, so blocks, scopes and pagination need nothing new."""
    slides = dict(mapping.slides)
    for slide_id in slide_ids:
        spec = mapping.for_slide(slide_id)
        outside = only_slides is not None and slide_id not in only_slides
        if outside or not spec.for_audience(audience):
            slides[str(slide_id)] = spec.model_copy(update={"action": "remove"})
    return mapping.model_copy(update={"slides": slides})


def repeat_targets(deck: DeckData, spec: SlideMapping) -> list[BoardDeck]:
    """Who the repeat block runs for: every board, or one board per space.

    Repeating per space is what a slide *about the space* wants — two boards of the same team
    would otherwise produce the same slide twice. RF-55: the slide's own scope filters the
    audience, which is how one template emits a different slide set per space or board.
    """
    pool = [b for b in deck.boards if spec.covers(b.space_key, b.board_id)]
    if spec.repeat_by != "espaco":
        return pool
    seen: dict[str, BoardDeck] = {}
    for board in pool:
        key = board.space_key or f"board:{board.board_id}"
        seen.setdefault(key, board)
    return list(seen.values())


def _out_of_scope(deck: DeckData, spec: SlideMapping, board: BoardDeck | None) -> bool:
    """RF-55: whether a non-repeating slide sits outside the scope somebody gave it.

    A board slide answers for its own board. A deck slide has no board, so it is kept when *any*
    board in this generation is in scope — a cover slide narrowed to one space belongs in a run
    that covers that space, and nowhere else.
    """
    if not spec.conditional:
        return False
    if board is not None:
        return not spec.covers(board.space_key, board.board_id)
    return not any(spec.covers(b.space_key, b.board_id) for b in deck.boards)


def _is_empty(deck: DeckData, board: BoardDeck | None, field: str) -> bool:
    kind, value = resolve(deck, board, field)
    # MISSING counts as empty here so the slide still disappears, but `_structure` reports the
    # name separately: hiding a slide because of a typo is the silent failure this rule exists
    # to stop.
    if kind is None or value is None or value is MISSING:
        return True
    if kind in {"text", "status"}:
        return not str(value).strip()
    if kind == "number":
        return not float(value)
    if kind == "row_groups":
        return not any(g.rows for g in value)
    if kind == "series":
        return not getattr(value, "categories", None)
    return not len(value)


def _structure(prs: Any, mapping: DeckMapping, deck: DeckData, result: GenerationResult) -> list[_Placed]:
    originals = list(prs.slides)
    boards = {b.board_id: b for b in deck.boards}
    # `hide_when_empty` swallows its field by design — an empty one makes the slide vanish — so an
    # unregistered name there would delete a slide and say nothing. Check the names up front.
    for index, original in enumerate(originals, start=1):
        field_key = mapping.for_slide(original.slide_id).hide_when_empty
        if field_key and spec_for(field_key, deck.variables) is None:
            result.problems.append(_unknown_problem(deck, field_key, index))
    placed: list[_Placed] = []
    i = 0
    while i < len(originals):
        slide = originals[i]
        spec = mapping.for_slide(slide.slide_id)
        if spec.action == "remove":
            px.delete_slide(prs, slide)
            i += 1
            continue
        if spec.scope == "repeat":
            block = [slide]
            j = i + 1
            while j < len(originals):
                nxt = mapping.for_slide(originals[j].slide_id)
                if nxt.action != "keep" or nxt.scope != "repeat":
                    break
                block.append(originals[j])
                j += 1
            audience = repeat_targets(deck, spec)
            if not audience:
                for s in block:
                    px.delete_slide(prs, s)
                why = (
                    "Nenhum board alcança o escopo deste bloco"
                    if spec.conditional
                    else "Nenhum board na geração"
                )
                result.problems.append(Problem("sem_dados", f"{why}: slides repetíveis removidos."))
            else:
                insert_at = px.slide_index(prs, block[-1]) + 1
                kept: list[tuple[Any, Any]] = []
                for position, board in enumerate(audience):
                    if position == 0:
                        targets = block
                    else:
                        targets = []
                        for source in block:
                            targets.append(px.duplicate_slide(prs, source, index=insert_at))
                            insert_at += 1
                    for target, source in zip(targets, block, strict=True):
                        source_spec = mapping.for_slide(source.slide_id)
                        # Each slide of the block is scoped on its own, so a block can run for a
                        # board while one of its slides sits that board out.
                        if not source_spec.covers(board.space_key, board.board_id):
                            kept.append((target, None))
                            continue
                        if source_spec.hide_when_empty and _is_empty(
                            deck, board, source_spec.hide_when_empty
                        ):
                            result.dropped.append(source.slide_id)
                            kept.append((target, None))
                            continue
                        placed.append(_Placed(target, source.slide_id, source_spec, board))
                for target, _ in kept:  # dropped after cloning, so indices stayed stable
                    px.delete_slide(prs, target)
            i = j
            continue
        board = None
        if spec.scope == "board":
            board = boards.get(spec.board_id or -1)
            if board is None:
                result.problems.append(
                    Problem(
                        "sem_dados",
                        f"Slide {i + 1}: o board {spec.board_id} não tem dados nesta geração.",
                        slide=i + 1,
                    )
                )
        if _out_of_scope(deck, spec, board):
            px.delete_slide(prs, slide)
            i += 1
            continue
        if spec.hide_when_empty and _is_empty(deck, board, spec.hide_when_empty):
            result.dropped.append(slide.slide_id)
            px.delete_slide(prs, slide)
            i += 1
            continue
        placed.append(_Placed(slide, slide.slide_id, spec, board))
        i += 1
    return placed


# ---------------------------------------------------------------- value helpers


@dataclass
class MarkerReport:
    """What went wrong inside a literal value, kept apart because the two are different faults.

    ``unregistered`` is a name the variable manager has never heard of — a typo, or a variable
    somebody forgot to register. ``unfilled`` is a registered variable that no target filled in and
    whose registration declares no default. Both are errors; neither may quietly become "".
    """

    unregistered: list[str] = field(default_factory=list)
    unfilled: list[str] = field(default_factory=list)

    def any(self) -> bool:
        return bool(self.unregistered or self.unfilled)


def expand_markers(
    text: str,
    deck: DeckData,
    board: BoardDeck | None,
    mapping: DeckMapping,
    report: MarkerReport | None = None,
    slide_id: int | None = None,
) -> str:
    """Replace ``{{campo}}`` inside a literal value, so a text written once on a slide becomes
    this space's words (RF-45).

    A name that does not resolve **stays visible**, whether it is unregistered or merely unfilled.
    That is the point: the marker survives into the deck, where ``validate`` refuses it, instead of
    collapsing to an empty string that looks like a deliberately blank line. What resolves to a
    real value — including a genuinely empty one somebody chose — is substituted as before.
    """

    def sub(name: str) -> str:
        kind, value, text = render_marker(deck, board, name, mapping, slide_id)
        if kind is None:
            if report is not None:
                report.unregistered.append(name)
            return f"{{{{{name}}}}}"
        if value is MISSING:
            if report is not None:
                report.unfilled.append(name)
            return f"{{{{{name}}}}}"
        return text or ""

    return MARKER.sub(lambda m: sub(m.group(1).strip()), text)


def _value(
    binding: Binding,
    deck: DeckData,
    board: BoardDeck | None,
    mapping: DeckMapping | None = None,
    report: MarkerReport | None = None,
    slide_id: int | None = None,
) -> tuple[str | None, Any]:
    if binding.value is not None:
        text = binding.value
        if mapping is not None and "{{" in text:
            text = expand_markers(text, deck, board, mapping, report, slide_id)
        return "text", text
    if not binding.key:
        return None, None
    return resolve(deck, board, binding.key, slide_id)


def render_marker(
    deck: DeckData, board: BoardDeck | None, raw: str, mapping: DeckMapping, slide_id: int | None = None
) -> tuple[str | None, Any, str | None]:
    """``(kind, value, text)`` for one marker, modifier included (RF-94).

    ``kind is None`` is an unknown name *or* an unknown modifier — both leave the marker visible so
    `validate` refuses it; ``value is MISSING`` is a registered variable with nothing to say.
    """
    base, modifier = split_marker(raw)
    if modifier is not None and modifier not in MODIFIERS:
        return None, None, None
    kind, value = resolve(deck, board, base, slide_id)
    if kind is None or value is MISSING:
        return kind, value, None
    return kind, value, _as_text(kind, value, mapping, modifier)


def _all_rows(kind: str | None, value: Any) -> list[Row]:
    if kind == "rows":
        return list(value or ())
    if kind == "row_groups":
        return [r for g in value or () for r in g.rows]
    return []


def _as_text(kind: str | None, value: Any, mapping: DeckMapping, modifier: str | None = None) -> str:
    if modifier is not None and value not in (None, MISSING):
        rows = _all_rows(kind, value)
        if modifier == "contagem":
            if kind in {"rows", "row_groups"}:
                return str(len(rows))
            if kind == "groups":
                return str(sum(len(g.items) for g in value or ()))
        elif modifier == "chaves" and kind in {"rows", "row_groups"}:
            return ", ".join(k for r in rows if (k := r.values.get("chave", "")))
        elif modifier == "lista" and kind in {"rows", "row_groups"}:
            return "\n".join(
                f"{r.values.get('chave', '')} {r.values.get('resumo', '')}".strip() for r in rows
            )
        elif modifier == "lista" and kind == "groups":
            return "\n".join(item for g in value or () for item in g.items)
    if value is MISSING:
        # Only reachable if a caller forgot to check; make it visible rather than blank.
        raise AssertionError("MISSING chegou a _as_text: o chamador tem de recusar antes")
    if value is None:
        return ""
    if kind == "status":
        return mapping.label(str(value))
    if kind == "groups":
        return "\n".join(f"{g.title} {', '.join(g.items)}" for g in value)
    if kind == "rows":
        return "\n".join(r.values.get("resumo", "") for r in value)
    return str(value)


def _variable_label(deck: DeckData, name: str) -> str:
    """How an error refers to the name: the marker, and the template it appeared in."""
    where = f" no modelo “{deck.template_name}”" if deck.template_name else ""
    return f"{{{{{name}}}}}{where}"


def _unknown_problem(deck: DeckData, name: str, number: int, target: str | None = None) -> Problem:
    """A name nothing in the catalogue or the registry answers to.

    For a ``texto.*`` name this is precisely "not registered": the registry is the only thing that
    makes a variable exist, so the message says how to fix it rather than just that it is unknown.
    """
    if variable_name(name) is not None:
        message = (
            f"Slide {number}: a variável {_variable_label(deck, name)} não está registrada. "
            "Registre-a no gerenciador de variáveis antes de usá-la no modelo."
        )
        kind = "variavel_nao_registrada"
    else:
        message = f"Slide {number}: o campo {name} não existe no catálogo."
        kind = "marcador_desconhecido"
    return Problem(kind, message, slide=number, target=target, field=name)


def _unfilled_problem(
    deck: DeckData, board: BoardDeck | None, name: str, number: int, target: str | None = None
) -> Problem:
    """Registered, but nobody ever said what it means here and the registration has no default."""
    where = (board.space_name or board.space_key or board.name) if board else "o relatório"
    return Problem(
        "variavel_sem_valor",
        f"Slide {number}: a variável {_variable_label(deck, name)} não tem valor para “{where}” "
        "e o registro não define um padrão.",
        slide=number,
        target=target,
        field=name,
    )


def _marker_resolver(
    deck: DeckData,
    board: BoardDeck | None,
    mapping: DeckMapping,
    result: GenerationResult,
    number: int,
    slide_id: int | None = None,
) -> Any:
    def resolve_marker(name: str) -> str:
        kind, value, rendered = render_marker(deck, board, name, mapping, slide_id)
        if kind is None and split_marker(name)[1] not in (None, *MODIFIERS):
            result.problems.append(
                Problem(
                    "marcador_desconhecido",
                    f"Slide {number}: {{{{{name}}}}} usa um modificador desconhecido; "
                    f"use {', '.join(MODIFIERS)}.",
                    slide=number,
                    field=name,
                )
            )
            return f"{{{{{name}}}}}"
        if kind is None:
            result.problems.append(_unknown_problem(deck, name, number))
            # Left in place on purpose: `validate` refuses a leftover marker, so a deck that
            # cannot be filled honestly cannot be delivered looking finished.
            return f"{{{{{name}}}}}"
        if value is MISSING:
            result.problems.append(_unfilled_problem(deck, board, name, number))
            return f"{{{{{name}}}}}"
        text = rendered or ""
        if not text:
            result.problems.append(
                Problem("vazio", f"Marcador {{{{{name}}}}} ficou vazio.", slide=number, field=name)
            )
        return text

    return resolve_marker


def expand_notes(
    prs: Any, produced: list[dict[str, Any]], deck: DeckData, mapping: DeckMapping, result: GenerationResult
) -> None:
    """RF-94: fill ``{{markers}}`` in speaker notes, with each slide's own board.

    Notes are written by the template and, later, by the manual edit layer; this runs after both,
    so a presenter can write "{{riscos.itens|contagem}} itens em risco: {{riscos.itens|chaves}}"
    once and have it say this week's numbers every week.
    """
    slides = list(prs.slides)
    boards = {b.board_id: b for b in deck.boards}
    for entry in produced:
        index = int(entry["index"]) - 1
        if not 0 <= index < len(slides) or not slides[index].has_notes_slide:
            continue
        frame = slides[index].notes_slide.notes_text_frame
        if "{{" not in frame.text:
            continue
        board = boards.get(entry.get("board_id"))
        resolver = _marker_resolver(deck, board, mapping, result, index + 1, entry["origin_slide_id"])
        for paragraph in frame.paragraphs:
            replace_in_paragraph(paragraph, resolver)


# ---------------------------------------------------------------- tables


def _table_rows(
    binding: Binding, kind: str | None, value: Any, mapping: DeckMapping, table: Any
) -> list[_TableRow]:
    columns = list(binding.columns or [])
    width = len(table.columns)
    columns += [None] * (width - len(columns))

    def item(row: Row) -> _TableRow:
        values: list[str | None] = []
        for name in columns[:width]:
            if name is None:
                values.append("")
            elif name == "situacao":
                values.append(mapping.label(row.status))
            else:
                # Only the summary carries the reader's prose, so only it gets clipped and cased;
                # a key or a date must stay verbatim to remain a key or a date.
                text = row.values.get(name, "")
                values.append(lay.decorate(text, binding) if name == "resumo" else text)
        fills = {}
        color = mapping.color(row.status)
        if binding.status_column is not None and color:
            fills[binding.status_column] = color
        colors: dict[int, str] = {}
        for rule in binding.rules:
            tested = rule.column if rule.column is not None else binding.status_column
            if tested is None or tested >= width or not lay.rule_matches_row(rule, values, values[tested]):
                continue
            target = fills if rule.apply == "preenchimento" else colors
            for col in range(width) if rule.whole_row else [tested]:
                target.setdefault(col, rule.color)
        return _TableRow("item", values, fills, colors)

    out: list[_TableRow] = []
    if kind == "rows":
        out = [item(r) for r in lay.shape_rows(value, binding)]
    elif kind == "row_groups":
        title_col = _title_column(binding, table, columns)
        for group in lay.shape_row_groups(value, binding):
            header: list[str | None] = [""] * width
            header[title_col] = group.title
            out.append(_TableRow("group", header, {}))
            out += [item(r) for r in group.rows]
    if binding.scales:
        items = [r for r in out if r.kind == "item"]
        for row, marks in zip(items, lay.scale_marks([r.values for r in items], binding.scales), strict=True):
            for col, mark in marks.items():
                if "text" in mark:
                    row.values[col] = mark["text"]
                if "fill" in mark:
                    row.fills[col] = mark["fill"]
    return out


def _title_column(binding: Binding, table: Any, columns: list[str | None]) -> int:
    rows = list(table.rows)
    proto = binding.group_row if binding.group_row is not None else binding.header_rows
    if proto < len(rows):
        for index, cell in enumerate(rows[proto].cells):
            if cell.text.strip():
                return index
    return columns.index("resumo") if "resumo" in columns else 0


def _prototypes(binding: Binding, table: Any) -> dict[str, Any]:
    trs = list(table._tbl.tr_lst)
    header = binding.header_rows
    if len(trs) <= header:  # header-only table: data rows imitate the last header row
        base = trs[-1]
        return {"group": copy.deepcopy(base), "item": copy.deepcopy(base)}
    if binding.mode == "grouped_rows":
        group_index = binding.group_row if binding.group_row is not None else header
        item_index = binding.item_row if binding.item_row is not None else min(group_index + 1, len(trs) - 1)
    else:
        item_index = binding.item_row if binding.item_row is not None else header
        group_index = item_index
    return {
        "group": copy.deepcopy(trs[min(group_index, len(trs) - 1)]),
        "item": copy.deepcopy(trs[min(item_index, len(trs) - 1)]),
    }


def _row_height(table: Any, proto: Any, row: _TableRow) -> int:
    probe = copy.deepcopy(proto)
    scratch = copy.deepcopy(table._tbl)
    for tr in list(scratch.tr_lst):
        scratch.remove(tr)
    scratch.append(probe)
    from pptx.table import Table

    shadow = Table(scratch, table._graphic_frame)
    px.rebuild_rows(shadow, 0, [(proto, row.values, {})])
    return px.estimate_row_height(shadow, scratch.tr_lst[0], int(proto.get("h") or 0))


def _paginate(
    rows: list[_TableRow], shape: Any, table: Any, protos: dict[str, Any], header_rows: int, slide_height: int
) -> list[list[_TableRow]]:
    heights = px.row_heights(table)
    available = slide_height * (1 - BOTTOM_MARGIN) - int(shape.top or 0) - sum(heights[:header_rows])
    pages: list[list[_TableRow]] = [[]]
    used = 0
    current_group: _TableRow | None = None
    for index, row in enumerate(rows):
        h = _row_height(table, protos[row.kind], row)
        if row.kind == "group":
            current_group = row
            nxt = rows[index + 1] if index + 1 < len(rows) else None
            if nxt is not None:  # never leave a group header alone at the bottom of a page
                h_next = _row_height(table, protos[nxt.kind], nxt)
                if pages[-1] and used + h + h_next > available:
                    pages.append([])
                    used = 0
        if pages[-1] and used + h > available:
            pages.append([])
            used = 0
            if row.kind == "item" and current_group is not None:
                cont = copy.deepcopy(current_group)
                cont.values = [v + CONTINUED if v else v for v in cont.values]
                pages[-1].append(cont)
                used += _row_height(table, protos["group"], cont)
        pages[-1].append(row)
        used += h
    return pages


# ---------------------------------------------------------------- binding application


def _paginate_and_bind(
    prs: Any,
    item: _Placed,
    mapping: DeckMapping,
    deck: DeckData,
    result: GenerationResult,
    theme: dict[str, Any],
    scratch: Path,
) -> list[_Placed]:
    slide = item.slide
    number = px.slide_index(prs, slide) + 1
    table_pages: dict[int, list[list[_TableRow]]] = {}
    prototypes: dict[int, dict[str, Any]] = {}
    for index, binding in enumerate(item.mapping.bindings):
        if binding.mode not in TABLE_MODES:
            continue
        shape = px.find_shape(slide, parse_target(binding.target).shape_id)
        if shape is None or not getattr(shape, "has_table", False) or not shape.has_table:
            continue
        kind, value = _value(binding, deck, item.board, mapping, slide_id=item.origin_id)
        rows = _table_rows(binding, kind, value, mapping, shape.table)
        if not any(r.kind == "item" for r in rows):
            # Keep one blank row so the table still looks like the template's, and say so.
            rows = [_TableRow("item", [""] * len(shape.table.columns), {})]
            result.problems.append(
                Problem(
                    "vazio",
                    f"Slide {number}: {binding.key} não tem linhas.",
                    slide=number,
                    target=binding.target,
                    field=binding.key,
                )
            )
        prototypes[index] = _prototypes(binding, shape.table)
        pages = _paginate(rows, shape, shape.table, prototypes[index], binding.header_rows, prs.slide_height)
        if item.mapping.overflow == "truncate" and len(pages) > 1:
            dropped = sum(len(p) for p in pages[1:])
            result.problems.append(
                Problem(
                    "truncado",
                    f"Slide {number}: {dropped} linha(s) não couberam e foram omitidas.",
                    slide=number,
                    target=binding.target,
                )
            )
            pages = pages[:1]
        table_pages[index] = pages

    count = max((len(p) for p in table_pages.values()), default=1)
    copies = [item]
    position = px.slide_index(prs, slide) + 1
    for page in range(1, count):
        dup = px.duplicate_slide(prs, slide, index=position)
        position += 1
        copies.append(_Placed(dup, item.origin_id, item.mapping, item.board, page))

    for placed in copies:
        slide_number = px.slide_index(prs, placed.slide) + 1
        for index, binding in enumerate(item.mapping.bindings):
            if binding.mode in TABLE_MODES and index in table_pages:
                pages = table_pages[index]
                rows = pages[placed.page] if placed.page < len(pages) else []
                _fill_table(placed.slide, binding, rows, prototypes[index])
                continue
            _apply(
                placed.slide,
                binding,
                deck,
                placed.board,
                mapping,
                result,
                slide_number,
                theme,
                scratch,
                placed.origin_id,
            )
    return copies


def _fill_table(slide: Any, binding: Binding, rows: list[_TableRow], protos: dict[str, Any]) -> None:
    shape = px.find_shape(slide, parse_target(binding.target).shape_id)
    if shape is None:
        return
    table = shape.table
    px.rebuild_rows(
        table,
        binding.header_rows,
        [(protos[r.kind], r.values, r.fills) for r in rows],
        colors=[r.colors for r in rows],
    )
    heights = []
    for tr in table._tbl.tr_lst:
        heights.append(px.estimate_row_height(table, tr, int(tr.get("h") or 0)))
    shape.height = sum(heights)


def _as_number(kind: str | None, value: Any) -> float | None:
    """A field as a number, whatever kind it came from: counts for lists, the value for numbers."""
    if kind == "number":
        return float(value) if value is not None else None
    if kind in {"rows", "groups"}:
        return float(len(value or ()))
    if kind == "row_groups":
        return float(sum(len(g.rows) for g in (value or ())))
    if value is None or value == "":
        return None
    try:
        return float(str(value).replace("%", "").replace(",", ".").strip())
    except ValueError:
        return None


def _style_frame(frame: Any, binding: Binding) -> None:
    """The layout options that are about the box rather than the words."""
    px.set_alignment(frame, binding.align)
    px.set_columns(frame, binding.text_columns)
    px.set_shrink_to_fit(frame, binding.shrink_to_fit)


def _text_target(shape: Any, binding: Binding) -> tuple[Any | None, Any | None]:
    target = parse_target(binding.target)
    if target.is_cell:
        if not (getattr(shape, "has_table", False) and shape.has_table):
            return None, None
        table = shape.table
        if target.row >= len(table.rows) or target.col >= len(table.columns):
            return None, None
        frame = table.cell(target.row, target.col).text_frame
    elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        frame = shape.text_frame
    else:
        return None, None
    if target.paragraph is not None:
        paragraphs = frame.paragraphs
        if target.paragraph >= len(paragraphs):
            return frame, None
        return frame, paragraphs[target.paragraph]
    return frame, None


def _apply_rule(shape: Any, target: Any, frame: Any, paragraph: Any, binding: Binding, text: str) -> None:
    """RF-92 on a single value: colour the words, or fill the box (or the cell) they sit in."""
    rule = lay.first_rule(binding.rules, text)
    if rule is None:
        return
    if rule.apply == "preenchimento":
        if target.is_cell:
            px.set_solid_fill(shape.table.cell(target.row, target.col), rule.color)
        else:
            px.set_solid_fill(shape, rule.color)
    elif paragraph is not None:
        for run in paragraph.runs:
            px.set_run_style(run, color=rule.color)
    else:
        px.set_text_color(frame, rule.color)


def _apply(
    slide: Any,
    binding: Binding,
    deck: DeckData,
    board: BoardDeck | None,
    mapping: DeckMapping,
    result: GenerationResult,
    number: int,
    theme: dict[str, Any],
    scratch: Path,
    origin_id: int | None = None,
) -> None:
    target = parse_target(binding.target)
    shape = px.find_shape(slide, target.shape_id)
    if shape is None:
        result.problems.append(
            Problem(
                "area_ausente",
                f"Slide {number}: a área {binding.target} não existe no modelo.",
                slide=number,
                target=binding.target,
            )
        )
        return
    if binding.key and binding.value is None and spec_for(binding.key, deck.variables) is None:
        result.problems.append(_unknown_problem(deck, binding.key, number, binding.target))
        return
    report = MarkerReport()
    kind, value = _value(binding, deck, board, mapping, report, origin_id)
    for name in dict.fromkeys(report.unregistered):
        result.problems.append(_unknown_problem(deck, name, number, binding.target))
    for name in dict.fromkeys(report.unfilled):
        result.problems.append(_unfilled_problem(deck, board, name, number, binding.target))
    if value is MISSING:
        # A bound variable with nothing to say. The marker goes into the shape instead of an empty
        # string: the deck is refused either way, and this way it is refused *visibly*.
        result.problems.append(_unfilled_problem(deck, board, binding.key or "", number, binding.target))
        kind, value = "text", f"{{{{{binding.key}}}}}"

    def empty() -> None:
        result.problems.append(
            Problem(
                "vazio",
                f"Slide {number}: {binding.key or 'valor'} ficou vazio.",
                slide=number,
                target=binding.target,
                field=binding.key,
            )
        )

    if binding.mode in {"text", "status_text", "number"}:
        if binding.mode == "status_text":
            text = mapping.label(value)
        elif binding.mode == "number":
            text = lay.format_number(_as_number(kind, value), binding)
        else:
            text = _as_text(kind, value, mapping)
        text = lay.decorate(text, binding) or binding.empty_text
        frame, paragraph = _text_target(shape, binding)
        if frame is None:
            result.problems.append(
                Problem(
                    "area_ausente",
                    f"Slide {number}: {binding.target} não aceita texto.",
                    slide=number,
                    target=binding.target,
                )
            )
            return
        if paragraph is not None:
            px.set_paragraph_text(paragraph, text)
        else:
            px.set_frame_text(frame, text)
        _style_frame(frame, binding)
        _apply_rule(shape, target, frame, paragraph, binding, text)
        if not text:
            empty()
    elif binding.mode in {"bullet_list", "inline_list"}:
        frame, paragraph = _text_target(shape, binding)
        if frame is None:
            return
        lines = lay.flatten(value, kind, binding)
        if binding.mode == "inline_list":
            joined = binding.separator.join(lines) or binding.empty_text
            if paragraph is not None:
                px.set_paragraph_text(paragraph, joined)
            else:
                px.set_frame_text(frame, joined)
        else:
            px.fill_lines(frame, lines, bullet=binding.bullet, empty_text=binding.empty_text)
        _style_frame(frame, binding)
        if not lines:
            empty()
    elif binding.mode == "grouped_list":
        frame, _ = _text_target(shape, binding)
        if frame is None:
            return
        groups: list[Group] = lay.shape_groups(value, binding) if kind == "groups" else []
        px.fill_grouped_list(
            frame,
            [(g.title, g.items) for g in groups],
            header_paragraph=binding.header_paragraph,
            item_paragraph=binding.item_paragraph,
            empty_text=binding.empty_text or "—",
        )
        _style_frame(frame, binding)
        if not groups:
            empty()
    elif binding.mode == "status_fill":
        color = mapping.color(str(value)) if value else None
        if not color:
            empty()
            return
        if target.is_cell and getattr(shape, "has_table", False) and shape.has_table:
            px.set_solid_fill(shape.table.cell(target.row, target.col), color)
        else:
            px.set_solid_fill(shape, color)
    elif binding.mode == "chart":
        if not isinstance(value, Series) or not value.categories:
            empty()
            return
        description = binding.alt_text or chart_description(binding.key, value, deck)
        if getattr(shape, "has_chart", False) and shape.has_chart and value.style not in RENDER_ONLY_STYLES:
            replace_chart_data(shape.chart, value)
            px.set_alt_text(shape, description)
        else:
            colors = [theme.get("colors", {}).get(f"accent{i}") for i in range(1, 7)]
            png = scratch / f"grafico-{slide.slide_id}-{shape.shape_id}.png"
            render_chart_png(
                value,
                png,
                int(shape.width),
                int(shape.height),
                colors=[c for c in colors if c],
                font=theme.get("fonts", {}).get("minor"),
            )
            picture = px.place_picture(slide, shape, str(png))
            px.set_alt_text(picture, description)
    elif binding.mode == "qr_code":
        text = lay.decorate(_as_text(kind, value, mapping), binding) or binding.empty_text
        if not text.strip():
            empty()
            return
        png = scratch / f"qr-{slide.slide_id}-{shape.shape_id}.png"
        render_qr_png(text, png, error=binding.qr_error)
        # Fitted, not stretched: a QR out of square stops scanning.
        picture = px.place_picture(slide, shape, str(png), fit=True)
        px.set_alt_text(picture, binding.alt_text or f"Código QR: {text}")


def chart_description(key: str | None, series: Series, deck: DeckData, most: int = 12) -> str:
    """#15: a chart said in words, for a screen reader — its name, then category: value pairs."""
    spec = spec_for(key or "", deck.variables) if key else None
    title = spec.label if spec else "Gráfico"
    if series.style == "gantt":
        lanes = ", ".join(series.categories[:most])
        return f"{title}: {len(series.marks)} barras em {len(series.categories)} linhas ({lanes})."
    if series.style == "scatter":
        points = "; ".join(f"{m.get('label', '')}: {m['y']:g}" for m in series.marks[:most])
        return f"{title}: {points}."
    parts = []
    for name, values in series.series.items():
        pairs = ", ".join(f"{c}: {v:g}" for c, v in list(zip(series.categories, values, strict=False))[:most])
        parts.append(f"{name} — {pairs}" if len(series.series) > 1 else pairs)
    return f"{title}: {'; '.join(parts)}."


def group_rows(groups: list[RowGroup]) -> int:
    return sum(len(g.rows) + 1 for g in groups)


def pages_for(rows: int, per_page: int) -> int:
    return max(1, math.ceil(rows / max(per_page, 1)))
