"""A small sample template used by the demo and the tests.

It deliberately contains what real templates contain: a marker split across three runs, labelled
text boxes ("SEMANA" / value), a status indicator shape with a legend, 1×1 tables used as text
areas, a data table with a group-header prototype row, a native chart and a plain shape that a
chart image must replace.
"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

INK = RGBColor(0x14, 0x20, 0x1D)
ACCENT = RGBColor(0x0E, 0x5C, 0x5B)


def _bullet(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    ppr.set("marL", "228600")
    ppr.set("indent", "-228600")
    bu = etree.SubElement(ppr, qn("a:buChar"))
    bu.set("char", "•")


def _label(slide, left, top, title: str, value: str | list[str], width=None):
    box = slide.shapes.add_textbox(left, top, width or Inches(3), Inches(0.8))
    frame = box.text_frame
    frame.paragraphs[0].text = title
    frame.paragraphs[0].runs[0].font.bold = True
    frame.paragraphs[0].runs[0].font.size = Pt(12)
    frame.paragraphs[0].runs[0].font.color.rgb = ACCENT
    second = frame.add_paragraph()
    for piece in value if isinstance(value, list) else [value]:
        run = second.add_run()
        run.text = piece
        run.font.size = Pt(16)
        run.font.color.rgb = INK
    return box


def _caption(slide, left, top, text: str, width=None):
    box = slide.shapes.add_textbox(left, top, width or Inches(5.8), Inches(0.4))
    box.text_frame.paragraphs[0].text = text
    run = box.text_frame.paragraphs[0].runs[0]
    run.font.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = ACCENT
    return box


def _fill(cell, hex_color: str) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor.from_string(hex_color)


def build_sample_template(path: str | Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 — cover, with a marker split in three runs (the first one bold)
    cover = prs.slides.add_slide(blank)
    title = cover.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(10), Inches(1.2))
    title.text_frame.paragraphs[0].text = "Acompanhamento semanal"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    title.text_frame.paragraphs[0].runs[0].font.bold = True
    _label(cover, Inches(0.8), Inches(3.6), "SEMANA", ["{{relatorio.", "sem", "ana}}"])
    cover.shapes[-1].text_frame.paragraphs[1].runs[0].font.bold = True

    # 2 — board slide (one per board)
    board_slide = prs.slides.add_slide(blank)
    _label(board_slide, Inches(0.5), Inches(0.3), "PROJETO", "Nome do espaço")
    _label(board_slide, Inches(4.0), Inches(0.3), "SPRINT", "{{sprint.nome}}")
    _label(board_slide, Inches(7.5), Inches(0.3), "SEMANA", "00/00/0000")
    status_box = board_slide.shapes.add_textbox(Inches(10.6), Inches(0.35), Inches(1.6), Inches(0.4))
    status_box.text_frame.paragraphs[0].text = "Status geral"
    indicator = board_slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(12.3), Inches(0.35), Inches(0.5), Inches(0.4)
    )
    indicator.fill.solid()
    indicator.fill.fore_color.rgb = RGBColor(0x80, 0x80, 0x80)
    indicator.line.fill.background()

    _caption(board_slide, Inches(0.5), Inches(1.4), "Entregas da semana")
    done = board_slide.shapes.add_table(1, 1, Inches(0.5), Inches(1.85), Inches(6), Inches(2.2)).table
    frame = done.cell(0, 0).text_frame
    frame.paragraphs[0].text = "Épico:"
    item = frame.add_paragraph()
    item.text = "Item entregue"
    _bullet(item)
    frame.add_paragraph()
    _fill(done.cell(0, 0), "FFFFFF")
    for p in frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(12)
            r.font.color.rgb = INK

    _caption(board_slide, Inches(0.5), Inches(4.4), "Bloqueios")
    blockers = board_slide.shapes.add_table(2, 3, Inches(0.5), Inches(4.85), Inches(6), Inches(0.8)).table
    for c, text in enumerate(["Item", "Ação", "Situação"]):
        blockers.cell(0, c).text = text
    for c in range(3):
        _fill(blockers.cell(1, c), "FFFFFF")

    chart_data = CategoryChartData()
    chart_data.categories = ["Semana 1", "Semana 2"]
    chart_data.add_series("Entregues", (0, 0))
    board_slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.0), Inches(1.4), Inches(5.8), Inches(2.8), chart_data
    ).name = "Gráfico vazão"
    area = board_slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(7.0), Inches(4.4), Inches(5.8), Inches(2.6)
    )
    area.name = "Área gráfico riscos"
    area.fill.solid()
    area.fill.fore_color.rgb = RGBColor(0xEE, 0xEE, 0xEE)
    area.line.fill.background()

    # 3 — risk tracking table: header, group-header prototype row, item prototype row
    risks = prs.slides.add_slide(blank)
    _caption(risks, Inches(0.5), Inches(0.4), "Riscos")
    legend = risks.shapes.add_table(1, 4, Inches(7.3), Inches(0.4), Inches(5.5), Inches(0.35)).table
    for c, (text, color) in enumerate(
        [("Situação:", "FFFFFF"), ("Concluído", "00B050"), ("Em risco", "FFC000"), ("Atrasado", "FF0000")]
    ):
        legend.cell(0, c).text = text
        _fill(legend.cell(0, c), color)
    tracking = risks.shapes.add_table(3, 7, Inches(0.5), Inches(1.0), Inches(12.3), Inches(1.2)).table
    for c, text in enumerate(["Sprint", "Chave", "Descrição", "Responsável", "Ação", "Prazo", "Situação"]):
        tracking.cell(0, c).text = text
    tracking.cell(1, 2).text = "Épico"
    tracking.cell(2, 6).text = "Concluído"
    _fill(tracking.cell(2, 6), "00B050")
    for r in (1, 2):
        tracking.rows[r].height = Emu(Inches(0.35))

    # 4 — closing
    closing = prs.slides.add_slide(blank)
    box = closing.shapes.add_textbox(Inches(0.8), Inches(3), Inches(11), Inches(1))
    box.text_frame.paragraphs[0].text = "Gerado pela Cadência em {{relatorio.data}}"

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out
