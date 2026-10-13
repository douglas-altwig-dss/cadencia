"""Turn the hand-filled dss+360 weekly report into a Cadência deck template.

The report pipeline never draws slides: it opens a template and refills the structures it finds
inside it. So this script does not redraw the deck — it takes the delivered report and gives it the
structures ``report/generate.py`` needs, keeping every font, fill and border the author chose.

What changes, and why:

* slides 4 and 5 (two more projects) are dropped, and the surviving project slide plus its
  Issues/Risks slide become the repeat block that ``_structure`` clones once per board;
* the hand-written prose that no field can fill ("Use Case Purpose", "Business Challenge") is
  replaced by the sprint goal and the ``resumo.*`` numbers, as inline ``{{markers}}``;
* every ``grouped_list`` cell gets the three prototype paragraphs ``fill_grouped_list`` rebuilds a
  list from: group header, item, blank separator;
* the risks table gets a group prototype row and an item prototype row, which is what
  ``_prototypes`` deep-copies for ``grouped_rows``;
* two chart areas are added in the space freed at the right of the risks table;
* real customer data (project names, report week, baked status colours, ``[{}]`` tokens) becomes
  neutral placeholder content.

Run it from ``backend/``; it reads the original at the repo root and never writes over it.
"""

from __future__ import annotations

import copy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from cadencia.report.core import pptx_utils as px

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "dss+360 - AI Applications Weekly Report Template.pptx"
TARGET = ROOT / "dss+360 - AI Applications Weekly Report - Cadencia Template.pptx"

# Slides carrying the two extra projects; the first project slide becomes the repeat prototype.
DROP_SLIDE_IDS = (2147197275, 2147197277)
PROJECT_SLIDE_ID = 2147197271
RISKS_SLIDE_ID = 2147197269

NEUTRAL = "BFBFBF"  # status indicators start grey: the generator paints the real colour
GROUP_BAND = "EDE4FC"  # light tint of the template's header purple (640EE6)
PLACEHOLDER = "Preencha antes de enviar."

# Sprint / #Jira / Description / Responsible / Action / Limit date / Status, narrowed to 12.9"
# so the charts get a readable column. Sprint is wide enough for a name like "Plataforma 3.5".
RISK_COLUMNS_IN = (1.40, 1.00, 3.55, 1.45, 3.05, 1.10, 1.35)
CHART_LEFT_IN = 13.44
CHART_WIDTH_IN = 5.45


def shape_by_id(slide, shape_id: int):
    return next(s for s, _ in px.iter_shapes(slide.shapes) if s.shape_id == shape_id)


def slide_by_id(prs, slide_id: int):
    return next(s for s in prs.slides if s.slide_id == slide_id)


def set_paragraph(slide, shape_id: int, index: int, text: str) -> None:
    frame = shape_by_id(slide, shape_id).text_frame
    px.set_paragraph_text(frame.paragraphs[index], text)


def fill(target, hex_color: str) -> None:
    px.set_solid_fill(target, hex_color)


def strip_bullet(paragraph) -> None:
    """Drop whatever bullet a copied paragraph carries and mark it explicitly bullet-less."""
    ppr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buChar", "a:buAutoNum", "a:buNone", "a:buFont", "a:buClrTx", "a:buSzTx"):
        for el in ppr.findall(qn(tag)):
            ppr.remove(el)
    ppr.append(etree.SubElement(ppr, qn("a:buNone")))


def restyle(paragraph, *, size_pt: float | None = None, bold: bool | None = None) -> None:
    for run in paragraph.runs:
        if size_pt is not None:
            run.font.size = Pt(size_pt)
        if bold is not None:
            run.font.bold = bold
    end = paragraph._p.find(qn("a:endParaRPr"))
    if end is not None and size_pt is not None:
        end.set("sz", str(int(size_pt * 100)))


def grouped_list_prototypes(prs) -> tuple[object, object]:
    """The template's own two bullet levels, taken from the Control table's comments cell.

    Level 0 (a "•" bullet) is how the author writes a heading inside a cell and level 1 (an "o"
    sub-bullet) is how they write the items under it — exactly the roles ``fill_grouped_list``
    needs for a group header and its items.
    """
    cell = shape_by_id(slide_by_id(prs, PROJECT_SLIDE_ID), 6).table.cell(2, 2)
    paragraphs = cell.text_frame.paragraphs
    return copy.deepcopy(paragraphs[0]._p), copy.deepcopy(paragraphs[1]._p)


def build_grouped_list(cell, protos, *, header: str, item: str, size_pt: float) -> None:
    """Give a cell the three paragraphs ``fill_grouped_list`` clones: header, item, separator.

    It reads paragraph 0 as the group header prototype, the first non-empty paragraph after it as
    the item prototype, and the first empty paragraph as the separator between groups.
    """
    from pptx.text.text import _Paragraph

    header_proto, item_proto = protos
    body = cell.text_frame._txBody
    for paragraph in list(cell.text_frame.paragraphs):
        body.remove(paragraph._p)

    def add(proto, text: str, *, bold: bool = False, size: float = size_pt):
        el = copy.deepcopy(proto)
        body.append(el)
        paragraph = _Paragraph(el, cell.text_frame)
        px.set_paragraph_text(paragraph, text)
        restyle(paragraph, size_pt=size, bold=bold)
        return paragraph

    add(header_proto, header, bold=True)
    add(item_proto, item)
    separator = add(header_proto, "", size=6)
    strip_bullet(separator)


def scale_columns(table, widths_in: tuple[float, ...]) -> None:
    for gc, width in zip(table._tbl.tblGrid.gridCol_lst, widths_in, strict=True):
        gc.set("w", str(int(Inches(width))))


def build() -> Path:
    prs = Presentation(str(SOURCE))
    protos = grouped_list_prototypes(prs)

    for slide_id in DROP_SLIDE_IDS:
        px.delete_slide(prs, slide_by_id(prs, slide_id))

    _cover(prs.slides[0])
    _project(slide_by_id(prs, PROJECT_SLIDE_ID), protos)
    _risks(slide_by_id(prs, RISKS_SLIDE_ID))

    prs.save(str(TARGET))
    return TARGET


# ---------------------------------------------------------------- slide 1


def _cover(slide) -> None:
    """The cover is deck-scoped, so only ``relatorio.*`` resolves on it."""
    box = shape_by_id(slide, 7)
    frame = box.text_frame
    week = copy.deepcopy(frame.paragraphs[2]._p)
    frame._txBody.append(week)
    from pptx.text.text import _Paragraph

    paragraph = _Paragraph(week, frame)
    px.set_paragraph_text(paragraph, "{{relatorio.semana}}")
    restyle(paragraph, size_pt=20)
    box.height = Emu(int(Inches(2.9)))


# ---------------------------------------------------------------- slide 2


def _header_band(slide, name_shape: int, week_shape: int, indicator_shape: int, status_shape: int) -> None:
    # A neutral project name is what keeps suggest_mapping from pinning the slide to one board,
    # which is the precondition for as_repeat() turning it into a per-board slide.
    set_paragraph(slide, name_shape, 1, "Project name")
    set_paragraph(slide, week_shape, 1, "00/00/0000")
    fill(shape_by_id(slide, indicator_shape), NEUTRAL)
    # "GENERALSTATUS" had no space to wrap on and was set too large for its 1.27" box, so it broke
    # mid-word. The words matter: suggest.STATUS_CAPTION is what finds the indicator beside them.
    status = shape_by_id(slide, status_shape).text_frame.paragraphs[0]
    px.set_paragraph_text(status, "GENERAL STATUS")
    restyle(status, size_pt=11)


def _project(slide, protos) -> None:
    _header_band(slide, name_shape=11, week_shape=34, indicator_shape=46, status_shape=44)

    # Use Case Purpose -> Sprint Goal, Business Challenge -> Sprint Summary: both boxes held prose
    # no collected field can produce, and both fields below already exist in the catalogue.
    set_paragraph(slide, 23, 0, "Sprint Goal")
    px.set_frame_text(shape_by_id(slide, 2).table.cell(0, 0).text_frame, "{{sprint.objetivo}}")

    set_paragraph(slide, 29, 0, "Sprint Summary")
    px.set_frame_text(
        shape_by_id(slide, 24).table.cell(0, 0).text_frame,
        "Delivered {{resumo.entregues}}   ·   Committed {{resumo.comprometidos}}   ·   "
        "Adherence {{resumo.aderencia}}\n"
        "Added {{resumo.escopo_adicionado}}   ·   Carried over {{resumo.arrastados}}",
    )

    # Both boxes hold one or two lines, where the prose they replaced ran to a paragraph. Pulling
    # the column up hands the reclaimed inch to the Control table, whose generated comments are
    # the only thing on this slide that grows with the data.
    for shape_id, top in ((2, 2.38), (29, 3.30), (24, 3.68), (49, 4.60), (6, 4.98)):
        shape_by_id(slide, shape_id).top = Emu(int(Inches(top)))
    for shape_id in (2, 24):
        box = shape_by_id(slide, shape_id)
        box.height = Emu(int(Inches(0.75)))
        # The row keeps its own height, and estimate_row_height treats that as the floor — leave
        # the old one behind and validate_deck reports the box as overflowing when it is not.
        box.table._tbl.tr_lst[0].set("h", str(int(Inches(0.75))))

    # "Completed Activities" and "Upcoming Activities" wrapped inside their 2.65" boxes and the
    # second line disappeared behind the table underneath. The captions are what suggest.CAPTIONS
    # matches, so they are widened rather than shortened.
    for shape_id in (17, 19):
        shape_by_id(slide, shape_id).width = Emu(int(Inches(4.50)))

    for shape_id, header, item in (
        (18, "Epic name:", "Delivered item"),
        (22, "Epic name:", "Upcoming item"),
    ):
        build_grouped_list(
            shape_by_id(slide, shape_id).table.cell(0, 0), protos, header=header, item=item, size_pt=10
        )

    _blockers(shape_by_id(slide, 21))
    _control(shape_by_id(slide, 6), protos)


def _blockers(shape) -> None:
    """table_rows takes its item prototype from the first row after the header, so keep one."""
    table = shape.table
    tbl = table._tbl
    for tr in list(tbl.tr_lst)[2:]:
        tbl.remove(tr)
    for col, text in enumerate(("Blocked item", "Action required", "Status")):
        px.set_frame_text(table.cell(1, col).text_frame, text)
    # The status cell stays white. When a board has no blockers the generator keeps one blank row
    # cloned from this prototype, and a prototype painted red would show as a phantom blocker.
    fill(table.cell(1, 2), "FFFFFF")
    tbl.tr_lst[1].set("h", str(int(Inches(0.30))))
    shape.height = Emu(sum(int(tr.get("h") or 0) for tr in tbl.tr_lst))


def _control(shape, protos) -> None:
    """Deadline and Risks are the two dimensions the product can evidence.

    Scope, Budget and Quality have no data source anywhere in the collection, so they ship neutral
    rather than green: a stale green light reads as a live signal.
    """
    table = shape.table
    tbl = table._tbl
    for row in range(1, 6):
        fill(table.cell(row, 1), NEUTRAL)
    for row, header, item in (
        (2, "Epic name:", "Data limite 00/00/00"),
        (5, "Epic name:", "Item at risk — evidence"),
    ):
        build_grouped_list(table.cell(row, 2), protos, header=header, item=item, size_pt=9)
    for row in (1, 3, 4):
        px.set_frame_text(table.cell(row, 2).text_frame, PLACEHOLDER)
        # One short line each, so give their inch back to Deadline and Risks, whose generated
        # comments push the table towards the bottom of the slide.
        tbl.tr_lst[row].set("h", str(int(Inches(0.35))))
    shape.height = Emu(sum(int(tr.get("h") or 0) for tr in tbl.tr_lst))


# ---------------------------------------------------------------- slide 3


def _risks(slide) -> None:
    _header_band(slide, name_shape=27, week_shape=35, indicator_shape=40, status_shape=38)

    shape = shape_by_id(slide, 4)
    table = shape.table
    tbl = table._tbl
    # grouped_rows clones row 1 for group headers and row 2 for items; the file ships only one.
    tbl.append(copy.deepcopy(tbl.tr_lst[1]))

    for col in range(7):
        px.set_frame_text(table.cell(1, col).text_frame, "")
        fill(table.cell(1, col), GROUP_BAND)
    # _title_column() takes the first non-empty cell of the group row: it must be Description.
    px.set_frame_text(table.cell(1, 2).text_frame, "Epic name")
    restyle(table.cell(1, 2).text_frame.paragraphs[0], bold=True)

    items = ("Sprint 1", "KEY-1", "Item summary", "Responsible", "Action", "00/00", "Status")
    for col, text in enumerate(items):
        px.set_frame_text(table.cell(2, col).text_frame, text)
    fill(table.cell(2, 6), "FFFFFF")  # neutral for the same reason as the blockers prototype

    scale_columns(table, RISK_COLUMNS_IN)
    shape.width = Emu(int(Inches(sum(RISK_COLUMNS_IN))))
    for tr in tbl.tr_lst[1:]:
        tr.set("h", str(int(Inches(0.30))))
    shape.height = Emu(sum(int(tr.get("h") or 0) for tr in tbl.tr_lst))

    # The logo ends at y≈3.7; the column freed by narrowing the table holds the two charts. They
    # sit beside the table, so its downward pagination cannot reach them.
    for name, top in (("Área gráfico vazão", 3.90), ("Área gráfico evolução", 7.10)):
        area = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(CHART_LEFT_IN),
            Inches(top),
            Inches(CHART_WIDTH_IN),
            Inches(3.00),
        )
        area.name = name
        area.fill.solid()
        area.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF2)
        area.line.fill.background()
        area.shadow.inherit = False


if __name__ == "__main__":
    print(build())
