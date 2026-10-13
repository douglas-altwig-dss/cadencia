"""Template manifest (RF-37): everything replaceable in a .pptx, with stable identifiers.

Target ids are ``s<slide_id>/sh<shape_id>`` plus an optional part: ``#p<n>`` (paragraph),
``#r<r>c<c>`` (table cell) or ``#r<r>c<c>p<n>`` (paragraph inside a cell). Slide ids (the
``sldId@id`` attribute) and shape ids (``cNvPr@id``) are persisted inside the file, so they
survive reordering and re-saving in PowerPoint; shapes inside groups are addressed the same way.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from cadencia.report.core.pptx_utils import iter_shapes
from cadencia.report.core.text import markers_in

MANIFEST_VERSION = 2
_TARGET = re.compile(r"^s(\d+)/sh(\d+)(?:#(?:p(\d+)|r(\d+)c(\d+)(?:p(\d+))?))?$")
_A = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


@dataclass(frozen=True)
class Target:
    slide_id: int
    shape_id: int
    paragraph: int | None = None
    row: int | None = None
    col: int | None = None

    @property
    def is_cell(self) -> bool:
        return self.row is not None


def parse_target(target: str) -> Target:
    match = _TARGET.match(target.strip())
    if not match:
        raise ValueError(f"identificador de área inválido: {target!r}")
    slide, shape, para, row, col, cell_para = match.groups()
    paragraph = para if para is not None else cell_para
    return Target(
        slide_id=int(slide),
        shape_id=int(shape),
        paragraph=int(paragraph) if paragraph is not None else None,
        row=int(row) if row is not None else None,
        col=int(col) if col is not None else None,
    )


def area_id(slide_id: int, shape_id: int, part: str = "") -> str:
    return f"s{slide_id}/sh{shape_id}" + (f"#{part}" if part else "")


def file_sha256(path: str | Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 16), b""):
            digest.update(chunk)
    return digest.hexdigest()


def fill_hex(obj: Any) -> str | None:
    """Solid RGB fill of a shape or table cell, if it has one."""
    try:
        fill = obj.fill
        if fill.type == MSO_FILL.SOLID and fill.fore_color.type == MSO_COLOR_TYPE.RGB:
            return str(fill.fore_color.rgb)
    except (AttributeError, TypeError, ValueError, NotImplementedError):
        return None
    return None


def _bullet(paragraph: Any) -> bool:
    ppr = paragraph._p.pPr
    if ppr is None:
        return False
    return ppr.find(qn("a:buChar")) is not None or ppr.find(qn("a:buAutoNum")) is not None


_ALIGN = {"l": "left", "ctr": "center", "r": "right", "just": "justify"}
_ANCHOR = {"t": "top", "ctr": "middle", "b": "bottom"}


def _run_color(run: Any) -> str | None:
    try:
        color = run.font.color
        if color.type == MSO_COLOR_TYPE.RGB:
            return str(color.rgb)
        if color.type == MSO_COLOR_TYPE.SCHEME:
            return f"scheme:{color.theme_color}"
    except (AttributeError, TypeError, ValueError, NotImplementedError):
        return None
    return None


def _typography(paragraph: Any) -> dict[str, Any]:
    """Enough of the paragraph's look for a preview to be believable rather than merely correct."""
    run = paragraph.runs[0] if paragraph.runs else None
    font = run.font if run is not None else None
    ppr = paragraph._p.pPr
    align = ppr.get("algn") if ppr is not None else None
    size = None
    for candidate in paragraph.runs:
        if candidate.font.size is not None:
            size = round(candidate.font.size.pt, 1)
            break
    return {
        "size_pt": size,
        "bold": bool(font.bold) if font is not None and font.bold is not None else None,
        "italic": bool(font.italic) if font is not None and font.italic is not None else None,
        "font": font.name if font is not None else None,
        "color": _run_color(run) if run is not None else None,
        "align": _ALIGN.get(align or ""),
    }


def _paragraphs(frame: Any) -> list[dict[str, Any]]:
    return [
        {
            "index": i,
            "text": p.text[:160],
            "level": p.level,
            "bullet": _bullet(p),
            "runs": len(p.runs),
            **_typography(p),
        }
        for i, p in enumerate(frame.paragraphs)
    ]


def _frame_box(shape: Any) -> dict[str, Any]:
    """Insets and vertical anchor, so a preview can put the text where PowerPoint will."""
    frame = shape.text_frame
    body = frame._txBody.bodyPr
    return {
        "anchor": _ANCHOR.get(body.get("anchor") or "t", "top"),
        "wrap": (body.get("wrap") or "square") != "none",
        "columns": int(body.get("numCol") or 1),
        "insets": {
            "l": int(frame.margin_left or 0),
            "r": int(frame.margin_right or 0),
            "t": int(frame.margin_top or 0),
            "b": int(frame.margin_bottom or 0),
        },
        "rotation": round(float(shape.rotation or 0), 2),
    }


def _kind(shape: Any) -> str:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if getattr(shape, "has_table", False) and shape.has_table:
        return "table"
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return "chart"
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        return "text"
    return "shape"


def _area(slide: Any, shape: Any, parent_id: int | None) -> dict[str, Any]:
    kind = _kind(shape)
    sid = slide.slide_id
    base = area_id(sid, shape.shape_id)
    area: dict[str, Any] = {
        "id": base,
        "shape_id": shape.shape_id,
        "name": shape.name,
        "kind": kind,
        "parent": area_id(sid, parent_id) if parent_id else None,
        "bbox": {
            "x": int(shape.left or 0),
            "y": int(shape.top or 0),
            "w": int(shape.width or 0),
            "h": int(shape.height or 0),
        },
        "targets": [],
        "markers": [],
    }
    if shape.is_placeholder:
        fmt = shape.placeholder_format
        area["placeholder"] = {"idx": fmt.idx, "type": str(fmt.type)}
    if kind in {"text", "shape"}:
        area["fill"] = fill_hex(shape)
    if kind in {"text", "shape"} and shape.has_text_frame:
        paragraphs = _paragraphs(shape.text_frame)
        area["paragraphs"] = paragraphs
        area["frame"] = _frame_box(shape)
        area["markers"] = markers_in(shape.text_frame.text)
        area["targets"] = [
            # "chart" is offered on any plain shape: the generator renders the chart as an image
            # over whatever area it is given, which is how a template reserves space for one.
            {
                "id": base,
                "label": "texto inteiro",
                "accepts": [
                    "text",
                    "number",
                    "status_text",
                    "bullet_list",
                    "inline_list",
                    "grouped_list",
                    "status_fill",
                    "chart",
                ],
            }
        ]
        area["targets"] += [
            {
                "id": area_id(sid, shape.shape_id, f"p{p['index']}"),
                "label": f"parágrafo {p['index'] + 1}: {p['text'][:40]}",
                "accepts": ["text", "number", "status_text", "inline_list"],
            }
            for p in paragraphs
        ]
    elif kind == "shape":
        area["targets"] = [{"id": base, "label": "forma", "accepts": ["status_fill", "chart"]}]
    elif kind == "table":
        table = shape.table
        cells = [[cell.text[:120] for cell in row.cells] for row in table.rows]
        area["table"] = {
            "rows": len(table.rows),
            "cols": len(table.columns),
            "header": cells[0] if cells else [],
            "cells": cells,
            "fills": [[fill_hex(cell) for cell in row.cells] for row in table.rows],
            "col_widths": [int(c.width or 0) for c in table.columns],
            "row_heights": [int(r.height or 0) for r in table.rows],
            "first_row_header": bool(table.first_row),
            "typography": [
                [_typography(cell.text_frame.paragraphs[0]) for cell in row.cells] for row in table.rows
            ],
        }
        area["markers"] = [m for row in cells for text in row for m in markers_in(text)]
        area["targets"] = [{"id": base, "label": "tabela inteira", "accepts": ["table_rows", "grouped_rows"]}]
        for r, row in enumerate(table.rows):
            for c, cell in enumerate(row.cells):
                area["targets"].append(
                    {
                        "id": area_id(sid, shape.shape_id, f"r{r}c{c}"),
                        "label": f"célula {r + 1},{c + 1}: {cell.text[:40]}",
                        "accepts": [
                            "text",
                            "number",
                            "status_text",
                            "bullet_list",
                            "inline_list",
                            "grouped_list",
                            "status_fill",
                        ],
                    }
                )
    elif kind == "chart":
        chart = shape.chart
        plots = list(chart.plots)
        area["chart"] = {
            "chart_type": str(chart.chart_type),
            "series": [s.name for p in plots for s in p.series],
            "categories": [str(c) for c in plots[0].categories] if plots else [],
        }
        area["targets"] = [{"id": base, "label": "gráfico nativo", "accepts": ["chart"]}]
    elif kind == "picture":
        area["targets"] = [{"id": base, "label": "área da imagem", "accepts": ["chart"]}]
    return area


def _theme(prs: Any) -> dict[str, Any]:
    try:
        part = prs.slide_masters[0].part.part_related_by(RT.THEME)
    except (KeyError, IndexError):
        return {"colors": {}, "fonts": {}}
    root = etree.fromstring(part.blob)
    colors: dict[str, str] = {}
    for el in root.iterfind(".//a:clrScheme/*", _A):
        if len(el):
            child = el[0]
            colors[etree.QName(el).localname] = child.get("val") or child.get("lastClr") or ""
    fonts = {}
    for key in ("majorFont", "minorFont"):
        latin = root.find(f".//a:{key}/a:latin", _A)
        if latin is not None:
            fonts["major" if key == "majorFont" else "minor"] = latin.get("typeface")
    return {"colors": colors, "fonts": fonts}


def extract_manifest(path: str | Path) -> dict[str, Any]:
    prs = Presentation(str(path))
    masters = []
    for mi, master in enumerate(prs.slide_masters):
        masters.append(
            {
                "index": mi,
                "layouts": [
                    {
                        "index": li,
                        "name": layout.name,
                        "placeholders": [
                            {
                                "idx": ph.placeholder_format.idx,
                                "type": str(ph.placeholder_format.type),
                                "name": ph.name,
                            }
                            for ph in layout.placeholders
                        ],
                    }
                    for li, layout in enumerate(master.slide_layouts)
                ],
            }
        )
    slides = []
    for index, slide in enumerate(prs.slides, start=1):
        areas = [_area(slide, shape, parent) for shape, parent in iter_shapes(slide.shapes)]
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        headline = next(
            (
                a["paragraphs"][-1]["text"]
                for a in areas
                if a.get("paragraphs") and len(a["paragraphs"]) > 1 and a["paragraphs"][-1]["text"]
            ),
            "",
        )
        slides.append(
            {
                "index": index,
                "slide_id": slide.slide_id,
                "layout": slide.slide_layout.name,
                "headline": headline,
                "notes": notes[:300],
                "areas": areas,
            }
        )
    return {
        "version": MANIFEST_VERSION,
        "template": {
            "filename": Path(path).name,
            "sha256": file_sha256(path),
            "slide_width": int(prs.slide_width),
            "slide_height": int(prs.slide_height),
            "slides": len(prs.slides),
        },
        "theme": _theme(prs),
        "masters": masters,
        "slides": slides,
    }


def target_ids(manifest: dict[str, Any]) -> set[str]:
    return {t["id"] for slide in manifest.get("slides", []) for a in slide["areas"] for t in a["targets"]}


def diff_manifests(old: dict[str, Any], new: dict[str, Any]) -> dict[str, list[str]]:
    before, after = target_ids(old), target_ids(new)
    return {"removed": sorted(before - after), "added": sorted(after - before)}
