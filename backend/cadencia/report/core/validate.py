"""Deck validation before delivery (RF-43).

Checks, in order of how badly they break the delivery: the OOXML package itself (RF-50 — a file
PowerPoint offers to repair is worthless whatever it contains), then areas the mapping points to
that no longer exist in the template, mapped fields that came out empty, markers left unreplaced,
text that does not fit its area, tables running off the slide, and a template that changed since
the mapping was saved. When LibreOffice is available the deck is rendered headlessly to PDF/PNG for
visual inspection.
"""

from __future__ import annotations

import json
import shutil
import subprocess
from pathlib import Path
from typing import Any

from pptx import Presentation

from cadencia.i18n.message import Msg
from cadencia.report.core import pptx_utils as px
from cadencia.report.core.generate import GenerationResult
from cadencia.report.core.manifest import target_ids
from cadencia.report.core.mapping import DeckMapping
from cadencia.report.core.package_check import check_package
from cadencia.report.core.text import leftover_markers

#: A deck with any of these cannot be delivered without ``?force=true``. The two variable kinds
#: are here because that is the whole of the rule: a name the registry does not know, or one
#: nobody filled in, used to render as an empty string and ship looking finished.
ERROR_KINDS = {
    "area_ausente",
    "campo_desconhecido",
    "marcador_desconhecido",
    "variavel_nao_registrada",
    "variavel_sem_valor",
}
TEXT_TOLERANCE = 1.3


def _overflow(path: Path) -> list[dict[str, Any]]:
    prs = Presentation(str(path))
    height = int(prs.slide_height)
    found = []
    for number, slide in enumerate(prs.slides, start=1):
        for shape, _parent in px.iter_shapes(slide.shapes):
            if getattr(shape, "has_table", False) and shape.has_table:
                table = shape.table
                needed = sum(
                    px.estimate_row_height(table, tr, int(tr.get("h") or 0)) for tr in table._tbl.tr_lst
                )
                if int(shape.top or 0) + needed > height:
                    found.append(
                        {
                            "slide": number,
                            "area": shape.name,
                            "message": Msg(
                                "deck_problem.table_overflow", slide=number, area=shape.name
                            ).to_json(),
                        }
                    )
                elif (
                    len(table.rows) == 1
                    and len(table.columns) == 1
                    and needed > int(shape.height or 0) * TEXT_TOLERANCE
                ):
                    found.append(
                        {
                            "slide": number,
                            "area": shape.name,
                            "message": Msg(
                                "deck_problem.text_exceeds_area", slide=number, area=shape.name
                            ).to_json(),
                        }
                    )
            elif (
                getattr(shape, "has_text_frame", False)
                and shape.has_text_frame
                and shape.text_frame.text.strip()
            ):
                needed_pt = px.estimate_text_height_pt(shape.text_frame, int(shape.width or 0))
                have_pt = int(shape.height or 0) / px.EMU_PER_PT
                if have_pt and needed_pt > have_pt * TEXT_TOLERANCE and needed_pt - have_pt > 12:
                    found.append(
                        {
                            "slide": number,
                            "area": shape.name,
                            "message": Msg(
                                "deck_problem.text_does_not_fit", slide=number, area=shape.name
                            ).to_json(),
                        }
                    )
    return found


def render_previews(pptx_path: Path, out_dir: Path, *, dpi: int = 40) -> dict[str, Any]:
    """Render every slide to PNG. ``dpi`` 40 is enough to eyeball a deck; the mapping screen asks
    for more, because it draws clickable areas on top of the image.

    A ``reason`` is a ``{"$msg": …}`` record: the caller renders it for whoever reads it."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return {"rendered": False, "reason": Msg("deck_problem.libreoffice_missing").to_json()}
    out_dir.mkdir(parents=True, exist_ok=True)
    try:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(pptx_path)],
            check=True,
            capture_output=True,
            timeout=240,
        )
    except (subprocess.SubprocessError, OSError) as exc:
        return {
            "rendered": False,
            "reason": Msg("deck_problem.libreoffice_failed", error=str(exc)).to_json(),
        }
    pdf = out_dir / f"{pptx_path.stem}.pdf"
    images: list[str] = []
    pdftoppm = shutil.which("pdftoppm")
    if pdftoppm and pdf.exists():
        subprocess.run(
            [pdftoppm, "-png", "-r", str(dpi), str(pdf), str(out_dir / "slide")], check=False, timeout=240
        )
        images = sorted(str(p) for p in out_dir.glob("slide*.png"))
    return {"rendered": pdf.exists(), "pdf": str(pdf) if pdf.exists() else None, "images": images}


def validate_deck(
    result: GenerationResult,
    mapping: DeckMapping,
    *,
    mapped_sha256: str | None,
    current_manifest: dict[str, Any],
    render_dir: Path | None = None,
) -> dict[str, Any]:
    # Every ``message`` (and every ``info`` entry) is a ``{"$msg": …}`` record, not text: the
    # report service stores it with a rendered copy beside it, and readers render it in their own
    # language (docs/i18n-backend.md, 2.7).
    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    info: list[dict[str, Any]] = []

    current_sha = current_manifest.get("template", {}).get("sha256")
    if mapped_sha256 and current_sha and mapped_sha256 != current_sha:
        warnings.append(
            {
                "kind": "modelo_alterado",
                "message": Msg("deck_problem.template_changed").to_json(),
            }
        )
    known = target_ids(current_manifest)
    for target in mapping.targets():
        if target not in known:
            errors.append(
                {
                    "kind": "area_ausente",
                    "target": target,
                    "message": Msg("deck_problem.area_missing_in_template", target=target).to_json(),
                }
            )

    for problem in result.problems:
        entry = problem.as_dict()
        (errors if problem.kind in ERROR_KINDS else warnings).append(entry)

    # RF-50: the package itself, before anything about its contents. A deck that PowerPoint would
    # refuse to open is not a deck, however good its numbers are, so these are always errors.
    package = check_package(result.path)
    errors += package.errors
    warnings += package.warnings

    prs = Presentation(str(result.path))
    for number, slide in enumerate(prs.slides, start=1):
        for name in leftover_markers(slide):
            errors.append(
                {
                    "kind": "marcador_restante",
                    "slide": number,
                    "message": Msg(
                        "deck_problem.leftover_marker", slide=number, marker=f"{{{{{name}}}}}"
                    ).to_json(),
                }
            )
    for item in _overflow(result.path):
        warnings.append({"kind": "estouro", **item})

    preview: dict[str, Any] = {"rendered": False}
    if render_dir is not None:
        preview = render_previews(result.path, render_dir)
        if not preview.get("rendered") and preview.get("reason"):
            info.append(preview["reason"])

    # De-duplicate identical messages (e.g. the same empty field on every cloned slide), compared
    # as the message's key and parameters rather than as text in any one language.
    def unique(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
        seen, out = set(), []
        for it in items:
            key = (it.get("kind"), json.dumps(it.get("message"), sort_keys=True, ensure_ascii=False))
            if key not in seen:
                seen.add(key)
                out.append(it)
        return out

    return {
        "ok": not errors,
        "errors": unique(errors),
        "warnings": unique(warnings),
        "info": info,
        "preview": preview,
        "slides": result.slide_count,
        "package": {"ok": package.ok, "parts": package.parts, "slides": package.slides},
    }
