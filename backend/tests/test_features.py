"""RF-71 … RF-100: forecasts, changes, digest, delivery, health, search, views, versions, comments,
find & replace, locks, tables, conditional formatting, templated notes, schedules, cancellation."""

import asyncio
import json
from datetime import date, timedelta

import httpx
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from cadencia import notify
from cadencia.api.app import create_app
from cadencia.api.deps import clock
from cadencia.collector import scheduler
from cadencia.collector.service import COLLECT_LEASE, LEASE_TTL, Collector
from cadencia.domain.forecast import business_dates, monte_carlo
from cadencia.jira.client import JiraClient
from cadencia.report.core import layout as lay
from cadencia.report.core import pptx_utils as px
from cadencia.report.core.context import Row
from cadencia.report.core.edits import (
    CellStyle,
    DeckEdits,
    Fill,
    Frame,
    ShapeEdit,
    SlideEdits,
    TableEdit,
    apply_shape_edits,
    shape_key,
)
from cadencia.report.core.generate import _as_text, _table_rows, generate_deck
from cadencia.report.core.manifest import extract_manifest, parse_target
from cadencia.report.core.mapping import Binding, DeckMapping, FormatCondition, FormatRule, Schedule
from cadencia.report.core.package_check import check_package
from cadencia.store import repo
from cadencia.store.db import session_scope
from tests.support import NOW, PT_BR, collect_demo
from tests.test_report import board, deck, mapping_for


@pytest.fixture
def client(demo, sessions, settings):
    asyncio.run(collect_demo(demo, sessions, settings))
    app = create_app(settings, sessions)
    app.dependency_overrides[clock] = lambda: NOW
    with TestClient(app, headers=PT_BR) as test_client:
        yield test_client


def ok(response):
    assert response.status_code < 300, response.text
    return response.json()


# ------------------------------------------------------------------ RF-72 forecast


def test_the_forecast_is_reproducible_and_honest_about_its_edges():
    start = date(2026, 9, 10)
    first = monte_carlo([0, 1, 2, 0, 3, 1], 6, start, horizon_days=5)
    assert first == monte_carlo([0, 1, 2, 0, 3, 1], 6, start, horizon_days=5)  # seeded from the inputs
    assert first.p50_days <= first.p85_days <= first.p95_days
    assert 0 <= first.probability <= 1 and first.low_confidence
    assert first.p50_date == business_dates(start, first.p50_days)[-1]
    assert monte_carlo([1, 1], 0, start).p50_days == 0  # nothing left is done today
    stuck = monte_carlo([0, 0, 0], 3, start)
    assert stuck.p50_days is None and "Nenhum item" in str(stuck.reason)
    certain = monte_carlo([2] * 20, 4, start, horizon_days=2)
    assert (certain.p95_days, certain.probability, certain.low_confidence) == (2, 1.0, False)


def test_sprint_and_epic_forecasts_are_served(client):
    sprint = ok(client.get("/api/sprints/12/forecast"))
    assert sprint["scope"] == "sprint" and sprint["forecast"]["remaining"] == len(sprint["remaining_keys"])
    assert sprint["forecast"]["horizon_days"] is not None
    epic = ok(client.get("/api/epics/CAD-1/forecast", params={"board_id": 1}))
    assert epic["epic"]["key"] == "CAD-1" and "caveat" in epic
    assert client.get("/api/epics/CAD-999/forecast", params={"board_id": 1}).status_code == 404


def test_a_kanban_board_is_forecast_against_its_own_open_queue(client):
    """RF-113: a board with no sprint still gets a Monte Carlo answer — against what is open on
    it right now, not against a horizon it has no sprint end date to measure."""
    board = ok(client.get("/api/boards/3/forecast"))
    assert board["scope"] == "board"
    assert board["forecast"]["remaining"] == len(board["remaining_keys"])
    assert board["forecast"]["horizon_days"] is None  # continuous flow has no end date to aim at


# ------------------------------------------------------------------ RF-74 / RF-73 changes and digest


def test_what_changed_since_the_last_checkpoint(client):
    changes = ok(client.get("/api/boards/1/changes"))
    assert changes["since"]["at"] == changes["checkpoints"][-1]["at"] and changes["until"]["id"] is None
    assert "CAD-14" in {e["key"] for e in changes["new_risks"]}
    risk = next(e for e in changes["new_risks"] if e["key"] == "CAD-14")
    assert risk["label"] == "Atrasado" and risk["summary"] and risk["url"].endswith("/browse/CAD-14")

    marks = changes["checkpoints"]
    between = ok(
        client.get("/api/boards/1/changes", params={"since": marks[0]["id"], "until": marks[-1]["id"]})
    )
    assert between["until"]["id"] == marks[-1]["id"] and between["entered"]
    backwards = client.get(
        "/api/boards/1/changes", params={"since": marks[-1]["id"], "until": marks[0]["id"]}
    )
    assert backwards.status_code == 400
    assert client.get("/api/boards/1/changes", params={"since": 99999}).status_code == 404


def test_the_digest_lists_new_risks_per_space(client):
    digest = ok(client.get("/api/digest"))
    produto = next(s for s in digest["spaces"] if s["space"]["name"] == "Produto Digital")
    plataforma = next(b for b in produto["boards"] if b["board"]["name"] == "Plataforma")
    assert [e["key"] for e in plataforma["new_risks"]] == ["CAD-14"]
    assert digest["total_new_risks"] >= 1


# ------------------------------------------------------------------ RF-71 delivery


def test_channels_are_masked_kept_and_validated(client):
    saved = ok(
        client.put(
            "/api/notifications/channels",
            json={
                "channels": [
                    {
                        "name": "Time",
                        "kind": "slack",
                        "target": "https://hooks.slack.com/services/T0/B0/segredo1234",
                    }
                ]
            },
        )
    )
    channel = saved["channels"][0]
    assert "segredo" not in json.dumps(saved) and channel["target"].endswith("…1234")
    # Sent back masked, as the screen would, the stored URL is kept rather than overwritten.
    again = ok(client.put("/api/notifications/channels", json={"channels": [channel]}))
    assert again["channels"][0]["id"] == channel["id"]
    refused = client.put(
        "/api/notifications/channels", json={"channels": [{"kind": "email", "target": "sem-arroba"}]}
    )
    assert refused.status_code == 400
    assert (
        client.put(
            "/api/notifications/channels", json={"channels": [{"kind": "slack", "target": "ftp://x"}]}
        ).status_code
        == 400
    )


def test_delivery_posts_each_subscribed_channel_and_survives_a_failing_one(sessions, settings):
    posted = []

    def handler(request: httpx.Request) -> httpx.Response:
        posted.append((str(request.url), json.loads(request.content)))
        return httpx.Response(500 if "quebrado" in str(request.url) else 200)

    with session_scope(sessions) as session:
        notify.save_channels(
            session,
            [
                {"kind": "slack", "target": "https://hooks.example/slack", "events": ["deck"]},
                {"kind": "teams", "target": "https://hooks.example/quebrado", "events": ["deck"]},
                {"kind": "webhook", "target": "https://hooks.example/digest-only", "events": ["digest"]},
            ],
        )
        message = notify.Message("deck", "Apresentação gerada", "Pronta.", link="https://cadencia/x")
        results = notify.deliver(session, settings, message, transport=httpx.MockTransport(handler))
    assert [r["ok"] for r in results] == [True, False]
    assert [url for url, _ in posted] == ["https://hooks.example/slack", "https://hooks.example/quebrado"]
    assert "Apresentação gerada" in posted[0][1]["text"] and posted[1][1]["@type"] == "MessageCard"
    with session_scope(sessions) as session:
        assert "HTTP 500" in notify.last_results(session)[results[1]["channel"]]["error"]


def test_the_digest_goes_out_once_per_space_with_news(demo, sessions, settings):
    asyncio.run(collect_demo(demo, sessions, settings))
    bodies = []
    transport = httpx.MockTransport(lambda r: bodies.append(json.loads(r.content)) or httpx.Response(200))
    with session_scope(sessions) as session:
        notify.save_channels(
            session, [{"kind": "webhook", "target": "https://hooks.example/d", "events": ["digest"]}]
        )
        result = notify.send_digest(session, settings, NOW, transport=transport)
    assert [s["space"]["name"] for s in result["sent"]] == ["Produto Digital"]
    assert len(bodies) == 1 and "CAD-14" in bodies[0]["text"]


# ------------------------------------------------------------------ RF-78, RF-80, RF-85


def test_the_health_check_names_real_problems(client):
    report = ok(client.get("/api/health-check", params={"board_id": "1"}))
    checks = {c["id"]: c for c in report["items"][0]["checks"]}
    assert report["items"][0]["board"]["name"] == "Plataforma"
    assert all(c["severity"] in {"erro", "aviso", "info"} for c in checks.values())
    if "sem_estimativa" in checks:
        assert all(k.startswith("CAD-") for k in checks["sem_estimativa"]["examples"])


def test_search_finds_boards_items_and_people(client):
    assert client.get("/api/search", params={"q": "x"}).json()["items"] == []  # too short to mean anything
    kinds = {i["type"] for i in ok(client.get("/api/search", params={"q": "plata"}))["items"]}
    assert "board" in kinds
    items = ok(client.get("/api/search", params={"q": "cad-14"}))["items"]
    assert items[0]["id"] == "CAD-14" and items[0]["path"] == "/issues/CAD-14"


def test_a_person_sees_their_week(client):
    issue = ok(client.get("/api/issues/CAD-14"))
    account = issue["issue"]["assignee_id"]
    week = ok(client.get(f"/api/people/{account}/week"))
    assert week["person"]["name"] == issue["issue"]["assignee"]
    assert "CAD-14" in {i["key"] for i in week["at_risk"]}
    assert week["counts"]["open"] == len(week["open"])


# ------------------------------------------------------------------ RF-75 saved views


def test_saved_views_round_trip(client):
    created = ok(
        client.post(
            "/api/views",
            json={"name": "Riscos do Ana", "path": "/flow", "query": "?person=ana", "pinned": True},
        )
    )
    assert created["url"] == "/flow?person=ana"
    ok(
        client.put(
            f"/api/views/{created['id']}",
            json={"name": "Riscos da Ana", "path": "/flow", "query": "person=ana"},
        )
    )
    assert [v["name"] for v in ok(client.get("/api/views"))["items"]] == ["Riscos da Ana"]
    assert client.post("/api/views", json={"name": "fora", "path": "https://evil.example"}).status_code == 400
    assert client.delete(f"/api/views/{created['id']}").status_code == 204
    assert client.delete(f"/api/views/{created['id']}").status_code == 404


# ------------------------------------------------------------------ RF-87 versions, RF-93 comments, RF-88 replace


@pytest.fixture
def template_id(client, tmp_path):
    from cadencia.report.core.sample_template import build_sample_template

    path = build_sample_template(tmp_path / "modelo.pptx")
    return ok(client.post("/api/reports/templates", files={"file": ("modelo.pptx", path.read_bytes())}))["id"]


def test_versions_are_kept_named_diffed_and_restored(client, template_id):
    base = f"/api/reports/templates/{template_id}"
    first = {"edits": {"version": 1, "slides": {"256:*:*": {"shapes": {"sh2": {"frame": {"x": 100}}}}}}}
    ok(client.put(f"{base}/edits", json=first))
    ok(client.put(f"{base}/edits", json={"edits": {"version": 1, "slides": {}}}))
    autos = ok(client.get(f"{base}/versions", params={"kind": "edits"}))["items"]
    assert len(autos) == 1 and autos[0]["auto"]  # autosave collapses into one rolling copy

    named = ok(
        client.post(
            f"{base}/versions",
            json={"kind": "edits", "label": "antes da reunião"},
            headers={"X-Remote-User": "ana"},
        )
    )
    assert named["author"] == "ana" and not named["auto"]
    ok(client.put(f"{base}/edits", json=first))
    detail = ok(client.get(f"/api/reports/versions/{named['id']}"))
    assert detail["diff"]["slides"][0]["slide"] == "256:*:*" and detail["diff"]["slides"][0]["shapes"] == [
        "sh2"
    ]

    ok(client.post(f"/api/reports/versions/{named['id']}/restore", json={}))
    assert ok(client.get(f"{base}/edits"))["edits"]["slides"] == {}
    labels = [v["label"] for v in ok(client.get(f"{base}/versions"))["items"]]
    assert any(label.startswith("antes de restaurar") for label in labels)

    mapping = ok(client.get(base))["suggestion"]
    ok(client.put(f"{base}/mapping", json=mapping))
    assert ok(client.get(f"{base}/versions", params={"kind": "mapping"}))["items"]


def test_comments_thread_resolve_and_delete(client, template_id):
    base = f"/api/reports/templates/{template_id}/comments"
    root = ok(
        client.post(
            base, json={"origin_slide_id": 256, "shape_key": "sh3", "body": "Número errado?", "author": "bia"}
        )
    )
    reply = ok(
        client.post(
            base, json={"parent_id": root["id"], "body": "Corrigido."}, headers={"X-Remote-User": "caio"}
        )
    )
    assert (reply["origin_slide_id"], reply["shape_key"], reply["author"]) == (256, "sh3", "caio")
    assert ok(client.get(base))["open"] == 1
    ok(client.patch(f"/api/reports/comments/{root['id']}", json={"resolved": True}))
    assert ok(client.get(base))["open"] == 0
    assert client.post(base, json={"origin_slide_id": 256, "body": "   "}).status_code == 400
    assert client.delete(f"/api/reports/comments/{root['id']}").status_code == 204
    assert ok(client.get(base))["items"] == []  # the replies went with it


def test_find_and_replace_previews_before_it_writes(client):
    created = ok(client.post("/api/reports/variables", json={"name": "produto", "default": "Cadencia Beta"}))
    ok(
        client.put(
            f"/api/reports/variables/{created['id']}/values",
            json={"scope": "global", "value": "Usa a cadencia"},
        )
    )
    preview = ok(
        client.post("/api/reports/variables/replace", json={"find": "cadencia", "replace": "Cadência"})
    )
    assert preview["count"] == 2 and not preview["applied"]
    assert ok(client.get("/api/reports/variables"))["items"][0]["default"] == "Cadencia Beta"
    applied = ok(
        client.post(
            "/api/reports/variables/replace", json={"find": "cadencia", "replace": "Cadência", "apply": True}
        )
    )
    assert applied["applied"]
    variable = ok(client.get("/api/reports/variables"))["items"][0]
    assert variable["default"] == "Cadência Beta" and variable["values"][0]["value"] == "Usa a Cadência"
    assert (
        ok(client.post("/api/reports/variables/replace", json={"find": "beta", "match_case": True}))["count"]
        == 0
    )


# ------------------------------------------------------------------ RF-89 locks, RF-91 tables


def test_locks_and_hiding_never_reach_the_file():
    edits = DeckEdits(
        slides={"1:0:1": SlideEdits(locked=["sh2"], hidden=["sh3"], groups={"g": ["sh2", "sh3"]})}
    )
    assert edits.touched() == set()
    template = DeckEdits(slides={"1:*:*": SlideEdits(locked=["sh2"])})
    template.slides["1:5:1"] = SlideEdits(locked=["sh4"], hidden=["sh2"])
    merged = template.for_slide("1:5:1")
    assert merged.locked == ["sh2", "sh4"] and merged.hidden == ["sh2"]


def test_a_background_set_on_the_template_reaches_every_board_unless_overridden():
    template = DeckEdits(slides={"1:*:*": SlideEdits(background=Fill(type="solid", color="AABBCC"))})
    assert template.for_slide("1:5:1").background == Fill(type="solid", color="AABBCC")

    template.slides["1:5:1"] = SlideEdits(background=Fill(type="solid", color="112233"))
    assert template.for_slide("1:5:1").background == Fill(type="solid", color="112233")


def test_inserting_a_line_writes_a_connector_not_a_rectangle():
    """A "line" preset used to fall back to a filled rectangle (`PRESETS["line"]` pointed at
    `MSO_SHAPE.RECTANGLE`); it must be a connector, corner to corner of its box, the way Slides'
    own Line tool and `geometry.py`'s reading of `prst="line"` both expect."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    edits = SlideEdits(
        shapes={
            "new:linha": ShapeEdit(
                op="insert", kind="shape", preset="line", frame=Frame(x=100, y=200, cx=300, cy=400)
            )
        }
    )
    assert apply_shape_edits(slide, edits, {}) == []
    (shape,) = list(slide.shapes)
    assert shape._element.tag.endswith("cxnSp")
    assert shape._element.find(".//{*}prstGeom").get("prst") == "line"
    assert (int(shape.left), int(shape.top), int(shape.width), int(shape.height)) == (100, 200, 300, 400)


def test_slide_background_is_a_real_file_change_not_an_editor_only_one():
    """Slides' "Background" button — a colour on this slide's own `<p:bg>`, through python-pptx's
    own `slide.background.fill` (the exact `_set_fill` a shape already goes through)."""
    from cadencia.report.core.scene import _background_of

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    assert _background_of(slide) is None  # nothing set yet: inherits the layout/master

    edits = SlideEdits(background=Fill(type="solid", color="112233"))
    assert edits.replays() is True  # unlike `groups`/`locked`, this changes the file
    assert apply_shape_edits(slide, edits, {}) == []
    assert _background_of(slide) == {"type": "solid", "color": "112233"}

    # An explicit "none" resets it back to inheriting, distinct from never having set it at all.
    apply_shape_edits(slide, SlideEdits(background=Fill(type="none")), {})
    assert _background_of(slide) is None


def test_table_widths_and_cell_styles_are_written_as_valid_ooxml(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1))
    shape.table.cell(0, 0).text = "Cabeçalho"
    key = shape_key(shape)
    edits = SlideEdits(
        tables={
            key: TableEdit(
                col_widths=[Inches(3), Inches(2)],
                cells={
                    "0:0": CellStyle(fill="1F4E79", bold=True, color="FFFFFF", align="center"),
                    "9:9": CellStyle(bold=True),
                },
            )
        }
    )
    assert DeckEdits(slides={"x": edits}).touched() == {"x"}
    problems = apply_shape_edits(slide, edits, {})
    assert [str(p) for p in problems] == [f"A célula 10,10 saiu da tabela {key}."]
    table = shape.table
    assert int(table.columns[0].width) == Inches(3) and int(shape.width) == Inches(5)
    cell = table.cell(0, 0)
    assert str(cell.fill.fore_color.rgb) == "1F4E79"
    run = cell.text_frame.paragraphs[0].runs[0]
    assert run.font.bold and str(run.font.color.rgb) == "FFFFFF"
    out = tmp_path / "tabela.pptx"
    prs.save(str(out))
    report = check_package(out)
    assert report.ok, report.errors


# ------------------------------------------------------------------ RF-92 conditional formatting


@pytest.mark.parametrize(
    ("when", "value", "shown", "expected"),
    [
        ("maior", "80", "85%", True),
        ("maior", "80", "75%", False),
        ("menor_igual", "1.000,50", "R$ 1.000,50", True),
        ("igual", "Atrasado", "atrasado", True),
        ("contem", "bloq", "Bloqueado X", True),
        ("maior", "3", "muitos", False),
        ("vazio", "", "  ", True),
        ("nao_vazio", "", "x", True),
    ],
)
def test_rules_read_the_value_the_way_the_slide_shows_it(when, value, shown, expected):
    assert lay.rule_matches(FormatRule(when=when, value=value, color="FF0000"), shown) is expected


def test_a_rule_colours_the_generated_text(tmp_path):
    from cadencia.report.core.sample_template import build_sample_template

    template = build_sample_template(tmp_path / "modelo.pptx")
    manifest = extract_manifest(template)
    mapping = mapping_for(manifest)
    board_slide = str(manifest["slides"][1]["slide_id"])
    spec = mapping.slides[board_slide]
    binding = next(b for b in spec.bindings if b.mode in {"text", "status_text"})
    binding.rules = [FormatRule(when="nao_vazio", color="C00000")]
    result = generate_deck(template, mapping, deck(board(1, "Plataforma")), tmp_path / "out.pptx")
    shape_id = parse_target(binding.target).shape_id
    slides = list(Presentation(str(result.path)).slides)
    copies = [slides[p["index"] - 1] for p in result.produced if str(p["origin_slide_id"]) == board_slide]
    colors = {
        str(run.font.color.rgb)
        for slide in copies
        if (shape := px.find_shape(slide, shape_id)) is not None and shape.has_text_frame
        for paragraph in [shape.text_frame.paragraphs[parse_target(binding.target).paragraph]]
        for run in paragraph.runs
        if run.text.strip()
    }
    assert colors == {"C00000"}


def test_table_rules_fill_the_tested_cell_or_the_whole_row():
    prs = Presentation()
    table = (
        prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_table(2, 3, 0, 0, Inches(3), Inches(1)).table
    )
    rows = [
        Row({"chave": "K-1", "resumo": "Bloqueado X", "prazo": "10/09"}, "risco"),
        Row({"chave": "K-2", "resumo": "Ok", "prazo": ""}, "em_dia"),
    ]
    binding = Binding(
        target="s1/sh1",
        mode="table_rows",
        columns=["chave", "resumo", "prazo"],
        rules=[
            FormatRule(
                when="contem", value="bloq", color="FF0000", apply="preenchimento", column=1, whole_row=True
            ),
            FormatRule(when="vazio", color="999999", column=2),
        ],
    )
    out = _table_rows(binding, "rows", rows, DeckMapping(), table)
    assert out[0].fills == {0: "FF0000", 1: "FF0000", 2: "FF0000"} and out[0].colors == {}
    assert out[1].fills == {} and out[1].colors == {2: "999999"}


def test_compound_rules_require_every_also_condition():
    prs = Presentation()
    table = (
        prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_table(2, 3, 0, 0, Inches(3), Inches(1)).table
    )
    rows = [
        Row({"chave": "K-1", "resumo": "Bloqueado X", "prazo": "alta"}, "risco"),
        Row({"chave": "K-2", "resumo": "Bloqueado Y", "prazo": "baixa"}, "risco"),
    ]
    binding = Binding(
        target="s1/sh1",
        mode="table_rows",
        columns=["chave", "resumo", "prazo"],
        rules=[
            FormatRule(
                when="contem",
                value="bloq",
                color="FF0000",
                apply="preenchimento",
                column=1,
                whole_row=True,
                also=[FormatCondition(when="igual", value="alta", column=2)],
            )
        ],
    )
    out = _table_rows(binding, "rows", rows, DeckMapping(), table)
    assert out[0].fills == {0: "FF0000", 1: "FF0000", 2: "FF0000"}
    assert out[1].fills == {}  # "bloq" matches but priority is not "alta", so the AND fails


# ------------------------------------------------------------------ RF-94 templated notes


def test_marker_modifiers_write_lists_into_running_text():
    rows = [
        Row({"chave": "CAD-11", "resumo": "Um"}, "risco"),
        Row({"chave": "CAD-14", "resumo": "Dois"}, "risco"),
    ]
    mapping = DeckMapping()
    assert _as_text("rows", rows, mapping, "contagem") == "2"
    assert _as_text("rows", rows, mapping, "chaves") == "CAD-11, CAD-14"
    assert _as_text("rows", rows, mapping, "lista") == "CAD-11 Um\nCAD-14 Dois"
    assert _as_text("text", "Sprint 3", mapping, "contagem") == "Sprint 3"  # nothing to count


def test_speaker_notes_are_filled_per_board(tmp_path):
    from cadencia.report.core.sample_template import build_sample_template

    template = build_sample_template(tmp_path / "modelo.pptx")
    prs = Presentation(str(template))
    prs.slides[
        1
    ].notes_slide.notes_text_frame.text = (
        "{{board.nome}}: {{riscos.bloqueios|contagem}} bloqueio(s). {{riscos.bloqueios|nada}}"
    )
    prs.save(str(template))
    manifest = extract_manifest(template)
    result = generate_deck(
        template,
        mapping_for(manifest),
        deck(board(1, "Plataforma"), board(2, "Dados")),
        tmp_path / "out.pptx",
    )
    notes = [
        s.notes_slide.notes_text_frame.text
        for s in Presentation(str(result.path)).slides
        if s.has_notes_slide
    ]
    assert any(n.startswith("Plataforma: 1 bloqueio(s).") for n in notes)
    assert any(n.startswith("Dados: 1 bloqueio(s).") for n in notes)
    # An unknown modifier stays visible and is reported, like an unknown name.
    assert all("{{riscos.bloqueios|nada}}" in n for n in notes if "bloqueio" in n)
    assert any(p.kind == "marcador_desconhecido" for p in result.problems)


# ------------------------------------------------------------------ RF-77 schedules


def test_template_schedules_become_cron_jobs_and_follow_changes(sessions, settings):
    from apscheduler.schedulers.background import BackgroundScheduler

    sched = BackgroundScheduler(timezone="UTC")
    with session_scope(sessions) as session:
        template = repo.add_template(
            session, name="Semanal", filename="s.pptx", sha256="s", path="", manifest={}
        )
        mapping = DeckMapping(schedule=Schedule(enabled=True, weekdays=[4], hour=10))
        repo.save_mapping(session, template.id, "s", mapping.model_dump())
        notify.save_digest_schedule(session, notify.DigestSchedule(enabled=True, weekdays=[0], hour=8))
        tid = template.id
    scheduler.sync_template_jobs(sched, settings, sessions)
    ids = {job.id for job in sched.get_jobs()}
    assert f"template:{tid}:fri:10:0:America/Sao_Paulo" in ids and "digest:mon:8:America/Sao_Paulo" in ids
    with session_scope(sessions) as session:
        repo.save_mapping(
            session, tid, "s", DeckMapping(schedule=Schedule(enabled=True, weekdays=[3], hour=9)).model_dump()
        )
        notify.save_digest_schedule(session, notify.DigestSchedule(enabled=False))
    scheduler.sync_template_jobs(sched, settings, sessions)
    assert {job.id for job in sched.get_jobs()} == {f"template:{tid}:thu:9:0:America/Sao_Paulo"}


def test_a_schedule_refuses_impossible_times():
    with pytest.raises(ValueError):
        Schedule(weekdays=[7])
    with pytest.raises(ValueError):
        Schedule(timezone="Marte/Base")


# ------------------------------------------------------------------ RF-82 progress and cancel


async def test_a_cancelled_collection_stops_cleanly(demo, sessions, settings):
    from tests.support import _no_sleep

    async with JiraClient.from_settings(settings, transport=demo.transport, sleep=_no_sleep) as client:
        collector = Collector(client, sessions, settings, now=lambda: NOW)
        await collector.discover()
        with session_scope(sessions) as session:
            repo.set_selected_boards(session, (1, 2, 3))
            repo.request_cancel(session, "r-cancel")
        demo.calls.clear()
        result = await collector.collect(mode="full", run_id="r-cancel")
    assert result["status"] == "cancelled"
    assert {b["status"] for b in result["boards"]} == {"cancelled"}
    assert not any("/changelog" in path for _m, path, _p in demo.calls)
    with session_scope(sessions) as session:
        assert repo.lease_holder(session, COLLECT_LEASE) is None and not repo.cancel_requested(
            session, "r-cancel"
        )
        assert all(b.last_error is None for b in repo.selected_boards(session))  # a decision, not a failure


def test_status_reports_the_running_collection_board_by_board(client, sessions):
    with session_scope(sessions) as session:
        repo.acquire_lease(session, COLLECT_LEASE, "r-now", LEASE_TTL)
        repo.start_run(session, "r-now", "full")
        repo.start_board_run(session, "r-now", 1)
        repo.board_progress(session, "r-now", 1, "historico", 10, 40)
    current = ok(client.get("/api/status"))["current"]
    assert current["run_id"] == "r-now" and current["mode"] == "full"
    assert (
        current["boards"][0] | {"name": "Plataforma", "phase": "historico", "done": 10, "total": 40}
        == current["boards"][0]
    )
    assert current["boards"][0]["eta_seconds"] is not None  # from the demo's own finished runs
    assert ok(client.post("/api/collect/cancel"))["cancelling"] == "r-now"
    assert ok(client.get("/api/status"))["current"]["cancel_requested"] is True
    with session_scope(sessions) as session:
        repo.release_lease(session, COLLECT_LEASE, "r-now")
    assert client.post("/api/collect/cancel").status_code == 409


def test_template_deletion_takes_versions_and_comments(client, template_id, sessions):
    ok(
        client.post(
            f"/api/reports/templates/{template_id}/comments", json={"origin_slide_id": 256, "body": "oi"}
        )
    )
    ok(client.post(f"/api/reports/templates/{template_id}/versions", json={"kind": "edits", "label": "v1"}))
    ok(client.delete(f"/api/reports/templates/{template_id}"))
    with session_scope(sessions) as session:
        assert (
            repo.list_versions(session, template_id) == [] and repo.list_comments(session, template_id) == []
        )


def test_business_dates_skip_weekends_and_holidays():
    friday = date(2026, 9, 4)
    assert business_dates(friday, 2, holidays=[date(2026, 9, 7)]) == [date(2026, 9, 8), date(2026, 9, 9)]
    assert business_dates(friday, 1)[0] - friday == timedelta(days=3)


# ------------------------------------------------------------------ sprint timeline


def _flatten(nodes):
    for node in nodes:
        yield node
        yield from _flatten(node["children"])


def test_the_sprint_timeline_lists_every_item_as_a_tree_with_its_time(client):
    view = ok(client.get("/api/sprints/12/timeline"))
    rows = {n["key"]: n for n in _flatten(view["nodes"])}
    assert view["window"]["start"] < view["window"]["today"] <= view["window"]["end"]
    # Epics, stories and their subtasks, and the items with no epic, all in one tree.
    assert {"CAD-1", "CAD-2", "CAD-10", "CAD-19", "CAD-20", "CAD-13", "CAD-21"} <= set(rows)
    assert [c["key"] for c in rows["CAD-10"]["children"]] == ["CAD-19", "CAD-20"]
    assert view["counts"] == {"epicos": 2, "itens": 10, "subtarefas": 2, "total": 14}
    # An item that is not in the sprint stays out (CAD-16 moved to the next one).
    assert "CAD-16" not in rows


def test_an_epic_s_own_initiative_is_collected_and_named_instead_of_a_placeholder(client):
    """RF-111: CAD-50 sits on no board of its own — CAD-2's parent link is the only way to it —
    so this proves the collector's second hop past an epic's own parent, not just the first."""
    rows = {n["key"]: n for n in _flatten(ok(client.get("/api/sprints/12/timeline"))["nodes"])}
    initiative = rows["CAD-50"]
    assert initiative["label"] == "Confiabilidade da plataforma"
    assert initiative["missing"] is False and initiative["synthetic"] is False
    assert initiative["issue"]["key"] == "CAD-50"
    assert [c["key"] for c in initiative["children"]] == ["CAD-2"]


def test_each_row_says_when_it_was_worked_on_and_in_what_state(client):
    rows = {n["key"]: n for n in _flatten(ok(client.get("/api/sprints/12/timeline"))["nodes"])}
    done = rows["CAD-10"]
    assert done["state"] == "concluido" and done["start"] < done["end"]
    assert (
        rows["CAD-13"]["state"] == "bloqueado" and rows["CAD-13"]["blocked"] and rows["CAD-13"]["end"] is None
    )
    assert rows["CAD-18"]["state"] == "bloqueado"  # waiting on CAD-11 by a link
    never = rows["CAD-17"]
    assert never["state"] == "aberto" and never["start"] is None and never["end"] is None
    # Joined after the sprint began, and carried from an earlier sprint.
    assert rows["CAD-15"]["added"] and not rows["CAD-10"]["added"]
    assert rows["CAD-11"]["carried"] and not rows["CAD-12"]["carried"]


def test_containers_span_their_children_and_end_only_when_all_are_done(client):
    rows = {n["key"]: n for n in _flatten(ok(client.get("/api/sprints/12/timeline"))["nodes"])}
    epic = rows["CAD-1"]
    starts = [c["start"] for c in epic["children"] if c["start"]]
    assert epic["start"] == min(starts) and epic["end"] is None and epic["blocked"]
    # A group with an item still open does not end, even though its other item finished.
    assert rows["__sem_epico__"]["end"] is None
    assert rows["CAD-2"]["end"] is None  # CAD-14 is still in progress
