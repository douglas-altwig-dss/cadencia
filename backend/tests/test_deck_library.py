"""Editor and template support: accessibility (#15), blocks (#16), image library (#17), table
scales (#18), presence (#19), per-slide history (#20), diagnostics (#24) and usage (#30)."""

from __future__ import annotations

import asyncio
import io
from datetime import UTC, datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

from cadencia.api.app import create_app
from cadencia.api.deps import clock
from cadencia.report.core.a11y import audit, contrast
from cadencia.report.core.context import Row
from cadencia.report.core.layout import scale_marks
from cadencia.report.core.mapping import CellScale
from cadencia.report.core.sample_template import build_sample_template
from cadencia.store import schema as s
from cadencia.store.db import session_scope
from tests.support import NOW, collect_demo


def ok(response):
    assert response.status_code < 300, response.text
    return response.json()


@pytest.fixture
def client(demo, sessions, settings):
    asyncio.run(collect_demo(demo, sessions, settings))
    app = create_app(settings, sessions)
    app.dependency_overrides[clock] = lambda: NOW
    with TestClient(app) as test_client:
        yield test_client


@pytest.fixture
def template(client, tmp_path):
    path = build_sample_template(tmp_path / "modelo.pptx")
    created = ok(client.post("/api/reports/templates", files={"file": ("modelo.pptx", path.read_bytes())}))
    return ok(client.get(f"/api/reports/templates/{created['id']}"))


# ------------------------------------------------------------------ accessibility (#15)


def test_the_audit_finds_silent_pictures_low_contrast_and_reading_order():
    scene = {
        "theme": {"background": "FFFFFF"},
        "slides": [
            {
                "index": 1,
                "shapes": [
                    {
                        "key": "sh2",
                        "kind": "text",
                        "text": {"paragraphs": [{"runs": [{"text": "a", "color": "DDDDDD", "size_pt": 12}]}]},
                    },
                    {
                        "key": "sh3",
                        "kind": "text",
                        "placeholder": {"type": "TITLE (1)"},
                        "text": {"paragraphs": [{"runs": [{"text": "T"}]}]},
                    },
                    {"key": "sh4", "kind": "image", "alt": ""},
                    {"key": "sh5", "kind": "chart", "alt": "Vazão por semana"},
                    {
                        "key": "sh6",
                        "kind": "text",
                        "fill": {"type": "solid", "color": "000000"},
                        "text": {"paragraphs": [{"runs": [{"text": "b", "color": "FFFFFF"}]}]},
                    },
                ],
            }
        ],
    }
    kinds = [(f["kind"], f["shape"]) for f in audit(scene)]
    assert ("contraste_baixo", "sh2") in kinds and ("sem_texto_alternativo", "sh4") in kinds
    assert ("ordem_de_leitura", "sh3") in kinds
    assert all(shape not in {"sh5", "sh6"} for _, shape in kinds)
    assert contrast("000000", "FFFFFF") == 21.0


def test_generated_charts_carry_alt_text_and_edits_can_set_it(client, template, sessions):
    board_slide = template["manifest"]["slides"][1]
    chart = next(a for a in board_slide["areas"] if a["kind"] == "chart")
    ok(
        client.put(
            f"/api/reports/templates/{template['id']}/mapping",
            json={
                "slides": {
                    str(board_slide["slide_id"]): {
                        "scope": "board",
                        "board_id": 1,
                        "bindings": [{"target": chart["id"], "mode": "chart", "field": "grafico.vazao"}],
                    }
                }
            },
        )
    )
    run = ok(client.post("/api/reports/generate", json={"template_id": template["id"], "render": False}))
    with session_scope(sessions) as session:
        path = session.get(s.ReportRunRow, run["id"]).path
    frame = next(
        sh for sh in Presentation(path).slides[1].shapes if getattr(sh, "has_chart", False) and sh.has_chart
    )
    assert frame._element.xpath("./*[1]/p:cNvPr")[0].get("descr").startswith("Vazão por janela da sprint:")
    report = ok(client.get(f"/api/reports/runs/{run['id']}/accessibility"))
    assert "sem_texto_alternativo" not in {
        f["kind"] for f in report["findings"] if f["shape"] == f"sh{chart['shape_id']}"
    }


# ------------------------------------------------------------------ blocks (#16), images (#17)


def test_blocks_keep_inserted_shapes_relative_to_their_corner(client):
    shapes = {
        "k1": {
            "op": "insert",
            "kind": "textbox",
            "frame": {"x": 1_000_000, "y": 2_000_000, "cx": 500_000, "cy": 200_000},
        },
        "k2": {
            "op": "insert",
            "kind": "shape",
            "preset": "rect",
            "frame": {"x": 1_600_000, "y": 2_100_000, "cx": 300_000, "cy": 300_000},
        },
        "sh9": {"op": "update", "frame": {"x": 5}},  # a template shape does not travel
    }
    block = ok(
        client.post(
            "/api/reports/blocks",
            json={"nome": "Legenda", "formas": shapes},
            headers={"X-Remote-User": "ana"},
        )
    )
    assert set(block["shapes"]) == {"k1", "k2"} and block["author"] == "ana"
    assert block["shapes"]["k1"]["frame"]["x"] == 0 and block["shapes"]["k2"]["frame"]["y"] == 100_000
    assert (block["width"], block["height"]) == (900_000, 400_000)
    assert [b["name"] for b in ok(client.get("/api/reports/blocks"))["items"]] == ["Legenda"]
    assert (
        client.post(
            "/api/reports/blocks", json={"nome": "Vazio", "formas": {"sh9": shapes["sh9"]}}
        ).status_code
        == 400
    )
    assert client.delete(f"/api/reports/blocks/{block['id']}").status_code == 204


def test_blocks_favour_the_most_used_over_the_alphabet(client):
    def make(name: str) -> dict:
        shapes = {"k1": {"op": "insert", "kind": "textbox", "frame": {"x": 0, "y": 0, "cx": 1, "cy": 1}}}
        return ok(client.post("/api/reports/blocks", json={"nome": name, "formas": shapes}))

    a, b = make("Alfa"), make("Beta")
    assert [x["name"] for x in ok(client.get("/api/reports/blocks"))["items"]] == ["Alfa", "Beta"]
    used = ok(client.post(f"/api/reports/blocks/{b['id']}/use"))
    assert used["usage_count"] == 1 and used["last_used_at"]
    ok(client.post(f"/api/reports/blocks/{a['id']}/use"))
    ok(client.post(f"/api/reports/blocks/{b['id']}/use"))
    # Beta was used twice, Alfa once — usage outranks the alphabet, most-used first.
    assert [x["name"] for x in ok(client.get("/api/reports/blocks"))["items"]] == ["Beta", "Alfa"]
    assert client.post("/api/reports/blocks/99999/use").status_code == 404


def _png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (20, 10), "blue").save(buffer, "PNG")
    return buffer.getvalue()


def test_images_are_labelled_tagged_and_searchable(client):
    asset = ok(client.post("/api/reports/assets", files={"file": ("logo-cliente.png", _png(), "image/png")}))
    ok(
        client.patch(
            f"/api/reports/assets/{asset['id']}",
            json={"rotulo": "Logo Acme", "etiquetas": ["logo", "cliente"]},
        )
    )
    assert [a["label"] for a in ok(client.get("/api/reports/assets", params={"q": "acme"}))["items"]] == [
        "Logo Acme"
    ]
    assert ok(client.get("/api/reports/assets", params={"tag": "LOGO"}))["items"]
    assert ok(client.get("/api/reports/assets", params={"q": "nada"}))["items"] == []
    assert set(ok(client.get("/api/reports/assets"))["tags"]) == {"cliente", "logo"}


# ------------------------------------------------------------------ table scales (#18)


def test_scales_draw_bars_arrows_and_colours_from_the_same_numbers():
    rows = [["A", "1"], ["B", "5"], ["C", "3"], ["D", "n/d"]]
    marks = scale_marks(
        rows,
        [
            CellScale(column=1, kind="barra"),
            CellScale(column=1, kind="icone"),
            CellScale(column=1, kind="cor"),
        ],
    )
    assert marks[0][1]["text"] == "▼ 1 ▱▱▱▱▱" and marks[1][1]["text"] == "▲ 5 ▰▰▰▰▰"
    assert marks[2][1]["text"] == "● 3 ▰▰▰▱▱" and marks[3] == {}
    assert marks[0][1]["fill"] == "F8D7DA" and marks[1][1]["fill"] == "D1E7DD"
    inverted = scale_marks(rows, [CellScale(column=1, kind="icone", inverted=True)])
    assert inverted[1][1]["text"].startswith("▼")  # more is worse: the biggest points down
    fixed = scale_marks([["A", "50"]], [CellScale(column=1, kind="barra", min=0, max=100)])
    assert fixed[0][1]["text"] == "50 ▰▰▰▱▱"
    _ = Row  # the generator and the preview feed rows of this shape


# ------------------------------------------------------------------ presence (#19), history (#20)


def test_presence_shows_who_else_is_editing(client, template, sessions):
    base = f"/api/reports/templates/{template['id']}/presence"
    assert ok(client.post(base, json={"slide": "256:*:*"}, headers={"X-Remote-User": "ana"}))["others"] == []
    [other] = ok(client.post(base, json={"slide": "257:*:*"}, headers={"X-Remote-User": "bia"}))["others"]
    assert other["author"] == "ana" and other["slide_key"] == "256:*:*"
    with session_scope(sessions) as session:  # ana closed the tab a while ago
        session.get(s.EditorPresenceRow, (template["id"], "ana")).seen_at = datetime.now(UTC) - timedelta(
            minutes=5
        )
    assert ok(client.post(base, json={"slide": "257:*:*"}, headers={"X-Remote-User": "bia"}))["others"] == []
    assert client.delete(base, headers={"X-Remote-User": "bia"}).status_code == 204


def test_one_slide_is_restored_without_touching_the_others(client, template):
    base = f"/api/reports/templates/{template['id']}"
    first = {
        "version": 1,
        "slides": {
            "256:*:*": {"shapes": {"sh2": {"frame": {"x": 100}}}},
            "257:*:*": {"shapes": {"sh2": {"frame": {"x": 1}}}},
        },
    }
    ok(client.put(f"{base}/edits", json={"edits": first}))
    [v1] = ok(client.get(f"{base}/versions", params={"kind": "edits"}))["items"]
    ok(client.post(f"{base}/versions", json={"kind": "edits", "label": "marco"}))
    second = {
        "version": 1,
        "slides": {
            "256:*:*": {"shapes": {"sh2": {"frame": {"x": 999}}}},
            "257:*:*": {"shapes": {"sh2": {"frame": {"x": 2}}}},
        },
    }
    ok(client.put(f"{base}/edits", json={"edits": second}))
    history = ok(client.get(f"{base}/slides/256:*:*/history"))["items"]
    assert len(history) >= 2 and history[-1]["version_id"] == v1["id"]
    restored = ok(client.post(f"{base}/slides/256:*:*/restore", json={"version_id": v1["id"]}))
    slides = restored["edits"]["slides"]
    assert slides["256:*:*"]["shapes"]["sh2"]["frame"]["x"] == 100
    assert slides["257:*:*"]["shapes"]["sh2"]["frame"]["x"] == 2  # untouched


# ------------------------------------------------------------------ diagnostics (#24), usage (#30)


def test_diagnostics_flag_orphan_markers_and_unbound_areas(client, template):
    report = ok(client.get(f"/api/reports/templates/{template['id']}/diagnostics"))
    checks = {c["id"]: c for c in report["checks"]}
    assert report["slides"] == len(template["manifest"]["slides"]) and report["size_mb"] > 0
    assert checks["areas_sem_vinculo"]["count"] > 0  # nothing is mapped yet
    assert {"marcadores_orfaos", "areas_sem_vinculo"} >= {
        c for c in checks if checks[c]["severity"] in {"erro", "info"}
    } - {"slides_sem_nome"}


def test_usage_counts_problems_dropped_and_skipped_slides(client, template, sessions):
    slide_id = template["manifest"]["slides"][0]["slide_id"]
    runs = [
        ok(client.post("/api/reports/generate", json={"template_id": template["id"], "render": False}))
        for _ in range(3)
    ]
    with session_scope(sessions) as session:
        for run in runs:
            row = session.get(s.ReportRunRow, run["id"])
            row.validation = {**row.validation, "dropped": [slide_id]}
    for run in runs[:2]:
        assert (
            client.post(
                f"/api/reports/runs/{run['id']}/presented", json={"segundos": {f"{slide_id}:0:1": 1.5}}
            ).status_code
            == 204
        )
    usage = ok(client.get(f"/api/reports/templates/{template['id']}/usage"))
    assert usage["runs"] == 3 and usage["dropped"][0] == {
        **usage["dropped"][0],
        "slide_id": slide_id,
        "count": 3,
    }
    assert usage["presented"][0]["avg_seconds"] == 1.5
    assert any("saiu vazio" in s for s in usage["suggestions"]) and any(
        "pulado" in s for s in usage["suggestions"]
    )
