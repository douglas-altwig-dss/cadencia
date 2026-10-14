"""Report use cases shared by the API, the CLI and the scheduler (RF-37 … RF-44)."""

from __future__ import annotations

import hashlib
import io
import re
import secrets
import shutil
import tempfile
import time
import zipfile
from collections.abc import Callable, Iterable, Mapping
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.parts.image import Image as PptxImage
from sqlalchemy import select
from sqlalchemy.orm import Session

from cadencia.config import Settings
from cadencia.domain.calendar import last_checkpoints
from cadencia.errors import Invalid, NotFound
from cadencia.i18n import DECK_LOCALE, STORAGE_LOCALE, Msg, localize, render, stored_text, tr, using_locale
from cadencia.logging import new_run_id
from cadencia.report.adapters.jira_source import JiraSource
from cadencia.report.core.a11y import audit
from cadencia.report.core.a11y import summary as a11y_summary
from cadencia.report.core.context import (
    FIELD_INDEX,
    FIELDS,
    FIXED_PREFIX,
    GROUP_IDS,
    MISSING,
    Kind,
    VariableSpec,
    catalog,
    fixed_text_spec,
    resolve,
    variable_name,
)
from cadencia.report.core.edits import DeckEdits, apply_deck_edits
from cadencia.report.core.generate import Problem, expand_notes, generate_deck
from cadencia.report.core.manifest import extract_manifest
from cadencia.report.core.mapping import DeckMapping
from cadencia.report.core.resolve_preview import slide_preview
from cadencia.report.core.scene import deck_scene, picture_blob
from cadencia.report.core.text import MARKER, marker_base
from cadencia.report.core.validate import render_previews, validate_deck
from cadencia.report.ports import DeckSource
from cadencia.store import repo
from cadencia.store import schema as s

JSONDict = dict[str, Any]


class _MissingBoard:
    """Stand-in for a board id that no longer resolves, so a space filter can skip it quietly."""

    space_key: str | None = None


_NO_BOARD = _MissingBoard()


def variable_lookup(
    session: Session, template_id: int | None = None
) -> Callable[[str | None, int | None], dict[str, str]]:
    """Per-board view of the variables a template may use (RF-52), read once per generation.

    ``board_id=None`` is the deck-level view: no slide, no board, no space, so the answer is the
    global value or the registration's default. A name with neither is simply absent from the
    result — that absence is what the generator refuses.
    """

    def lookup(space_key: str | None, board_id: int | None) -> dict[str, str]:
        return repo.resolve_variables(
            session, template_id=template_id, space_key=space_key, board_id=board_id
        )

    return lookup


def load_mapping(session: Session, template_id: int) -> DeckMapping:
    """A template's mapping with the shared status configuration underneath it (RF-62).

    Every caller that resolves a colour or a semáforo label has to go through here. The mapping on
    disk holds only what *this* template overrides; the rest comes from the variable manager, which
    is the single place the three colours are configured.
    """
    row = repo.get_mapping(session, template_id)
    mapping = DeckMapping.model_validate(row.mapping) if row else DeckMapping()
    palette, labels = repo.deck_palette(session)
    return mapping.inheriting(palette, labels)


def deck_source(session: Session, settings: Settings) -> DeckSource:
    """RF-60: the source this install fills decks from.

    One function, so a project carrying the module elsewhere replaces a single line rather than
    hunting through the orchestration. Everything below asks the port, never `analytics` directly.
    """
    return JiraSource(session, settings)


def source_kind(key: str | None) -> Kind:
    """RF-59: the kind a variable inherits from the field it reads, or ``text`` when it reads none.

    Derived rather than stored on purpose. A stored kind can disagree with its source the moment
    somebody repoints the variable, and a variable claiming to be ``rows`` over a ``text`` field
    would fill a table with nothing and say nothing about why.
    """
    if not key:
        return "text"
    spec = next((f for f in FIELDS if f.key == key), None)
    return spec.kind if spec is not None else "text"


def variable_registry(session: Session, template_id: int | None = None) -> dict[str, VariableSpec]:
    """The registered variables, which is the single source of truth for which names exist."""
    return {
        name: VariableSpec(
            name=name,
            label=reg.label,
            default=reg.default,
            source_field=reg.source_field,
            kind=source_kind(reg.source_field),
        )
        for name, reg in repo.registered_variables(session, template_id).items()
    }


def check_writable(session: Session, variable_id: int) -> None:
    """RF-59: refuse a hand-written value on a variable that carries something nobody can type.

    Lives here, not in the store, because deciding it needs the field catalogue — and a second
    copy of "which kinds are collections" would drift from `FIELDS` the first time one is added.
    """
    row = repo.get_variable(session, variable_id)
    if row is None or not row.source_field:
        return
    kind = source_kind(row.source_field)
    if kind != "text":
        raise Invalid(Msg("deck_service.writable.typed", name=row.name, source=row.source_field, kind=kind))


def check_source_field(key: str) -> str:
    """Validate a variable's source field, or say why it cannot be one.

    RF-59 lifted the old text/number/status restriction: a variable is typed by its source, so a
    ``rows`` field makes a ``rows`` variable that fills a table. What stays refused is a ``texto.*``
    source — a variable reading from a variable could be made to point at itself, and a cycle in
    `resolve` is an infinite loop on a user's deck, not a warning.
    """
    if variable_name(key) is not None:
        raise Invalid(Msg("deck_service.source.is_variable", key=key))
    spec = next((f for f in FIELDS if f.key == key), None)
    if spec is None:
        raise Invalid(Msg("deck_service.source.unknown", key=key))
    return spec.key


def catalog_extras(session: Session, template_id: int | None = None) -> list[object]:
    """The variables as catalogue entries, so they appear in the picker like any field.

    RF-59: each carries the kind it inherits from its source, which is what lets a variable be
    offered for a table or a chart binding instead of only for a text one.
    """
    return [
        fixed_text_spec(spec.name, spec.label, spec.kind, spec.source_field)
        for spec in variable_registry(session, template_id).values()
    ]


#: How far a template may expand when unzipped, as a multiple of the upload limit. Real decks sit
#: around 1.2-1.5x (their images are already compressed); a zip bomb is thousands.
UNPACKED_FACTOR = 10


def check_package(content: bytes, max_bytes: int) -> None:
    """Refuse a .pptx that is not a zip, or that would unpack past ``max_bytes`` in python-pptx."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as package:
            unpacked = sum(info.file_size for info in package.infolist())
    except zipfile.BadZipFile as exc:
        raise Invalid(Msg("deck_service.upload.not_zip")) from exc
    if unpacked > max_bytes:
        raise Invalid(Msg("deck_service.upload.too_big", mb=max_bytes // (1024 * 1024)))


def store_template(
    session: Session, settings: Settings, filename: str, content: bytes, name: str | None = None
) -> s.TemplateRow:
    check_package(content, settings.template_max_mb * UNPACKED_FACTOR * 1024 * 1024)
    digest = hashlib.sha256(content).hexdigest()
    existing = session.scalar(select(s.TemplateRow).where(s.TemplateRow.sha256 == digest))
    if existing is not None:
        return existing
    settings.templates_dir.mkdir(parents=True, exist_ok=True)
    path = settings.templates_dir / f"{digest}.pptx"
    if not path.exists():
        path.write_bytes(content)
    try:
        manifest = extract_manifest(path)
    except Exception as exc:
        path.unlink(missing_ok=True)
        raise Invalid(Msg("deck_service.upload.not_pptx", error=str(exc))) from exc
    manifest["template"]["filename"] = filename
    return repo.add_template(
        session,
        name=name or Path(filename).stem,
        filename=filename,
        sha256=digest,
        path=str(path),
        manifest=manifest,
    )


def remove_template(session: Session, settings: Settings, template: s.TemplateRow) -> dict[str, object]:
    """Delete a template, its mapping, its edits and the files that only it uses.

    The .pptx is content-addressed, so the row and the file are one to one and the file can go with
    the row. Generated decks are left alone — see :func:`cadencia.store.repo.delete_template`.
    """
    runs = repo.runs_for_template(session, template.id)
    stored = Path(template.path) if template.path else None
    previews = settings.templates_dir / template.sha256
    repo.delete_template(session, template.id)
    removed = []
    if stored is not None and stored.is_file():
        stored.unlink(missing_ok=True)
        removed.append(stored.name)
    if previews.is_dir():
        shutil.rmtree(previews, ignore_errors=True)
        removed.append(f"{template.sha256}/")
    return {"deleted": True, "id": template.id, "name": template.name, "runs_kept": runs, "files": removed}


def template_previews(settings: Settings, template: s.TemplateRow) -> dict[str, object]:
    """Slide images of the template itself, so the mapping screen can draw its areas on them.

    Rendered once per template version and cached beside the .pptx: the file is content-addressed,
    so a directory named after the digest can never hold images of a different deck.
    """
    target = settings.templates_dir / template.sha256
    images = sorted(p.name for p in target.glob("slide-*.png")) if target.is_dir() else []
    if images:
        return {"rendered": True, "images": images}
    if not Path(template.path).exists():
        return {"rendered": False, "images": [], "reason": tr("deck_service.previews.file_missing")}

    # Render into a scratch directory and move it into place, so two requests racing on the first
    # view of a template cannot serve a half-written set of images.
    settings.templates_dir.mkdir(parents=True, exist_ok=True)
    scratch = Path(tempfile.mkdtemp(prefix=f"{template.sha256[:12]}-", dir=settings.templates_dir))
    try:
        result = render_previews(Path(template.path), scratch, dpi=90)
        if not any(scratch.glob("slide-*.png")):
            reason = render(result.get("reason") or Msg("deck_service.previews.nothing"))
            return {"rendered": False, "images": [], "reason": reason}
        try:
            scratch.rename(target)
        except OSError:  # another request rendered it first
            shutil.rmtree(scratch, ignore_errors=True)
        return {"rendered": True, "images": sorted(p.name for p in target.glob("slide-*.png"))}
    finally:
        shutil.rmtree(scratch, ignore_errors=True)


def _sample(kind: str | None, value: object, limit: int = 6) -> dict[str, object]:
    """One field's value as short lines the mapping screen can show beside the field picker."""
    lines: list[str] = []
    count: int | None = None
    if kind in {"text", "status"}:
        lines = [str(value)] if value else []
    elif kind == "number":
        lines = [] if value is None else [f"{float(value):g}"]
    elif kind == "groups":
        groups = list(value or [])
        count = len(groups)
        for group in groups:
            lines.append(group.title)
            lines += [f"• {item}" for item in group.items]
    elif kind == "rows":
        rows = list(value or [])
        count = len(rows)
        lines = [" · ".join(filter(None, (r.values.get("chave"), r.values.get("resumo")))) for r in rows]
    elif kind == "row_groups":
        groups = list(value or [])
        count = sum(len(g.rows) for g in groups)
        for group in groups:
            lines.append(group.title)
            for row in group.rows:
                lines.append(f"• {row.values.get('chave', '')} {row.values.get('resumo', '')}".strip())
    elif kind == "series":
        series = value
        if series is not None and getattr(series, "categories", None):
            count = len(series.categories)
            lines = [f"{name}: {', '.join(str(v) for v in values)}" for name, values in series.series.items()]
    return {
        "kind": kind,
        "empty": not lines,
        "count": count,
        "lines": [line[:120] for line in lines[:limit]],
        "truncated": len(lines) > limit,
    }


def preview_fields(
    session: Session,
    settings: Settings,
    *,
    board_id: int | None,
    reference: datetime | None,
    now: datetime,
) -> dict[str, object]:
    """What every catalogue field would produce right now, for one board (RF-38).

    The mapping screen shows this next to the field picker: "entregas.por_epico" and
    "proximas.por_epico" are indistinguishable by name alone, but never by their contents.
    """
    board_row = session.get(s.BoardRow, board_id) if board_id is not None else None
    if board_row is None:
        board_row = next(iter(repo.selected_boards(session)), None)
    if board_row is None:
        return {"board": None, "reference": None, "fields": {}}

    source = deck_source(session, settings)
    reference = reference or default_reference(session, settings, [board_row.id], now)
    deck = source.deck_data(
        [board_row.id],
        reference=reference,
        now=now,
        registry=variable_registry(session),
        variables=variable_lookup(session),
    )
    board = deck.boards[0] if deck.boards else None
    specs = [*source.catalog(), *catalog_extras(session)]
    with using_locale(DECK_LOCALE):  # samples of deck content, as the deck will print them
        fields = {spec.key: _sample(*resolve(deck, board, spec.key)) for spec in specs}
    return {
        "board": {
            "id": board_row.id,
            "name": board_row.name,
            "space_key": board_row.space_key,
            "space_name": board_row.space_name,
        },
        "reference": reference.isoformat(),
        "fields": fields,
    }


def preview_slide(
    session: Session,
    settings: Settings,
    template: s.TemplateRow,
    *,
    slide_id: int,
    mapping: DeckMapping,
    board_id: int | None,
    reference: datetime | None,
    now: datetime,
) -> dict[str, object]:
    """Resolve one slide against one board, for the live preview (RF-38).

    The draft mapping is passed in rather than read from the database, so the preview follows the
    edit being made instead of the last save.
    """
    manifest_slide = next(
        (sl for sl in template.manifest.get("slides", []) if sl.get("slide_id") == slide_id), None
    )
    if manifest_slide is None:
        raise NotFound(slide_id, Msg("deck_service.slide.not_found", slide=slide_id, template=template.id))
    spec = mapping.for_slide(slide_id)
    # A slide pinned to one board previews that board, whatever the picker says: that is what will
    # be generated.
    wanted = spec.board_id if spec.scope == "board" and spec.board_id else board_id
    board_row = session.get(s.BoardRow, wanted) if wanted is not None else None
    if board_row is None:
        board_row = next(iter(repo.selected_boards(session)), None)

    deck_board = None
    if board_row is not None:
        reference = reference or default_reference(session, settings, [board_row.id], now)
        deck = deck_source(session, settings).deck_data(
            [board_row.id],
            reference=reference,
            now=now,
            anonymize=mapping.anonymize,
            registry=variable_registry(session, template.id),
            variables=variable_lookup(session, template.id),
            slide_variables=repo.slide_variable_values(session, template.id),
            template_name=template.name,
        )
        deck_board = deck.boards[0] if deck.boards else None
    else:
        reference = reference or now
        deck = deck_source(session, settings).deck_data(
            [],
            reference=reference,
            now=now,
            registry=variable_registry(session, template.id),
            variables=variable_lookup(session, template.id),
            slide_variables=repo.slide_variable_values(session, template.id),
            template_name=template.name,
        )

    # What the slide will say is deck content, resolved exactly as generation does — in the deck's
    # language whatever the request's (docs/i18n-backend.md, 2.8). Only the warnings are the
    # reader's, rendered after the block.
    with using_locale(DECK_LOCALE):
        result = slide_preview(manifest_slide, spec, mapping, deck, deck_board)
    result["warnings"] = [render(w) for w in result["warnings"]]
    result["board"] = (
        {
            "id": board_row.id,
            "name": board_row.name,
            "space_key": board_row.space_key,
            "space_name": board_row.space_name or board_row.space_key,
        }
        if board_row
        else None
    )
    result["reference"] = reference.isoformat()
    return result


def load_edits(session: Session, template_id: int) -> DeckEdits:
    row = session.get(s.DeckEditRow, template_id)
    return DeckEdits.model_validate(row.edits) if row else DeckEdits()


def save_edits(
    session: Session, template_id: int, edits: DeckEdits, now: datetime, *, author: str | None = None
) -> s.DeckEditRow:
    row = session.get(s.DeckEditRow, template_id)
    if row is None:
        row = s.DeckEditRow(template_id=template_id, revision=0, edits={})
        session.add(row)
    row.edits = edits.model_dump()
    row.revision = (row.revision or 0) + 1
    row.updated_at = now
    session.flush()
    # RF-87: every save leaves a restorable copy behind (one rolling copy per ten minutes).
    repo.record_version(session, template_id, "edits", row.edits, author=author, now=now)
    return row


# ---------------------------------------------------------------- versions (RF-87)


def current_document(session: Session, template_id: int, kind: str) -> dict[str, Any]:
    if kind == "edits":
        return load_edits(session, template_id).model_dump()
    row = repo.get_mapping(session, template_id)
    return dict(row.mapping) if row else {}


def version_diff(kind: str, current: dict[str, Any], version: dict[str, Any]) -> dict[str, Any]:
    """What restoring ``version`` would change, in words a person can decide on."""
    if kind == "edits":
        now_slides = current.get("slides") or {}
        old_slides = version.get("slides") or {}
        changed = []
        for key in sorted(now_slides.keys() | old_slides.keys()):
            a, b = now_slides.get(key) or {}, old_slides.get(key) or {}
            if a == b:
                continue
            shapes_a, shapes_b = a.get("shapes") or {}, b.get("shapes") or {}
            shapes = sorted(
                k for k in shapes_a.keys() | shapes_b.keys() if shapes_a.get(k) != shapes_b.get(k)
            )
            other = sorted(
                f
                for f in ("notes", "cells", "ungrouped", "groups", "locked", "hidden", "tables")
                if a.get(f) != b.get(f)
            )
            changed.append({"slide": key, "shapes": shapes, "other": other})
        deck = sorted(f for f in ("order", "deleted", "duplicates") if current.get(f) != version.get(f))
        return {"slides": changed, "deck": deck, "same": not changed and not deck}
    now_slides = current.get("slides") or {}
    old_slides = version.get("slides") or {}
    changed = [
        {
            "slide": key,
            "bindings": len((old_slides.get(key) or {}).get("bindings") or []),
            "current_bindings": len((now_slides.get(key) or {}).get("bindings") or []),
        }
        for key in sorted(now_slides.keys() | old_slides.keys())
        if now_slides.get(key) != old_slides.get(key)
    ]
    deck = sorted(
        k for k in (current.keys() | version.keys()) - {"slides"} if current.get(k) != version.get(k)
    )
    return {"slides": changed, "deck": deck, "same": not changed and not deck}


def restore_version(
    session: Session,
    template: s.TemplateRow,
    version: s.DeckVersionRow,
    now: datetime,
    author: str | None = None,
) -> dict[str, Any]:
    """Put ``version`` back, keeping what is being replaced as a named version of its own."""
    current = current_document(session, template.id, version.kind)
    repo.record_version(
        session,
        template.id,
        version.kind,
        current,
        # Written once in the requester's language and not re-translated (decision 8).
        label=tr("deck_service.version.before_restore", version=version.id),
        auto=False,
        author=author,
        now=now,
    )
    if version.kind == "edits":
        saved = save_edits(
            session, template.id, DeckEdits.model_validate(version.document), now, author=author
        )
        return {"kind": "edits", "revision": saved.revision}
    mapping = DeckMapping.model_validate(version.document)
    repo.save_mapping(session, template.id, template.sha256, mapping.model_dump())
    return {"kind": "mapping"}


def asset_paths(session: Session) -> dict[str, str]:
    """Every stored image by id, for the edit applier to place on a slide."""
    return {row.id: row.path for row in session.scalars(select(s.AssetRow))}


#: The image types python-pptx can place on a slide, which is everything `store_asset` accepts.
IMAGE_TYPES = frozenset({"image/png", "image/jpeg", "image/gif", "image/bmp", "image/tiff", "image/x-wmf"})


def store_asset(session: Session, settings: Settings, filename: str, content: bytes) -> s.AssetRow:
    """Keep an image the editor placed on a slide.

    Only what python-pptx can put on a slide is accepted, and the content type comes from the bytes,
    never from the client: the asset is served back from the app's own origin, so a file that merely
    *claimed* to be an image could be an HTML page running as part of the app.
    """
    try:
        image = PptxImage.from_blob(content)
        width, height = image.size
    except Exception as exc:  # PIL has its own zoo of errors for bytes that are not an image
        raise Invalid(Msg("deck_service.asset.not_image")) from exc
    digest = hashlib.sha256(content).hexdigest()
    existing = session.get(s.AssetRow, digest)
    if existing is not None:
        return existing
    assets = settings.data_dir / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    path = assets / f"{digest}.{image.ext}"
    if not path.exists():
        path.write_bytes(content)
    row = s.AssetRow(
        id=digest,
        filename=filename,
        content_type=image.content_type,
        width=width,
        height=height,
        path=str(path),
        uploaded_at=datetime.now(UTC),
    )
    session.add(row)
    session.flush()
    return row


def run_scene(session: Session, run: s.ReportRunRow) -> dict[str, object]:
    """The generated deck read back as the editable scene (RF-46)."""
    if not run.path or not Path(run.path).exists():
        raise NotFound(run.id, Msg("deck_service.run.file_missing", run=run.id))
    # The template may have been deleted since; the deck still opens, it just loses the theme and
    # the "this shape is data-bound" markers, which both live in the template.
    template = session.get(s.TemplateRow, run.template_id)
    mapping = load_mapping(session, run.template_id)
    scene = deck_scene(
        run.path,
        list(run.validation.get("produced") or []),
        mapping,
        load_edits(session, run.template_id),
        theme=template.manifest.get("theme") if template else None,
    )
    scene["template_deleted"] = template is None
    scene["run_id"] = run.id
    scene["template_id"] = run.template_id
    scene["reference_at"] = run.reference_at.isoformat()
    return scene


def reassociate_run(session: Session, run: s.ReportRunRow, template_id: int) -> s.ReportRunRow:
    """#12: point an orphaned run (its own template deleted) at a re-uploaded one.

    Only for a genuinely orphaned run — a run whose template still exists already has a home, and
    relinking it would silently swap which mapping and theme it answers to. The new template does
    not have to be a byte-for-byte match: `deck_scene` looks up each produced slide's theme and
    bound-field marks by its *origin slide id*, so a template with different ids simply keeps
    showing none of that, exactly as it did before this call — this can only help, never corrupt.
    """
    if session.get(s.TemplateRow, run.template_id) is not None:
        raise Invalid(Msg("deck_service.run.has_template"))
    if session.get(s.TemplateRow, template_id) is None:
        raise NotFound(template_id, Msg("deck_service.template.not_found"))
    run.template_id = template_id
    session.flush()
    return run


def template_scene(session: Session, template: s.TemplateRow) -> dict[str, object]:
    """The **template** read back as a scene, through the same code a generated run uses (RF-61).

    This is what makes one renderer possible. The mapping screen used to draw a LibreOffice PNG
    with clickable boxes over it — a picture of a slide, at 90 dpi, obviously an image — while the
    editor drew the real geometry. Two renderings of the same slide disagree the moment either
    changes. A template is just a .pptx with nothing filled in yet, so `deck_scene` reads it with
    no `produced` entries and the canvas draws it exactly as it draws a run.

    The mapping still travels with it, so a shape already bound is marked as such before anything
    has been generated — which is the one thing the old hotspot view was genuinely good at.
    """
    if not Path(template.path).exists():
        raise NotFound(template.id, Msg("deck_service.template.file_missing", template=template.id))
    mapping = load_mapping(session, template.id)
    scene = deck_scene(template.path, [], mapping, None, theme=template.manifest.get("theme"))
    scene["template_id"] = template.id
    scene["template_name"] = template.name
    # No run, so nothing has been filled: the canvas shows the template's own words, which is
    # precisely what "marcador" mode (RF-57) wants to show anyway.
    scene["run_id"] = None
    return scene


def template_picture(template: s.TemplateRow, shape_key: str) -> tuple[bytes, str] | None:
    """A picture embedded in the template, for the canvas drawing the template."""
    return picture_blob(template.path, shape_key) if Path(template.path).exists() else None


def default_reference(
    session: Session, settings: Settings, board_ids: Iterable[int], now: datetime
) -> datetime:
    """Most recent checkpoint (e.g. last Friday 18h) of the first board, or now."""
    for board_id in board_ids:
        profile = repo.board_profile(session, board_id, repo.profile_defaults(settings))
        marks = last_checkpoints(profile, now, 1)
        if marks:
            return marks[-1]
    return now


def run_boards(session: Session, mapping: DeckMapping, board_ids: Iterable[int] | None = None) -> list[int]:
    """The boards a run of this mapping fills: the ones asked for, else the mapping's, else every
    selected board; narrowed by space, plus any board a slide is pinned to."""
    wanted = list(board_ids or mapping.board_ids or [b.id for b in repo.selected_boards(session)])
    # RF-55: a chosen space narrows the run to its boards. Empty means every space, so a mapping
    # saved before this existed covers exactly what it always covered.
    if mapping.space_keys and board_ids is None:
        allowed = set(mapping.space_keys)
        wanted = [b for b in wanted if (session.get(s.BoardRow, b) or _NO_BOARD).space_key in allowed]
    for spec in mapping.slides.values():
        if spec.scope == "board" and spec.board_id and spec.board_id not in wanted:
            wanted.append(spec.board_id)
    return wanted


def generate_report(
    session: Session,
    settings: Settings,
    template_id: int,
    *,
    reference: datetime | None = None,
    board_ids: Iterable[int] | None = None,
    anonymize: bool | None = None,
    render: bool = True,
    now: datetime | None = None,
    kind: str = "entrega",
    audience: str = "",
    only_slides: Iterable[int] | None = None,
    parent_run_id: str | None = None,
) -> s.ReportRunRow:
    """Generate a deck. ``kind`` "previa" is one board looked at before generating (#9) and never
    listed; "parcial" regenerates ``only_slides`` for some boards, from ``parent_run_id`` (#10);
    ``audience`` keeps the slides meant for that variant (#25)."""
    started = time.monotonic()
    now = now or datetime.now(UTC)
    template = session.get(s.TemplateRow, template_id)
    if template is None:
        raise NotFound(template_id, Msg("deck_service.template.not_found_id", template=template_id))
    mapping_row = repo.get_mapping(session, template_id)
    mapping = load_mapping(session, template_id)

    wanted = run_boards(session, mapping, board_ids)
    reference = reference or default_reference(session, settings, wanted, now)
    deck = deck_source(session, settings).deck_data(
        wanted,
        reference=reference,
        now=now,
        anonymize=mapping.anonymize if anonymize is None else anonymize,
        registry=variable_registry(session, template_id),
        variables=variable_lookup(session, template_id),
        slide_variables=repo.slide_variable_values(session, template_id),
        template_name=template.name,
    )
    run_id = new_run_id()
    settings.decks_dir.mkdir(parents=True, exist_ok=True)
    out = settings.decks_dir / f"{run_id}.pptx"
    if audience and audience not in mapping.audiences:
        raise Invalid(Msg("deck_service.audience.unknown", audience=audience))
    edits = load_edits(session, template_id)
    # Everything that writes into the .pptx runs in the deck's language (R5): `deck_data` already
    # pins it while the data is built, and this pins it again for anything a value renders lazily
    # on its way into a shape. The problems it reports are messages, rendered further down.
    with using_locale(DECK_LOCALE):
        result = generate_deck(
            template.path,
            mapping,
            deck,
            out,
            theme=template.manifest.get("theme"),
            audience=audience,
            only_slides=set(only_slides) if only_slides is not None else None,
        )
        # The data is in; now replay the manual edits on top, which is what makes a re-run keep them.
        if edits.touched():
            prs = Presentation(str(out))
            problems, placement = apply_deck_edits(prs, result.produced, edits, asset_paths(session))
            prs.save(str(out))
            result.slide_count = len(prs.slides)
            for entry in result.produced:
                key = f"{entry['origin_slide_id']}:{entry.get('board_id') or 0}:{entry.get('page', 1)}"
                if key in placement:
                    entry["index"] = placement[key]
            result.produced.sort(key=lambda r: r["index"])
            for message in problems:
                result.problems.append(Problem("edicao_perdida", message))
            # RF-94: notes typed in the editor may carry markers too; fill them for each slide's board.
            prs = Presentation(str(out))
            expand_notes(prs, result.produced, deck, mapping, result)
            prs.save(str(out))
    validation = validate_deck(
        result,
        mapping,
        mapped_sha256=mapping_row.template_sha256 if mapping_row else None,
        current_manifest=template.manifest,
        render_dir=settings.decks_dir / run_id if render else None,
    )
    if not mapping_row:
        validation["warnings"].insert(
            0,
            {"kind": "sem_mapeamento", "message": Msg("deck_service.no_mapping").to_json()},
        )
    validation["produced"] = result.produced
    validation["edited_slides"] = sorted(edits.touched())
    validation["dropped"] = result.dropped
    findings = audit(deck_scene(out, result.produced, mapping, edits))
    validation["accessibility"] = {"summary": a11y_summary(findings), "findings": findings[:60]}
    # Not ``kind``: that name is this run's kind (entrega, previa, parcial) and is read below.
    for finding_kind, count in a11y_summary(findings).items():
        validation["warnings"].append(
            {
                "kind": f"acessibilidade_{finding_kind}",
                "message": Msg(f"{A11Y_SUMMARY}{finding_kind}", n=count).to_json(),
            }
        )
    if kind == "previa":
        prune_previews(session, settings, template_id)
    return repo.add_report_run(
        session,
        id=run_id,
        template_id=template_id,
        template_sha256=template.sha256,
        reference_at=reference,
        board_ids=[b.board_id for b in deck.boards],
        sprint_ids=[b.sprint_id for b in deck.boards if b.sprint_id],
        status="ok" if validation["ok"] else "invalid",
        path=str(out),
        validation=stored_validation(validation),
        anonymized=mapping.anonymize if anonymize is None else anonymize,
        kind=kind,
        audience=audience,
        parent_run_id=parent_run_id,
        seconds=round(time.monotonic() - started, 2),
    )


#: The a11y warnings' keys: ``deck_service.a11y_summary.<finding kind>`` (one per kind `a11y.audit`
#: reports: sem_texto_alternativo, contraste_baixo, ordem_de_leitura).
A11Y_SUMMARY = "deck_service.a11y_summary."


def run_accessibility(session: Session, run: s.ReportRunRow) -> dict[str, Any]:
    findings = audit(run_scene(session, run))
    return {"summary": a11y_summary(findings), "findings": localize(findings)}


def template_accessibility(session: Session, template: s.TemplateRow) -> dict[str, Any]:
    findings = audit(template_scene(session, template))
    return {"summary": a11y_summary(findings), "findings": localize(findings)}


# ---------------------------------------------------------------- validation text (R4)


def _stored_entry(entry: Any) -> Any:
    """One validation entry as it is persisted: the ``$msg`` record in ``msg`` and its pt-BR text
    in ``message``, so old readers, rollbacks and ad-hoc SQL keep reading a sentence."""
    if not isinstance(entry, dict) or not Msg.is_json(entry.get("message")):
        return entry
    return {**entry, "msg": entry["message"], "message": render(entry["message"], STORAGE_LOCALE)}


def stored_validation(validation: Mapping[str, Any]) -> dict[str, Any]:
    """``validate_deck``'s report (messages as ``$msg`` records) in the shape ``report_run`` keeps
    (docs/i18n-backend.md, 2.7). ``info`` keeps its strings, with the records in ``info_msgs``;
    ``preview.reason`` likewise in ``reason_msg``."""
    out = dict(validation)
    for name in ("errors", "warnings"):
        if name in out:
            out[name] = [_stored_entry(e) for e in out[name] or []]
    info = list(out.get("info") or [])
    if any(Msg.is_json(i) for i in info):
        out["info_msgs"] = [i if Msg.is_json(i) else None for i in info]
        out["info"] = [render(i, STORAGE_LOCALE) for i in info]
    preview = out.get("preview")
    if isinstance(preview, dict) and Msg.is_json(preview.get("reason")):
        out["preview"] = {
            **preview,
            "reason_msg": preview["reason"],
            "reason": render(preview["reason"], STORAGE_LOCALE),
        }
    accessibility = out.get("accessibility")
    if isinstance(accessibility, dict) and accessibility.get("findings"):
        out["accessibility"] = {
            **accessibility,
            "findings": [_stored_entry(f) for f in accessibility["findings"]],
        }
    return out


def _read_entry(entry: Any) -> Any:
    if not isinstance(entry, dict):
        return entry
    out = {k: v for k, v in entry.items() if k != "msg"}
    if "message" in entry:
        out["message"] = stored_text(entry.get("msg"), entry.get("message"))
    return out


def localized_validation(validation: Mapping[str, Any] | None) -> dict[str, Any]:
    """A stored validation for a reader: every message rendered from its record in the ambient
    language, the records themselves left out. Rows stored before the records existed keep their
    stored pt-BR text (decision 7)."""
    out = dict(validation or {})
    for name in ("errors", "warnings"):
        if name in out:
            out[name] = [_read_entry(e) for e in out[name] or []]
    records = out.pop("info_msgs", None)
    if records is not None:
        stored = list(out.get("info") or [])
        out["info"] = [
            stored_text(record, stored[i] if i < len(stored) else None) for i, record in enumerate(records)
        ]
    preview = out.get("preview")
    if isinstance(preview, dict) and "reason_msg" in preview:
        out["preview"] = {k: v for k, v in preview.items() if k != "reason_msg"}
        out["preview"]["reason"] = stored_text(preview.get("reason_msg"), preview.get("reason"))
    accessibility = out.get("accessibility")
    if isinstance(accessibility, dict) and accessibility.get("findings"):
        out["accessibility"] = {
            **accessibility,
            "findings": [_read_entry(f) for f in accessibility["findings"]],
        }
    return out


#: Previews kept per template; older ones (and their files) go — they were never meant to last.
KEEP_PREVIEWS = 3


def prune_previews(session: Session, settings: Settings, template_id: int) -> None:
    rows = session.scalars(
        select(s.ReportRunRow)
        .where(s.ReportRunRow.template_id == template_id, s.ReportRunRow.kind == "previa")
        .order_by(s.ReportRunRow.created_at.desc())
    ).all()
    for row in rows[KEEP_PREVIEWS - 1 :]:
        if row.path and Path(row.path).exists():
            Path(row.path).unlink()
        shutil.rmtree(settings.decks_dir / row.id, ignore_errors=True)
        session.delete(row)
    session.flush()


def run_dict(row: s.ReportRunRow) -> dict[str, object]:
    return {
        "id": row.id,
        "template_id": row.template_id,
        "template_sha256": row.template_sha256,
        "created_at": row.created_at.isoformat(),
        "reference_at": row.reference_at.isoformat(),
        "board_ids": row.board_ids,
        "sprint_ids": row.sprint_ids,
        "status": row.status,
        "validation": localized_validation(row.validation),
        "anonymized": row.anonymized,
        "kind": row.kind,
        "parent_run_id": row.parent_run_id,
        "audience": row.audience,
        "review": {
            "status": row.review_status,
            "by": row.reviewed_by,
            "at": row.reviewed_at.isoformat() if row.reviewed_at else None,
            "note": row.review_note,
            "history": row.review_history,
        },
        "seconds": row.seconds,
    }


# ---------------------------------------------------------------- review (#12)

#: action → the states it may start from.
REVIEW_FROM = {
    "enviar": ("rascunho", "reprovado"),
    "aprovar": ("em_revisao",),
    "reprovar": ("em_revisao",),
    "reabrir": ("aprovado", "reprovado", "em_revisao"),
}
REVIEW_TO = {"enviar": "em_revisao", "aprovar": "aprovado", "reprovar": "reprovado", "reabrir": "rascunho"}


def open_comments(session: Session, template_id: int) -> int:
    return sum(
        1 for c in repo.list_comments(session, template_id) if c.resolved_at is None and c.parent_id is None
    )


def review_run(
    session: Session, run: s.ReportRunRow, action: str, author: str | None, note: str = ""
) -> None:
    """Move a deck through rascunho → em_revisao → aprovado | reprovado.

    Approval needs a deck that passed validation and a template with no open comment: approving
    what the validation refused, or while a reviewer's question is unanswered, is exactly the
    decision the workflow exists to prevent."""
    if action not in REVIEW_FROM:
        raise Invalid(Msg("deck_service.review.unknown_action", action=action))
    if run.kind != "entrega":
        raise Invalid(Msg("deck_service.review.only_delivery"))
    if run.review_status not in REVIEW_FROM[action]:
        raise Invalid(Msg("deck_service.review.wrong_state", action=action, state=run.review_status))
    if action == "aprovar":
        if run.status != "ok":
            raise Invalid(Msg("deck_service.review.not_valid"))
        pending = open_comments(session, run.template_id)
        if pending:
            raise Invalid(Msg("deck_service.review.open_comments", count=pending))
    if action == "reprovar" and not note.strip():
        raise Invalid(Msg("deck_service.review.note_required"))
    now = datetime.now(UTC)
    run.review_status = REVIEW_TO[action]
    run.reviewed_by = author
    run.reviewed_at = now
    run.review_note = note.strip()
    # A later decision must not erase an earlier one's author and reason — only the single
    # `review_*` columns above are "the current state"; this list is the record of how it got there.
    run.review_history = [
        *run.review_history,
        {"action": action, "by": author, "at": now.isoformat(), "note": note.strip()},
    ]
    session.flush()


def download_blocker(session: Session, run: s.ReportRunRow) -> dict[str, Any] | None:
    """Why a download needs ``force``: validation first (RF-70), then approval (#12)."""
    if run.status != "ok":
        return {
            "cause": tr("deck_service.blocker.invalid.cause"),
            "action": tr("deck_service.blocker.invalid.action"),
            "validation": localized_validation(run.validation),
        }
    needs_approval = run.kind == "entrega" and load_mapping(session, run.template_id).require_approval
    if needs_approval and run.review_status != "aprovado":
        return {
            "cause": tr("deck_service.blocker.unapproved.cause"),
            "action": tr("deck_service.blocker.unapproved.action"),
        }
    return None


# ---------------------------------------------------------------- PDF (#8)


def run_pdf(settings: Settings, run: s.ReportRunRow) -> Path:
    """The run's PDF: the one validation already rendered, or rendered now."""
    if not run.path or not Path(run.path).exists():
        raise NotFound(run.id, Msg("deck_service.run.file_missing", run=run.id))
    folder = settings.decks_dir / run.id
    existing = folder / f"{Path(run.path).stem}.pdf"
    if existing.exists():
        return existing
    rendered = render_previews(Path(run.path), folder)
    if not rendered.get("pdf"):
        reason = rendered.get("reason")
        raise Invalid(Msg.from_json(reason) if Msg.is_json(reason) else Msg("deck_service.pdf.failed"))
    return Path(rendered["pdf"])


# ---------------------------------------------------------------- compare two runs (#11)


def _cell_fill_hex(cell: Any) -> str | None:
    """A table cell's own solid fill, the way a status colour or a rule is written (RF-92):
    ``None`` for anything inherited or themed, which a diff has no business calling a change."""
    try:
        if cell.fill.type != MSO_FILL_TYPE.SOLID:
            return None
        return str(cell.fill.fore_color.rgb)
    except (AttributeError, TypeError):
        return None


def _chart_lines(shape: Any) -> list[str]:
    """#9: a chart's own data as comparable lines — without this, a value changing behind a bar
    or a status colour changing behind a cell both passed a text-only diff unnoticed."""
    try:
        chart = shape.chart
        categories = list(chart.plots[0].categories)
    except (AttributeError, IndexError, ValueError):
        return []
    lines = []
    for plot in chart.plots:
        for series in plot.series:
            values = ", ".join(f"{cat}={val}" for cat, val in zip(categories, series.values, strict=False))
            lines.append(tr("deck_service.compare.chart", name=str(series.name), values=values))
    return lines


def _run_texts(run: s.ReportRunRow) -> dict[str, dict[str, Any]]:
    """Each produced slide's words, keyed by where it came from (template slide, board, page).

    Table cell fills and chart data are folded in as synthetic lines: a status colour changing
    behind an unchanged word, or a bar's own value changing, are both real differences a
    text-only diff would otherwise call "same"."""
    if not run.path or not Path(run.path).exists():
        raise NotFound(run.id, Msg("deck_service.run.file_missing", run=run.id))
    slides = list(Presentation(run.path).slides)
    out: dict[str, dict[str, Any]] = {}
    for entry in run.validation.get("produced") or []:
        index = int(entry["index"]) - 1
        if not 0 <= index < len(slides):
            continue
        lines: list[str] = []
        for shape in slides[index].shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                lines += [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
                    fills = [_cell_fill_hex(c) for c in row.cells]
                    if any(fills):
                        lines.append(
                            tr("deck_service.compare.colors", fills=" | ".join(f or "-" for f in fills))
                        )
            if getattr(shape, "has_chart", False) and shape.has_chart:
                lines += _chart_lines(shape)
        key = f"{entry.get('origin_slide_id')}:{entry.get('board_id') or 0}:{entry.get('page', 1)}"
        out[key] = {"index": index + 1, "board_id": entry.get("board_id"), "lines": lines}
    return out


def compare_runs(
    before: s.ReportRunRow, after: s.ReportRunRow, board_names: Mapping[int, str]
) -> dict[str, Any]:
    """#11: what the deck says differently between two generations, slide by slide."""
    old, new = _run_texts(before), _run_texts(after)
    slides = []
    for key in sorted(old.keys() | new.keys(), key=lambda k: [int(x) for x in k.split(":")]):
        was, now = old.get(key), new.get(key)
        board_id = (now or was or {}).get("board_id")
        entry: dict[str, Any] = {
            "key": key,
            "board": board_names.get(board_id, "") if board_id else "",
            "before_index": was["index"] if was else None,
            "after_index": now["index"] if now else None,
        }
        if was is None or now is None:
            entry["change"] = "novo" if was is None else "removido"
            slides.append(entry)
            continue
        removed = [line for line in was["lines"] if line not in now["lines"]]
        added = [line for line in now["lines"] if line not in was["lines"]]
        if removed or added:
            entry.update(change="alterado", removed=removed, added=added)
            slides.append(entry)
    return {"before": run_dict(before), "after": run_dict(after), "slides": slides, "same": not slides}


# ---------------------------------------------------------------- share links (#14)

SHARE_DAYS = 14


def create_share(
    session: Session, run: s.ReportRunRow, author: str | None, days: int = SHARE_DAYS
) -> s.DeckShareRow:
    if run.kind != "entrega":
        raise Invalid(Msg("deck_service.share.only_delivery"))
    if not 1 <= days <= 90:
        raise Invalid(Msg("deck_service.share.days"))
    now = datetime.now(UTC)
    row = s.DeckShareRow(
        token=secrets.token_urlsafe(24),
        run_id=run.id,
        created_at=now,
        expires_at=now + timedelta(days=days),
        created_by=author,
    )
    session.add(row)
    session.flush()
    return row


def shared_run(session: Session, token: str) -> s.ReportRunRow:
    """The run behind a share token. Expired, revoked and unknown tokens all answer the same
    NotFound, so probing learns nothing about which tokens ever existed."""
    share = session.get(s.DeckShareRow, token)
    if share is None or share.revoked_at is not None or share.expires_at <= datetime.now(UTC):
        raise NotFound("link", Msg("deck_service.share.not_found"))
    run = session.get(s.ReportRunRow, share.run_id)
    if run is None:
        raise NotFound("link", Msg("deck_service.share.not_found"))
    return run


def share_dict(row: s.DeckShareRow, settings: Settings) -> dict[str, Any]:
    base = settings.public_url.rstrip("/")
    return {
        "token": row.token,
        "run_id": row.run_id,
        "created_at": row.created_at.isoformat(),
        "expires_at": row.expires_at.isoformat(),
        "created_by": row.created_by,
        "revoked": row.revoked_at is not None,
        "views": row.views,
        "path": f"/shared/{row.token}",
        "url": f"{base}/shared/{row.token}" if base else None,
    }


# ---------------------------------------------------------------- the variable audit


def template_variable_names(template: s.TemplateRow, mapping: DeckMapping) -> dict[str, list[str]]:
    """Every ``texto.*`` name this template would ask for, and where it asks for it.

    Four places, because a variable can enter a deck four ways: written into the slide itself as a
    ``{{marker}}``, bound to an area as a field, written inside a binding's literal value, or named
    as a slide's ``hide_when_empty``. Missing any of them would make the audit lie.
    """
    used: dict[str, list[str]] = {}

    def note(key: str, where: str) -> None:
        if variable_name(key) is not None:
            used.setdefault(key, []).append(where)

    for slide in template.manifest.get("slides", []):
        for area in slide.get("areas", []):
            for name in area.get("markers", []):
                note(marker_base(name), f"slide {slide.get('index')} · {area.get('id')}")
    for slide_id, spec in mapping.slides.items():
        if spec.hide_when_empty:
            note(spec.hide_when_empty, f"slide {slide_id} · {tr('deck_service.audit.hide_when_empty')}")
        for binding in spec.bindings:
            if binding.key and binding.value is None:
                note(binding.key, f"slide {slide_id} · {binding.target}")
            for marker in MARKER.finditer(binding.value or ""):
                note(marker_base(marker.group(1)), f"slide {slide_id} · {binding.target}")
    return used


def audit_variables(session: Session, settings: Settings, template_id: int | None = None) -> JSONDict:
    """Which templates would stop generating under the registration rule, and why.

    Read-only on purpose: it names what is wrong and leaves the fixing to a person. Silently
    registering whatever a template happens to mention would recreate, one layer up, exactly the
    implicit registration this rule removes.
    """
    rows = list(session.scalars(select(s.TemplateRow).order_by(s.TemplateRow.id)))
    if template_id is not None:
        rows = [r for r in rows if r.id == template_id]
    boards = list(repo.selected_boards(session))
    results: list[JSONDict] = []
    for template in rows:
        mapping_row = repo.get_mapping(session, template.id)
        mapping = DeckMapping.model_validate(mapping_row.mapping) if mapping_row else DeckMapping()
        used = template_variable_names(template, mapping)
        registry = variable_registry(session, template.id)
        lookup = variable_lookup(session, template.id)
        unregistered = [
            {"name": key, "used_in": where}
            for key, where in sorted(used.items())
            if variable_name(key) not in registry
        ]
        unfilled: list[JSONDict] = []
        for key, where in sorted(used.items()):
            name = variable_name(key)
            if name is None or name not in registry:
                continue
            # A default or a Jira field both mean the name always has somewhere to get a value.
            if registry[name].default is not None or registry[name].source_field:
                continue
            short = [
                b.name
                for b in boards
                if name not in lookup(b.space_key, b.id)
                # A slide-scoped value still saves it, but only on that slide, so it is not a
                # board-wide answer and the audit keeps reporting the board.
            ]
            if short:
                unfilled.append({"name": key, "used_in": where, "boards": short})
        results.append(
            {
                "template_id": template.id,
                "template": template.name,
                "ok": not unregistered and not unfilled,
                "variables_used": sorted(used),
                "unregistered": unregistered,
                "unfilled": unfilled,
            }
        )
    return {
        "ok": all(r["ok"] for r in results),
        "templates": len(results),
        "breaking": [r["template"] for r in results if not r["ok"]],
        "results": results,
        "boards_checked": [b.name for b in boards],
        "data_dir": str(settings.data_dir),
    }


# ---------------------------------------------------------------- the manager's own catalogue


def variable_conflicts(session: Session, template_id: int | None = None) -> list[JSONDict]:
    """Names that resolve to more than one thing, reported rather than silently resolved.

    Two cases can happen. A name registered **twice** — once shared, once bound to this template —
    is legal and useful (a template overriding a shared wording), but only one of them answers, so
    the manager says which. A registered name that collides with a **catalogue key** would make the
    field picker ambiguous; it cannot happen through the API today, and the check is here so it
    cannot start happening quietly either.
    """
    out: list[JSONDict] = []
    by_name: dict[str, list[s.VariableRow]] = {}
    for row in repo.list_variables(session, template_id):
        by_name.setdefault(row.name, []).append(row)
    for name, rows in sorted(by_name.items()):
        if len(rows) > 1:
            winner = sorted(rows, key=lambda r: (r.template_id is None, r.id))[0]
            out.append(
                {
                    "kind": "sombreada",
                    "name": name,
                    "message": tr(
                        "deck_service.conflict.shadowed",
                        name=name,
                        count=len(rows),
                        template_wins=winner.template_id is not None,
                    ),
                    "winner_template_id": winner.template_id,
                }
            )
        key = f"{FIXED_PREFIX}{name}"
        if any(f.key == key for f in FIELDS):
            out.append(
                {
                    "kind": "colide_com_campo",
                    "name": name,
                    "message": tr("deck_service.conflict.field", key=key),
                    "winner_template_id": None,
                }
            )
    return out


def field_catalog(session: Session, template_id: int | None = None) -> list[JSONDict]:
    """The Jira field catalogue as the variable manager browses it (RF-56).

    It is the same `FIELDS` the generator resolves against — name, kind, scope, label, group and
    hint all carried through untouched — plus the registered variables, so one screen answers
    "what can I put on a slide". The two stay *typed*: a catalogue field can be rows, groups, a
    series or a status, and collapsing them into text variables would throw away every table and
    chart binding in the deck.
    """
    return [
        {**entry, "origin": "variavel" if str(entry["key"]).startswith(FIXED_PREFIX) else "jira"}
        for entry in localized_catalog(catalog(catalog_extras(session, template_id)))
    ]


def localized_catalog(entries: Iterable[Mapping[str, Any]]) -> list[JSONDict]:
    """`context.catalog()` entries for a person, in the ambient language (D7).

    `FIELDS` keeps its pt-BR labels because the deck prints them (chart alt text); the picker shows
    ``fields.label.<key>`` instead. A variable keeps its registered label, which is user data —
    only its hint is translated. ``group_id`` is the stable group to sort and group by, now that
    ``group`` is a translated word."""
    out: list[JSONDict] = []
    for entry in entries:
        item = dict(entry)
        key = str(item.get("key", ""))
        if key in FIELD_INDEX:
            item["label"] = tr(f"fields.label.{key}")
        group_id = GROUP_IDS.get(str(item.get("group", "")))
        if group_id is not None:
            item["group"] = tr(f"fields.group.{group_id}")
        item["group_id"] = group_id
        out.append(localize(item))
    return out


def resolve_source_fields(
    session: Session,
    settings: Settings,
    *,
    template_id: int | None,
    board_id: int | None,
    now: datetime,
) -> dict[str, str]:
    """RF-58: what the variables wired to a Jira field say for one board.

    ``repo.resolve_variables`` cannot answer this — it has no board data — so the manager asks
    here, and asks through the very same `resolve` the generator uses. That is the point: the
    "resolvido" line under a connected variable has to be the value the deck will really carry,
    including the registration's default when the field is empty for this board.

    Returns only the names that are actually connected, and only when a board is chosen; a field
    without a board has nothing to say.
    """
    registry = variable_registry(session, template_id)
    connected = {name: spec for name, spec in registry.items() if spec.source_field}
    if not connected or board_id is None:
        return {}
    if session.get(s.BoardRow, board_id) is None:
        return {}
    deck = deck_source(session, settings).deck_data(
        [board_id],
        reference=default_reference(session, settings, [board_id], now),
        now=now,
        registry=registry,
        variables=variable_lookup(session, template_id),
    )
    board = deck.boards[0] if deck.boards else None
    out: dict[str, str] = {}
    with using_locale(DECK_LOCALE):  # what the deck will really carry, so in the deck's language
        for name in connected:
            _kind, value = resolve(deck, board, f"{FIXED_PREFIX}{name}")
            if value is not None and not isinstance(value, type(MISSING)):
                out[name] = str(value)
    return out


# ---------------------------------------------------------------- RF-59: field -> variable


def _variable_slug(key: str) -> str:
    """``riscos.itens`` -> ``riscos_itens``: the key, spelled the way a variable name must be."""
    return re.sub(r"[^a-z0-9]+", "_", key.lower()).strip("_")[:64]


def plan_binding_migration(session: Session, template_id: int | None = None) -> JSONDict:
    """What converting every legacy ``field:`` binding into a registered variable would do.

    RF-59 makes the manager the only place a name exists, which means the bindings that still
    address the catalogue directly have to become variables. This works it out and reports it; it
    writes nothing. The names are slugged from the key rather than the label, so it stays obvious
    which field each one came from, and they are registered **shared** — both of this project's
    templates bind many of the same fields, so one set of variables covers them instead of a
    near-duplicate per template.

    A field whose name is already taken by a *different* registration is reported as a conflict and
    left alone: silently reusing somebody's variable because the names collide is precisely the
    overwriting this module keeps refusing to do.
    """
    rows = list(session.scalars(select(s.TemplateRow).order_by(s.TemplateRow.id)))
    if template_id is not None:
        rows = [r for r in rows if r.id == template_id]
    registered = repo.registered_variables(session)
    by_source = {reg.source_field: name for name, reg in registered.items() if reg.source_field}

    uses: dict[str, list[JSONDict]] = {}
    for template in rows:
        mapping_row = repo.get_mapping(session, template.id)
        if mapping_row is None:
            continue
        mapping = DeckMapping.model_validate(mapping_row.mapping)
        for slide_id, spec in mapping.slides.items():
            for binding in spec.bindings:
                if binding.variable or binding.value is not None or not binding.field:
                    continue
                uses.setdefault(binding.field, []).append(
                    {
                        "template_id": template.id,
                        "template": template.name,
                        "slide_id": slide_id,
                        "target": binding.target,
                        "mode": binding.mode,
                    }
                )

    create: list[JSONDict] = []
    reuse: list[JSONDict] = []
    conflicts: list[JSONDict] = []
    for key, where in sorted(uses.items()):
        kind = source_kind(key)
        spec = next((f for f in FIELDS if f.key == key), None)
        entry = {"field": key, "kind": kind, "label": spec.label if spec else key, "used_by": where}
        if key in by_source:
            reuse.append({**entry, "name": by_source[key]})
            continue
        name = _variable_slug(key)
        existing = registered.get(name)
        if existing is not None:
            conflicts.append(
                {
                    **entry,
                    "name": name,
                    "reason": tr(
                        "deck_service.migration.name_taken", name=name, source=existing.source_field
                    ),
                }
            )
            continue
        create.append({**entry, "name": name})
    return {
        "templates": [{"id": r.id, "name": r.name} for r in rows],
        "bindings": sum(len(v) for v in uses.values()),
        "distinct_fields": len(uses),
        "create": create,
        "reuse": reuse,
        "conflicts": conflicts,
        "ok": not conflicts,
    }


def apply_binding_migration(session: Session, template_id: int | None = None) -> JSONDict:
    """Register the variables :func:`plan_binding_migration` planned, and repoint the bindings.

    Refuses to do anything if the plan reports a conflict: converting half of a template would
    leave a mapping that is neither one model nor the other. The rewrite itself is narrow — a
    binding loses ``field`` and gains ``variable``, and nothing else about it moves — so the
    generated deck is unchanged, which is what the golden snapshots are there to prove.
    """
    plan = plan_binding_migration(session, template_id)
    if not plan["ok"]:
        raise Invalid(Msg("deck_service.migration.conflicts"))

    created: list[JSONDict] = []
    for entry in plan["create"]:
        row = repo.create_variable(
            session,
            name=str(entry["name"]),
            label=str(entry["label"]),
            description=tr("deck_service.migration.description", field=str(entry["field"])),
            source_field=str(entry["field"]),
        )
        created.append({"id": row.id, "name": row.name, "source_field": row.source_field})

    # field key -> variable name, including the ones that already existed.
    names = {str(e["field"]): str(e["name"]) for e in (*plan["create"], *plan["reuse"])}
    touched: list[JSONDict] = []
    for template in plan["templates"]:
        mapping_row = repo.get_mapping(session, int(template["id"]))
        if mapping_row is None:
            continue
        # Raw, *not* `load_mapping`: this writes the mapping back, and saving a merged one would
        # bake the shared palette in as this template's override (RF-62).
        mapping = DeckMapping.model_validate(mapping_row.mapping)
        changed = 0
        for spec in mapping.slides.values():
            for binding in spec.bindings:
                if binding.variable or binding.value is not None or not binding.field:
                    continue
                name = names.get(binding.field)
                if name is None:
                    continue
                binding.variable, binding.field = name, None
                changed += 1
        if changed:
            repo.save_mapping(
                session, int(template["id"]), mapping_row.template_sha256, mapping.model_dump(mode="json")
            )
        touched.append({**template, "bindings_converted": changed})
    return {
        "created": created,
        "templates": touched,
        "converted": sum(t["bindings_converted"] for t in touched),
    }


# ---------------------------------------------------------------- RF-62: palette -> the manager


def plan_palette_migration(session: Session) -> JSONDict:
    """Which templates carry a status palette that is only a copy of the shared one.

    Before RF-62 every mapping was saved with a **full** palette, whether or not the template's
    legend actually said anything. Those entries now read as deliberate overrides, which means
    changing the colours centrally would do nothing for them. This finds the copies so they can be
    dropped and the template can inherit again — and leaves alone the entries that differ, because
    those came from the template's own legend and are a real choice.
    """
    shared_palette, shared_labels = repo.deck_palette(session)
    results: list[JSONDict] = []
    for template in session.scalars(select(s.TemplateRow).order_by(s.TemplateRow.id)):
        row = repo.get_mapping(session, template.id)
        if row is None:
            continue
        mapping = DeckMapping.model_validate(row.mapping)
        redundant = sorted(k for k, v in mapping.palette.items() if shared_palette.get(k) == v)
        kept = sorted(k for k, v in mapping.palette.items() if shared_palette.get(k) != v)
        redundant_labels = sorted(k for k, v in mapping.status_labels.items() if shared_labels.get(k) == v)
        if not redundant and not redundant_labels:
            continue
        results.append(
            {
                "template_id": template.id,
                "template": template.name,
                "inherit": redundant,
                "inherit_labels": redundant_labels,
                "keeps_own": {k: mapping.palette[k] for k in kept},
            }
        )
    return {"templates": results, "ok": not results}


def apply_palette_migration(session: Session) -> JSONDict:
    """Drop the redundant copies so those templates inherit the shared configuration again."""
    plan = plan_palette_migration(session)
    for entry in plan["templates"]:
        row = repo.get_mapping(session, int(entry["template_id"]))
        if row is None:
            continue
        mapping = DeckMapping.model_validate(row.mapping)
        for key in entry["inherit"]:
            mapping.palette.pop(key, None)
        for key in entry["inherit_labels"]:
            mapping.status_labels.pop(key, None)
        repo.save_mapping(
            session, int(entry["template_id"]), row.template_sha256, mapping.model_dump(mode="json")
        )
    return plan


def template_readiness(session: Session, settings: Settings, template_id: int, now: datetime) -> JSONDict:
    """What would be wrong with this template's deck, **before** generating it (RF-66).

    `audit_variables` already answers "which names have nowhere to get a value"; this adds where on
    the slides that bites, so the editor can mark the affected thumbnails instead of handing over a
    list of names. Read-only and fast: it resolves variables, it does not build a .pptx.
    """
    template = session.get(s.TemplateRow, template_id)
    if template is None:
        raise NotFound(template_id, Msg("deck_service.template.not_found_id", template=template_id))
    mapping = load_mapping(session, template_id)
    registry = variable_registry(session, template_id)
    lookup = variable_lookup(session, template_id)
    boards = [
        b
        for b in repo.selected_boards(session)
        if not mapping.space_keys or b.space_key in set(mapping.space_keys)
    ]

    slides: list[JSONDict] = []
    for slide in template.manifest.get("slides", []):
        origin = int(slide["slide_id"])
        spec = mapping.for_slide(origin)
        if spec.action == "remove":
            continue
        used = sorted(
            {
                key
                for key in (
                    *(b.key for b in spec.bindings if b.key and b.value is None),
                    *(marker_base(m.group(1)) for b in spec.bindings for m in MARKER.finditer(b.value or "")),
                    *(
                        marker_base(name)
                        for area in slide.get("areas", [])
                        for name in area.get("markers", [])
                    ),
                )
                if variable_name(key) is not None
            }
        )
        problems: list[JSONDict] = []
        for key in used:
            name = variable_name(key) or ""
            spec_for_name = registry.get(name)
            if spec_for_name is None:
                problems.append({"kind": "nao_registrada", "variable": key, "boards": []})
                continue
            if spec_for_name.source_field or spec_for_name.default is not None:
                continue
            # A hand-filled variable with no default: which of the boards in scope is short of it.
            short = [
                b.name
                for b in boards
                if spec.covers(b.space_key, b.id) and name not in lookup(b.space_key, b.id)
            ]
            if short:
                problems.append({"kind": "sem_valor", "variable": key, "boards": short})
        unbound = [b.target for b in spec.bindings if b.value is None and not b.key]
        slides.append(
            {
                "slide_id": origin,
                "index": slide.get("index"),
                "name": spec.name or slide.get("headline") or "",
                "variables": used,
                "problems": problems,
                "unbound_targets": unbound,
                "ok": not problems and not unbound,
            }
        )
    return {
        "template_id": template_id,
        "reference": now.isoformat(),
        "boards": [b.name for b in boards],
        "slides": slides,
        "ok": all(entry["ok"] for entry in slides),
    }


# ---------------------------------------------------------------- find & replace (RF-88)


def replace_in_variables(
    session: Session,
    find: str,
    replace: str,
    *,
    match_case: bool = False,
    template_id: int | None = None,
    apply: bool = False,
) -> JSONDict:
    """Find ``find`` in every variable value and registration default, and optionally replace it.

    Without ``apply`` nothing is written: the list of matches is the preview the manager shows
    before anybody commits to renaming a product across thirty values.
    """
    if not find:
        raise Invalid(Msg("deck_service.find.required"))
    pattern = re.compile(re.escape(find), 0 if match_case else re.IGNORECASE)

    def swap(text: str) -> str:
        return pattern.sub(lambda _match: replace, text)

    matches: list[JSONDict] = []
    for variable in repo.list_variables(session, template_id):
        if variable.default and pattern.search(variable.default):
            matches.append(
                {
                    "variable": variable.name,
                    "scope": "registro",
                    "scope_key": "",
                    "before": variable.default,
                    "after": swap(variable.default),
                }
            )
            if apply:
                variable.default = swap(variable.default)
        for value in repo.list_variable_values(session, variable.id):
            if value.value and pattern.search(value.value):
                matches.append(
                    {
                        "variable": variable.name,
                        "scope": value.scope,
                        "scope_key": value.scope_key,
                        "before": value.value,
                        "after": swap(value.value),
                    }
                )
                if apply:
                    value.value = swap(value.value)
    if apply:
        session.flush()
    return {"matches": matches, "count": len(matches), "applied": apply}
