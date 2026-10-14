import asyncio
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from cadencia.api import routes_setup
from cadencia.api.app import create_app
from cadencia.api.deps import clock
from cadencia.i18n import Msg
from cadencia.jira.errors import JiraAuthError
from cadencia.report.core.sample_template import build_sample_template
from tests.support import NOW, PT_BR, collect_demo


@pytest.fixture
def client(demo, sessions, settings):
    asyncio.run(collect_demo(demo, sessions, settings))
    app = create_app(settings, sessions)
    app.dependency_overrides[clock] = lambda: NOW
    with TestClient(app, headers=PT_BR) as test_client:
        yield test_client


def ok(response):
    assert response.status_code < 300, response.text
    return response.json()


def test_health_and_route_inventory(client):
    health = ok(client.get("/api/health"))
    assert health["status"] == "ok" and health["jira_configured"]
    assert len(ok(client.get("/openapi.json"))["paths"]) >= 35


def test_dashboard_has_one_row_per_board_naming_its_space(client):
    rows = {r["board"]["name"]: r for r in ok(client.get("/api/dashboard"))["rows"]}
    assert set(rows) == {"Plataforma", "Dados", "Suporte"}
    plataforma = rows["Plataforma"]
    assert plataforma["sprint"]["name"] == "Plataforma 3.5"
    assert plataforma["window"]["index"] == 2 and plataforma["window"]["total"] == 2
    assert plataforma["risk_counts"]["arrastado"] == 1
    assert [t["label"] for t in plataforma["trend"]] == ["Plataforma 3.3", "Plataforma 3.4", "Plataforma 3.5"]
    assert plataforma["freshness"]["last_success_at"]
    # kanban: four closed weekly windows plus the open one
    assert rows["Suporte"]["sprint"] is None and len(rows["Suporte"]["trend"]) == 5
    # The team is the space above the board, and two boards share one.
    assert plataforma["board"]["space_name"] == "Produto Digital"
    assert rows["Suporte"]["board"]["space_name"] == "Produto Digital"
    assert rows["Dados"]["board"]["space_name"] == "Inteligência de Dados"


def test_sprint_overview_windows_and_commitment(client):
    overview = ok(client.get("/api/sprints/12"))
    first, second = overview["timeline"]["windows"]
    assert first["state"] == "fechada" and second["state"] == "aberta"
    assert first["diff"]["added"] == ["CAD-15"]
    assert second["diff"]["removed"] == ["CAD-16"]
    assert "CAD-10" in second["diff"]["completed"]
    sets = overview["commitment"]["sets"]
    assert "CAD-15" in sets["added"] and "CAD-16" in sets["removed"]
    assert overview["progress"]["total"] == len(sets["committed"]) - len(sets["removed"]) + len(sets["added"])


def test_sprint_overview_includes_each_item_s_own_epic(client):
    """RF-112: the window lists carry only standard items, but the frontend groups them under
    their epic — which needs that epic's own brief, not just its key."""
    overview = ok(client.get("/api/sprints/12"))
    second = overview["timeline"]["windows"][1]
    assert "CAD-10" in second["diff"]["completed"]
    assert overview["items"]["CAD-10"]["parent_key"] == "CAD-1"
    assert overview["items"]["CAD-1"]["summary"] == "Checkout novo"


def test_risks_are_classified_with_evidence(client):
    items = {r["key"]: r for r in ok(client.get("/api/sprints/12/risks"))["items"]}
    situation = {k: v["risk"]["situation"] for k, v in items.items()}
    assert situation["CAD-11"] == "arrastado"
    assert situation["CAD-12"] == "envelhecendo"
    assert situation["CAD-13"] == "bloqueado"
    assert situation["CAD-14"] == "atrasado"
    assert situation["CAD-18"] == "bloqueado"
    assert situation["CAD-17"] == "em_dia"
    carried = items["CAD-11"]["risk"]["evidence"]
    assert [e["rule"] for e in carried][:1] == ["arrasto"]
    stagnation = next(e for e in carried if e["rule"] == "estagnacao")
    assert any("Bruno Lima para Carla Mendes" in d for d in stagnation["details"])
    assert stagnation["last_transition"]["author"] == "Bruno Lima"
    assert "aguardando CAD-11" in items["CAD-18"]["risk"]["evidence"][0]["message"]
    only_blocked = ok(client.get("/api/sprints/12/risks", params={"situation": "bloqueado"}))["items"]
    assert {r["key"] for r in only_blocked} == {"CAD-13", "CAD-18"}


def test_tree_epic_and_item_drill_down(client):
    nodes = {n["key"]: n for n in ok(client.get("/api/sprints/12/tree"))["nodes"]}
    story = next(c for c in nodes["CAD-1"]["children"] if c["key"] == "CAD-10")
    assert [c["key"] for c in story["children"]] == ["CAD-19", "CAD-20"]
    assert nodes["__sem_epico__"]["label"] == "Sem épico"
    epic = ok(client.get("/api/epics/CAD-1"))
    assert epic["tree"]["progress"]["total"] > len(
        nodes["CAD-1"]["children"]
    )  # all sprints, not just this one
    detail = ok(client.get("/api/issues/CAD-22"))
    assert detail["metrics"]["reopened"] and detail["metrics"]["backward_moves"] == 1
    assert detail["issue"]["url"] == "https://demo.atlassian.net/browse/CAD-22"
    assert isinstance(detail["description"], str)
    assert isinstance(epic["description"], str)
    # The epic page lives at /epics/<key>: the view names its board so the page can link back to it.
    assert epic["board"] == {"id": 1, "name": "Plataforma"}


def test_a_space_lists_its_boards_and_its_risks_across_them(client):
    space = ok(client.get("/api/spaces/CAD"))
    assert space["space"] == {"key": "CAD", "name": "Produto Digital"}
    assert [b["name"] for b in space["boards"]] == ["Plataforma", "Suporte"]

    risks = ok(client.get("/api/spaces/CAD/risks"))
    keys = [i["key"] for i in risks["items"]]
    assert len(keys) == len(set(keys)) and "CAD-14" in keys
    assert {i["board"]["name"] for i in risks["items"]} <= {"Plataforma", "Suporte"}
    severities = [i["risk"]["severity"] for i in risks["items"]]
    assert severities == sorted(severities, reverse=True)
    assert risks["freshness"] is not None
    only = ok(client.get("/api/spaces/CAD/risks", params={"situation": "atrasado"}))
    assert only["items"] and {i["risk"]["situation"] for i in only["items"]} == {"atrasado"}

    assert client.get("/api/spaces/NOPE").status_code == 404
    assert client.get("/api/spaces/NOPE/risks").status_code == 404


def test_a_kanban_board_gets_the_same_shape_of_tree_a_sprint_does(client):
    """RF-113: `board_tree` walks the whole board instead of one sprint's membership, but the
    grouping (by epic, "Sem épico" for the rest) is the same `build_tree` a sprint gets."""
    nodes = {n["key"]: n for n in ok(client.get("/api/boards/3/tree"))["nodes"]}
    assert "__sem_epico__" in nodes
    # SUP-7's epic (CAD-2) is on a different board entirely — the cross-project fetch (and the
    # initiative above it) has to have made it into this kanban board's own tree too.
    cad2 = next(
        n for n in nodes.values() if n["key"] == "CAD-2" or any(c["key"] == "CAD-2" for c in n["children"])
    )
    epic_node = cad2 if cad2["key"] == "CAD-2" else next(c for c in cad2["children"] if c["key"] == "CAD-2")
    assert any(c["key"] == "SUP-7" for c in epic_node["children"])


def test_a_kanban_board_s_overview_uses_calendar_weeks_not_a_sprint(client):
    """RF-113: `board_overview` has no sprint to report progress/commitment against — it carries
    only the weekly windows (calendar-based), risk counts, and the items needed to group by epic."""
    overview = ok(client.get("/api/boards/3/overview"))
    assert overview["board"]["id"] == 3
    assert "sprint" not in overview and "commitment" not in overview and "progress" not in overview
    assert len(overview["timeline"]["windows"]) >= 2
    assert overview["timeline"]["windows"][-1]["state"] == "aberta"
    assert all(w["state"] == "fechada" for w in overview["timeline"]["windows"][:-1])


def test_flow_views(client):
    view = ok(client.get("/api/flow"))
    assert view["summary"]["completed"] > 0
    names = [p["name"] for p in view["people"]]
    assert names == sorted(names, key=str.casefold)  # alphabetical, never ranked
    assert view["summary"]["cycle_time"]["n"] == view["summary"]["completed"]
    cfd = ok(client.get("/api/boards/1/cfd"))
    assert [c["name"] for c in cfd["columns"]] == ["A fazer", "Em andamento", "Code Review", "Concluído"]
    assert ok(client.get("/api/boards/1/cycle-scatter"))["points"]
    aging = ok(client.get("/api/boards/1/aging"))
    assert "CAD-12" in {i["key"] for i in aging["items"]}
    assert aging["columns"] == ["Em andamento", "Code Review"]
    cvd = ok(client.get("/api/boards/1/committed-vs-delivered"))["items"]
    assert [r["sprint"]["name"] for r in cvd] == ["Plataforma 3.3", "Plataforma 3.4", "Plataforma 3.5"]
    person = ok(client.get("/api/people/demo-ana"))
    assert person["person"]["name"] == "Ana Souza" and "mediana" in person["note"]


def test_planning_allocation_and_simulation(client):
    capacity = ok(client.get("/api/sprints/12/capacity"))
    assert capacity["items"] and capacity["focus_factor"] == 0.7
    assert ok(client.get("/api/sprints/12/allocation"))["plan"] is None
    created = ok(client.post("/api/sprints/12/allocation"))
    committed = ok(client.get("/api/sprints/12"))["commitment"]["sets"]["committed"]
    assert len(created["plan"]["assignments"]) == len(committed)
    assert "heurística" in created["disclaimer"]
    stored = ok(client.get("/api/sprints/12/allocation"))
    assert stored["plan"]["assignments"] == created["plan"]["assignments"]
    simulated = ok(client.post("/api/sprints/12/allocation/simulate", json={"focus_factor": 0.3}))
    assert simulated["simulated"] and simulated["plan"]["parameters"]["focus_factor"] == 0.3
    assert ok(client.get("/api/sprints/12/allocation"))["id"] == stored["id"]  # simulation not persisted


def test_exports(client):
    response = client.get("/api/export/risks", params={"sprint_id": 12})
    assert response.status_code == 200 and response.headers["content-type"].startswith("text/csv")
    assert response.text.startswith("﻿") and ";" in response.text.splitlines()[0]
    rows = ok(client.get("/api/export/flow", params={"format": "json"}))
    assert isinstance(rows, list) and rows


def test_board_selection_and_settings(client):
    listing = ok(client.get("/api/boards", params={"q": "dad"}))
    assert [b["name"] for b in listing["items"]] == ["Dados"]
    ok(client.put("/api/boards/1/settings", json={"stagnation_checkpoints": 2}))
    assert ok(client.get("/api/boards/1"))["settings"]["stagnation_checkpoints"] == 2
    assert client.put("/api/boards/1/settings", json={"nope": 1}).status_code == 400
    assert [b["name"] for b in ok(client.get("/api/boards", params={"q": "produto"}))["items"]] == [
        "Plataforma",
        "Suporte",
    ]
    status = ok(client.get("/api/status"))
    assert {b["board"]["id"] for b in status["boards"]} == {1, 2, 3}


def test_boards_are_grouped_under_the_space_that_owns_them(client):
    listing = ok(client.get("/api/boards"))
    groups = {g["space"]["name"]: [b["name"] for b in g["boards"]] for g in listing["spaces"]}
    assert groups == {
        "Inteligência de Dados": ["Dados"],
        "Produto Digital": ["Plataforma", "Suporte"],
    }
    spaces = ok(client.get("/api/spaces"))
    assert [g["space"]["key"] for g in spaces["items"]] == ["DAD", "CAD"]
    assert spaces["total"] == 2


def test_status_rows_say_whether_they_are_a_space_or_a_board(client):
    entities = ok(client.get("/api/status"))["entities"]
    assert [(e["entity"], e["name"]) for e in entities] == [
        ("space", "Inteligência de Dados"),
        ("board", "Dados"),
        ("space", "Produto Digital"),
        ("board", "Plataforma"),
        ("board", "Suporte"),
    ]
    produto = next(e for e in entities if e["name"] == "Produto Digital")
    boards = [e for e in entities if e["entity"] == "board" and e["space"]["key"] == "CAD"]
    assert produto["boards"] == 2
    # The space rolls its boards up by the oldest successful collection, and keeps the old columns.
    assert produto["last_success_at"] == min(b["last_success_at"] for b in boards)
    assert produto["last_full_at"] and produto["last_run"]["status"] == "success"


def test_people_are_grouped_by_space_and_may_appear_in_more_than_one(client):
    people = ok(client.get("/api/people"))
    by_space = {g["space"]["name"]: {p["name"] for p in g["people"]} for g in people["spaces"]}
    assert set(by_space) == {"Inteligência de Dados", "Produto Digital"}
    shared = set.intersection(*by_space.values())
    assert shared, "a demo tem gente trabalhando em mais de um espaço"
    everyone = {p["name"] for p in people["items"]}
    assert set().union(*by_space.values()) | {p["name"] for p in people["unassigned"]} == everyone
    for group in people["spaces"]:  # alphabetical, never ranked
        names = [p["name"] for p in group["people"]]
        assert names == sorted(names, key=str.casefold)


def test_connection_test_reports_the_cause(client, monkeypatch):
    async def failing(*_args, **_kwargs):
        raise JiraAuthError("HTTP 401")

    monkeypatch.setattr(routes_setup, "run_connection_test", failing)
    body = ok(client.post("/api/connection/test"))
    assert body["ok"] is False and body["error"]["cause"] == "credencial inválida"


def test_report_template_mapping_generation_and_download(client, tmp_path):
    path = build_sample_template(tmp_path / "modelo.pptx")
    with open(path, "rb") as fh:
        uploaded = ok(client.post("/api/reports/templates", files={"file": ("modelo.pptx", fh)}))
    detail = ok(client.get(f"/api/reports/templates/{uploaded['id']}"))
    suggestion = detail["suggestion"]
    for spec in suggestion["slides"].values():
        if spec["scope"] == "board" and not spec.get("board_id"):
            spec["scope"] = "repeat"
    suggestion.pop("notes")
    ok(client.put(f"/api/reports/templates/{uploaded['id']}/mapping", json=suggestion))
    run = ok(client.post("/api/reports/generate", json={"template_id": uploaded["id"], "render": False}))
    assert run["status"] == "ok", run["validation"]["errors"]
    assert run["validation"]["slides"] >= 8
    download = client.get(f"/api/reports/runs/{run['id']}/download")
    assert download.status_code == 200 and download.content[:2] == b"PK"
    # A generation has its own page, so it can be read alone.
    alone = ok(client.get(f"/api/reports/runs/{run['id']}"))
    assert alone["id"] == run["id"] and alone["template_id"] == uploaded["id"]
    assert client.get("/api/reports/runs/nope").status_code == 404
    bad = client.put(
        f"/api/reports/templates/{uploaded['id']}/mapping",
        json={"slides": {"1": {"bindings": [{"target": "s1/sh999", "mode": "text", "field": "board.nome"}]}}},
    )
    assert bad.status_code == 400


# ---------------------------------------------------------------- presentation (RF-38, RF-45)


def define_project_variables(client, template_id=None):
    """One variable name, a different value per project, plus a global fallback."""
    definicao = ok(
        client.post(
            "/api/reports/variables",
            json={"name": "definicao", "label": "Definição do projeto", "template_id": template_id},
        )
    )
    for key, value in (("CAD", "Checkout e pagamentos."), ("DAD", "Ingestão e data warehouse.")):
        ok(
            client.put(
                f"/api/reports/variables/{definicao['id']}/values",
                json={"scope": "space", "scope_key": key, "value": value},
            )
        )
    aviso = ok(client.post("/api/reports/variables", json={"name": "aviso"}))
    ok(
        client.put(
            f"/api/reports/variables/{aviso['id']}/values",
            json={"scope": "global", "value": "Interno."},
        )
    )
    return definicao, aviso


def test_variables_are_managed_on_their_own_and_resolve_per_project(client):
    definicao, aviso = define_project_variables(client)
    listing = ok(client.get("/api/reports/variables"))
    assert [v["name"] for v in listing["items"]] == ["aviso", "definicao"]
    assert listing["items"][1]["placeholder"] == "{{texto.definicao}}"
    assert listing["scopes"] == ["global", "space", "board", "slide"]
    assert {v["scope"] for v in listing["items"][1]["values"]} == {"space"}
    # With no target chosen the answer is the global default, not an empty map.
    assert listing["resolved"] == {"aviso": "Interno."}

    for space, expected in (("CAD", "Checkout e pagamentos."), ("DAD", "Ingestão e data warehouse.")):
        resolved = ok(client.get("/api/reports/variables", params={"space_key": space}))["resolved"]
        assert resolved["definicao"] == expected
        assert resolved["aviso"] == "Interno."  # the global default reaches both

    # The picker offers them like any other field...
    catalog = {f["key"]: f for f in ok(client.get("/api/reports/catalog"))["fields"]}
    assert catalog["texto.definicao"]["scope"] == "board"
    assert catalog["texto.definicao"]["label"] == "Definição do projeto"
    # ...and the preview shows the resolved text, never the raw placeholder.
    for board_id, expected in ((1, "Checkout e pagamentos."), (2, "Ingestão e data warehouse.")):
        fields = ok(client.get("/api/reports/preview", params={"board_id": board_id}))["fields"]
        assert fields["texto.definicao"]["lines"] == [expected]

    assert client.post("/api/reports/variables", json={"name": "com espaço"}).status_code == 400

    # A default belongs to the registration, and "no default" is not "a default of empty".
    fallback = ok(client.post("/api/reports/variables", json={"name": "rodape", "default": "Uso interno."}))
    assert fallback["default"] == "Uso interno."
    assert ok(client.get("/api/reports/variables"))["resolved"]["rodape"] == "Uso interno."
    emptied = ok(client.patch(f"/api/reports/variables/{fallback['id']}", json={"default": ""}))
    assert emptied["default"] == ""
    assert ok(client.get("/api/reports/variables"))["resolved"]["rodape"] == ""
    cleared = ok(client.patch(f"/api/reports/variables/{fallback['id']}", json={"clear_default": True}))
    assert cleared["default"] is None
    # No value anywhere and no default: the name is absent, which is what generation refuses.
    assert "rodape" not in ok(client.get("/api/reports/variables"))["resolved"]
    assert client.delete(f"/api/reports/variables/{fallback['id']}").status_code == 204

    assert client.post("/api/reports/variables", json={"name": "definicao"}).status_code == 400
    renamed = ok(client.patch(f"/api/reports/variables/{definicao['id']}", json={"name": "descricao"}))
    assert renamed["placeholder"] == "{{texto.descricao}}"
    assert client.delete(f"/api/reports/variables/{aviso['id']}").status_code == 204
    assert [v["name"] for v in ok(client.get("/api/reports/variables"))["items"]] == ["descricao"]


def test_a_variable_can_be_assigned_to_one_template(client, tmp_path):
    """RF-52: detaching variables from the project means binding them to a template."""
    first = build_sample_template(tmp_path / "um.pptx")
    with open(first, "rb") as fh:
        a = ok(client.post("/api/reports/templates", files={"file": ("um.pptx", fh)}, data={"name": "A"}))
    shared, _ = define_project_variables(client)
    only_a = ok(
        client.post(
            "/api/reports/variables",
            json={"name": "rodape_do_a", "label": "Rodapé do A", "template_id": a["id"]},
        )
    )
    ok(
        client.put(
            f"/api/reports/variables/{only_a['id']}/values",
            json={"scope": "global", "value": "Só neste modelo."},
        )
    )
    mine = ok(client.get("/api/reports/variables", params={"template_id": a["id"]}))
    assert {v["name"] for v in mine["items"]} == {"aviso", "definicao", "rodape_do_a"}
    # Another template does not see it, which is the whole point of the binding.
    others = ok(client.get("/api/reports/variables", params={"template_id": a["id"] + 999}))
    assert "rodape_do_a" not in {v["name"] for v in others["items"]}
    catalog = {
        f["key"] for f in ok(client.get("/api/reports/catalog", params={"template_id": a["id"]}))["fields"]
    }
    assert "texto.rodape_do_a" in catalog
    plain = {
        f["key"]
        for f in ok(client.get("/api/reports/catalog", params={"template_id": a["id"] + 999}))["fields"]
    }
    assert "texto.rodape_do_a" not in plain

    ok(client.patch(f"/api/reports/variables/{only_a['id']}", json={"detach": True}))
    everywhere = ok(client.get("/api/reports/variables", params={"template_id": a["id"] + 999}))
    assert "rodape_do_a" in {v["name"] for v in everywhere["items"]}
    assert shared["template_id"] is None


def test_a_variable_set_on_one_slide_wins_over_the_board_and_the_project(client, tmp_path):
    """The slide level: the narrowest target the user can pick, so it answers first."""
    source = build_sample_template(tmp_path / "modelo.pptx")
    with open(source, "rb") as fh:
        uploaded = ok(client.post("/api/reports/templates", files={"file": ("modelo.pptx", fh)}))
    detail = ok(client.get(f"/api/reports/templates/{uploaded['id']}"))
    slide = detail["manifest"]["slides"][1]
    definicao, _ = define_project_variables(client, template_id=uploaded["id"])
    ok(
        client.put(
            f"/api/reports/variables/{definicao['id']}/values",
            json={"scope": "slide", "scope_key": str(slide["slide_id"]), "value": "Só neste slide."},
        )
    )
    resolved = ok(
        client.get(
            "/api/reports/variables",
            params={
                "template_id": uploaded["id"],
                "space_key": "CAD",
                "board_id": 1,
                "slide_id": slide["slide_id"],
            },
        )
    )["resolved"]
    assert resolved["definicao"] == "Só neste slide."
    # Any other slide still gets the project's wording.
    elsewhere = ok(
        client.get(
            "/api/reports/variables",
            params={"template_id": uploaded["id"], "space_key": "CAD", "board_id": 1},
        )
    )["resolved"]
    assert elsewhere["definicao"] == "Checkout e pagamentos."

    mapping = detail["suggestion"]
    mapping.pop("notes", None)
    spec = mapping["slides"][str(slide["slide_id"])]
    spec["scope"] = "board"
    spec["board_id"] = 1
    target = next(t["id"] for a in slide["areas"] if a.get("paragraphs") for t in a["targets"])
    spec["bindings"].append({"target": target, "mode": "text", "value": "{{texto.definicao}}"})
    ok(client.put(f"/api/reports/templates/{uploaded['id']}/mapping", json=mapping))
    view = ok(
        client.post(
            f"/api/reports/templates/{uploaded['id']}/slides/{slide['slide_id']}/preview",
            json={"mapping": mapping, "board_id": 1},
        )
    )
    literal = next(i for i in view["items"] if i.get("literal"))
    assert literal["text"] == "Só neste slide."


def test_slide_preview_resolves_placeholders_and_follows_the_chosen_project(client, tmp_path):
    define_project_variables(client)
    template = build_sample_template(tmp_path / "modelo.pptx")
    with open(template, "rb") as fh:
        uploaded = ok(client.post("/api/reports/templates", files={"file": ("modelo.pptx", fh)}))
    detail = ok(client.get(f"/api/reports/templates/{uploaded['id']}"))
    slide = detail["manifest"]["slides"][1]
    mapping = detail["suggestion"]
    mapping.pop("notes", None)
    spec = mapping["slides"][str(slide["slide_id"])]
    spec["scope"] = "repeat"
    target = next(t["id"] for a in slide["areas"] if a.get("paragraphs") for t in a["targets"])
    spec["bindings"].append(
        {"target": target, "mode": "text", "value": "Sobre: {{texto.definicao}} ({{texto.aviso}})"}
    )

    seen = {}
    for board_id in (1, 2):
        body = {"mapping": mapping, "board_id": board_id}
        view = ok(
            client.post(
                f"/api/reports/templates/{uploaded['id']}/slides/{slide['slide_id']}/preview", json=body
            )
        )
        literal = next(i for i in view["items"] if i.get("literal"))
        seen[view["board"]["space_name"]] = literal["text"]
        assert "{{" not in literal["text"]  # the preview shows the resolved text, not the marker
        assert literal["raw"] == "Sobre: {{texto.definicao}} ({{texto.aviso}})"
    # Switching the project repopulates the placeholder with that project's words.
    assert seen == {
        "Produto Digital": "Sobre: Checkout e pagamentos. (Interno.)",
        "Inteligência de Dados": "Sobre: Ingestão e data warehouse. (Interno.)",
    }


def test_slide_preview_carries_the_resolved_content_of_every_mode(client, tmp_path):
    template = build_sample_template(tmp_path / "modelo.pptx")
    with open(template, "rb") as fh:
        uploaded = ok(client.post("/api/reports/templates", files={"file": ("modelo.pptx", fh)}))
    detail = ok(client.get(f"/api/reports/templates/{uploaded['id']}"))
    slide = detail["manifest"]["slides"][1]
    mapping = detail["suggestion"]
    mapping.pop("notes", None)
    mapping["slides"][str(slide["slide_id"])]["scope"] = "repeat"
    view = ok(
        client.post(
            f"/api/reports/templates/{uploaded['id']}/slides/{slide['slide_id']}/preview",
            json={"mapping": mapping, "board_id": 1},
        )
    )
    by_mode = {}
    for item in view["items"]:
        by_mode.setdefault(item["mode"], []).append(item)
    assert "text" in by_mode
    # A grouped list comes back as real groups with real items, not as a count.
    grouped = [i for i in by_mode.get("grouped_list", []) if not i["empty"]]
    assert grouped and grouped[0]["groups"][0]["items"]
    fills = [i for i in by_mode.get("status_fill", []) if not i["empty"]]
    assert fills and len(fills[0]["fill"]) == 6  # a hex colour ready to paint
    assert view["board"]["name"] == "Plataforma"
    assert view["scope"] == "repeat"


def test_template_refresh_upgrades_an_older_manifest(client, tmp_path):
    template = build_sample_template(tmp_path / "modelo.pptx")
    with open(template, "rb") as fh:
        uploaded = ok(client.post("/api/reports/templates", files={"file": ("modelo.pptx", fh)}))
    result = ok(client.post(f"/api/reports/templates/{uploaded['id']}/refresh"))
    assert result["version"] == 2
    manifest = ok(client.get(f"/api/reports/templates/{uploaded['id']}"))["manifest"]
    area = next(a for s in manifest["slides"] for a in s["areas"] if a.get("paragraphs"))
    assert "size_pt" in area["paragraphs"][0] and "anchor" in area["frame"]


def generated_run(client, tmp_path):
    template = build_sample_template(tmp_path / "modelo.pptx")
    with open(template, "rb") as fh:
        uploaded = ok(client.post("/api/reports/templates", files={"file": ("modelo.pptx", fh)}))
    detail = ok(client.get(f"/api/reports/templates/{uploaded['id']}"))
    mapping = detail["suggestion"]
    mapping.pop("notes", None)
    for spec in mapping["slides"].values():
        if spec["scope"] == "repeat":
            spec["scope"] = "board"
            spec["board_id"] = 1
    ok(client.put(f"/api/reports/templates/{uploaded['id']}/mapping", json=mapping))
    run = ok(client.post("/api/reports/generate", json={"template_id": uploaded["id"], "render": False}))
    return uploaded["id"], run


def test_a_generated_deck_opens_as_an_editable_scene(client, tmp_path):
    template_id, run = generated_run(client, tmp_path)
    scene = ok(client.get(f"/api/reports/runs/{run['id']}/scene"))
    assert scene["template_id"] == template_id and scene["slide_width"] > 0
    assert len(scene["slides"]) == run["validation"]["slides"]
    shapes = [s for slide in scene["slides"] for s in slide["shapes"]]
    # Geometry comes straight out of the file, in EMU, never rounded to pixels.
    assert all(isinstance(s["frame"]["x"], int) for s in shapes)
    # Data-bound shapes are marked, so the editor can show what a re-run will overwrite.
    assert any(s["bound"] and s["bound_field"] for s in shapes)
    text = next(s for s in shapes if s.get("text"))
    assert "paragraphs" in text["text"] and "insets_pt" in text["text"]


def test_edits_are_saved_and_replayed_on_the_next_generation(client, tmp_path):
    """The acceptance criterion: refresh the data, keep the manual work."""
    template_id, run = generated_run(client, tmp_path)
    scene = ok(client.get(f"/api/reports/runs/{run['id']}/scene"))
    slide = scene["slides"][1]
    box = next(s for s in slide["shapes"] if s.get("text"))
    edits = {
        "version": 1,
        "slides": {
            slide["key"]: {
                "shapes": {
                    box["key"]: {
                        "op": "update",
                        "frame": {"x": 1500000, "y": 900000, "cx": 3000000, "cy": 800000, "rotation": 7.5},
                    },
                    "new:nota": {
                        "op": "insert",
                        "kind": "textbox",
                        "frame": {"x": 600000, "y": 5200000, "cx": 5000000, "cy": 700000},
                        "text": {"paragraphs": [{"text": "Revisado à mão", "style": {"size_pt": 16}}]},
                    },
                }
            }
        },
        "order": [],
        "deleted": [],
        "duplicates": {},
    }
    saved = ok(
        client.put(f"/api/reports/templates/{template_id}/edits", json={"edits": edits, "base_revision": 0})
    )
    assert saved["revision"] == 1
    # Saving against a stale revision is refused rather than silently overwriting.
    stale = client.put(
        f"/api/reports/templates/{template_id}/edits", json={"edits": edits, "base_revision": 0}
    )
    assert stale.status_code == 409

    again = ok(client.post("/api/reports/generate", json={"template_id": template_id, "render": False}))
    assert again["status"] == "ok"
    after = ok(client.get(f"/api/reports/runs/{again['id']}/scene"))
    edited = next(s for s in after["slides"] if s["key"] == slide["key"])
    moved = next(s for s in edited["shapes"] if s["key"] == box["key"])
    assert moved["frame"] == {"x": 1500000, "y": 900000, "cx": 3000000, "cy": 800000, "rotation": 7.5}
    inserted = next(s for s in edited["shapes"] if s["key"] == "new:nota")
    assert inserted["text"]["paragraphs"][0]["text"] == "Revisado à mão"
    assert edited["edited"] is True

    # Clearing the layer goes back to what the data alone produces.
    assert client.delete(f"/api/reports/templates/{template_id}/edits").status_code == 204
    plain = ok(client.post("/api/reports/generate", json={"template_id": template_id, "render": False}))
    scene_after = ok(client.get(f"/api/reports/runs/{plain['id']}/scene"))
    back = next(s for s in scene_after["slides"] if s["key"] == slide["key"])
    assert not any(s["key"] == "new:nota" for s in back["shapes"])
    assert next(s for s in back["shapes"] if s["key"] == box["key"])["frame"] == box["frame"]


def test_the_editor_serves_metric_compatible_fonts(client):
    plan = ok(client.get("/api/reports/fonts"))
    assert "families" in plan and "faces" in plan
    by_family = {f["family"].lower(): f for f in plan["families"]}
    # Whatever the machine has, the contract is the same: a face or an honest "approximate".
    for entry in plan["families"]:
        assert entry["metric_compatible"] == (entry["face"] is not None)
    if "calibri" in by_family and by_family["calibri"]["face"]:
        face = by_family["calibri"]["face"]
        assert client.get(f"/api/reports/fonts/{face}/400").status_code == 200
    assert client.get("/api/reports/fonts/NaoExiste/400").status_code == 404


def test_a_malformed_package_fails_the_generation_instead_of_reaching_the_user(client, tmp_path, monkeypatch):
    """RF-50: the package check runs on every generation, and a failure invalidates the run."""
    template = build_sample_template(tmp_path / "modelo.pptx")
    with open(template, "rb") as fh:
        uploaded = ok(client.post("/api/reports/templates", files={"file": ("modelo.pptx", fh)}))

    healthy = ok(client.post("/api/reports/generate", json={"template_id": uploaded["id"], "render": False}))
    assert healthy["status"] == "ok"
    assert healthy["validation"]["package"]["ok"] is True
    assert healthy["validation"]["package"]["slides"] == healthy["validation"]["slides"]

    # A deck that would make PowerPoint offer to repair it must never come back as "ok".
    from cadencia.report.core import validate as validate_module

    real = validate_module.check_package

    def broken(path):
        report = real(path)
        report.fail(
            "ordem_de_esquema",
            Msg(
                "package.ordem_de_esquema",
                part="ppt/slides/slide1.xml",
                element="pPr",
                children=["defRPr", "buChar"],
            ),
        )
        return report

    monkeypatch.setattr(validate_module, "check_package", broken)
    bad = ok(client.post("/api/reports/generate", json={"template_id": uploaded["id"], "render": False}))
    assert bad["status"] == "invalid"
    assert any(e["kind"] == "ordem_de_esquema" for e in bad["validation"]["errors"])
    # And the download is refused without an explicit override.
    assert client.get(f"/api/reports/runs/{bad['id']}/download").status_code == 409


def test_templates_can_be_uploaded_listed_and_deleted(client, tmp_path, settings):
    """RF-51: a template is something you own, not a one-way upload."""
    source = build_sample_template(tmp_path / "modelo.pptx")
    with open(source, "rb") as fh:
        first = ok(
            client.post(
                "/api/reports/templates", files={"file": ("modelo.pptx", fh)}, data={"name": "Semanal"}
            )
        )
    assert first["name"] == "Semanal"

    listing = ok(client.get("/api/reports/templates"))["items"]
    entry = next(t for t in listing if t["id"] == first["id"])
    assert entry["slides"] and entry["size"] and entry["missing_file"] is False
    assert entry["runs"] == 0 and entry["has_mapping"] is False and entry["has_edits"] is False

    # Give it a mapping, an edit and a generated deck, so deletion has something to clean up.
    detail = ok(client.get(f"/api/reports/templates/{first['id']}"))
    mapping = detail["suggestion"]
    mapping.pop("notes", None)
    ok(client.put(f"/api/reports/templates/{first['id']}/mapping", json=mapping))
    run = ok(client.post("/api/reports/generate", json={"template_id": first["id"], "render": False}))
    scene = ok(client.get(f"/api/reports/runs/{run['id']}/scene"))
    edits = {
        "version": 1,
        "slides": {scene["slides"][0]["key"]: {"shapes": {}, "notes": "nota"}},
        "order": [],
        "deleted": [],
        "duplicates": {},
    }
    ok(client.put(f"/api/reports/templates/{first['id']}/edits", json={"edits": edits, "base_revision": 0}))

    entry = next(t for t in ok(client.get("/api/reports/templates"))["items"] if t["id"] == first["id"])
    assert entry["has_mapping"] and entry["has_edits"] and entry["runs"] == 1

    stored = Path(detail["manifest"]["template"]["filename"])
    assert stored  # the row carries the original name
    result = ok(client.delete(f"/api/reports/templates/{first['id']}"))
    assert result["deleted"] and result["runs_kept"] == 1 and result["files"]

    assert client.get(f"/api/reports/templates/{first['id']}").status_code == 404
    assert first["id"] not in [t["id"] for t in ok(client.get("/api/reports/templates"))["items"]]
    assert not (settings.templates_dir / f"{first['sha256']}.pptx").exists()
    # The generated deck survives: it is a finished deliverable, not part of the template.
    assert client.get(f"/api/reports/runs/{run['id']}/download").status_code == 200
    after = ok(client.get(f"/api/reports/runs/{run['id']}/scene"))
    assert after["template_deleted"] is True and after["slides"]
    assert client.delete(f"/api/reports/templates/{first['id']}").status_code == 404


def test_an_orphaned_run_can_be_reassociated_with_a_re_uploaded_template(client, tmp_path):
    """#12: the template deleted, the run relinked to a fresh upload — theme and marks come back."""
    source = build_sample_template(tmp_path / "modelo.pptx")
    with open(source, "rb") as fh:
        original = ok(client.post("/api/reports/templates", files={"file": ("modelo.pptx", fh)}))
    run = ok(client.post("/api/reports/generate", json={"template_id": original["id"], "render": False}))

    # Uploaded *before* the delete, so another row is still alive when `original`'s id is freed —
    # otherwise the store's id sequence resets and the replacement could land on the same id.
    # An extra slide keeps the bytes (and so the content hash) distinct from `source`'s.
    other = tmp_path / "outro.pptx"
    other_prs = Presentation(str(source))
    other_prs.slides.add_slide(other_prs.slide_layouts[6])
    other_prs.save(str(other))
    with open(other, "rb") as fh:
        replacement = ok(client.post("/api/reports/templates", files={"file": ("outro.pptx", fh)}))
    assert replacement["id"] != original["id"]

    ok(client.delete(f"/api/reports/templates/{original['id']}"))
    assert ok(client.get(f"/api/reports/runs/{run['id']}/scene"))["template_deleted"] is True
    assert (
        client.post(f"/api/reports/runs/{run['id']}/reassociate", json={"template_id": 999999}).status_code
        == 404
    )
    assert (
        client.post(
            "/api/reports/runs/naoexiste/reassociate", json={"template_id": replacement["id"]}
        ).status_code
        == 404
    )

    relinked = ok(
        client.post(f"/api/reports/runs/{run['id']}/reassociate", json={"template_id": replacement["id"]})
    )
    assert relinked["template_deleted"] is False and relinked["template_id"] == replacement["id"]
    # A run that already has a template cannot be relinked out from under it.
    again = client.post(f"/api/reports/runs/{run['id']}/reassociate", json={"template_id": replacement["id"]})
    assert again.status_code == 400


def test_a_mapping_can_be_cloned_to_a_redesigned_template(client, tmp_path):
    """#13: a plan reports what would move; nothing is saved until `aplicar` is set."""

    def build(path, extra_shape):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if extra_shape:  # shifts the shape_id below without touching its name or slide position
            slide.shapes.add_textbox(0, 0, 100, 100)
        title = slide.shapes.add_textbox(100, 100, 400, 100)
        title.name = "Titulo"
        title.text_frame.paragraphs[0].text = "x"
        prs.save(str(path))

    original_path = tmp_path / "original.pptx"
    build(original_path, extra_shape=False)
    with open(original_path, "rb") as fh:
        original = ok(client.post("/api/reports/templates", files={"file": ("original.pptx", fh)}))
    manifest = ok(client.get(f"/api/reports/templates/{original['id']}"))["manifest"]
    title_area = next(a for a in manifest["slides"][0]["areas"] if a["name"] == "Titulo")
    mapping = {
        "slides": {
            str(manifest["slides"][0]["slide_id"]): {
                "bindings": [{"target": f"{title_area['id']}#p0", "mode": "text", "field": "sprint.nome"}]
            }
        }
    }
    ok(client.put(f"/api/reports/templates/{original['id']}/mapping", json=mapping))

    redesigned_path = tmp_path / "redesigned.pptx"
    build(redesigned_path, extra_shape=True)
    with open(redesigned_path, "rb") as fh:
        redesigned = ok(client.post("/api/reports/templates", files={"file": ("redesigned.pptx", fh)}))

    plan = ok(
        client.post(
            f"/api/reports/templates/{redesigned['id']}/clone-mapping",
            json={"from_template_id": original["id"]},
        )
    )
    assert plan["matched"] == 1 and plan["unmatched"] == [] and "saved" not in plan
    # Reading the plan changes nothing until it is applied.
    assert ok(client.get(f"/api/reports/templates/{redesigned['id']}"))["mapping"] is None

    applied = ok(
        client.post(
            f"/api/reports/templates/{redesigned['id']}/clone-mapping",
            json={"from_template_id": original["id"], "aplicar": True},
        )
    )
    assert applied["saved"] is True
    saved = ok(client.get(f"/api/reports/templates/{redesigned['id']}"))["mapping"]
    new_manifest = ok(client.get(f"/api/reports/templates/{redesigned['id']}"))["manifest"]
    new_title = next(a for a in new_manifest["slides"][0]["areas"] if a["name"] == "Titulo")
    new_slide_key = str(new_manifest["slides"][0]["slide_id"])
    assert saved["slides"][new_slide_key]["bindings"][0]["target"] == f"{new_title['id']}#p0"

    assert (
        client.post(
            f"/api/reports/templates/{redesigned['id']}/clone-mapping", json={"from_template_id": 999999}
        ).status_code
        == 404
    )


def test_uploading_the_same_file_twice_reuses_the_template(client, tmp_path):
    """The store is content-addressed, so the same bytes must not pile up as separate templates."""
    source = build_sample_template(tmp_path / "modelo.pptx")
    ids = []
    for _ in range(2):
        with open(source, "rb") as fh:
            ids.append(ok(client.post("/api/reports/templates", files={"file": ("modelo.pptx", fh)}))["id"])
    assert ids[0] == ids[1]
    assert len([t for t in ok(client.get("/api/reports/templates"))["items"] if t["id"] == ids[0]]) == 1
    assert client.post("/api/reports/templates", files={"file": ("x.txt", b"nope")}).status_code == 400


def test_the_audit_flags_a_template_that_would_break_instead_of_registering_for_it(
    client, sessions, settings, tmp_path
):
    """RF-45: registration is the source of truth, so the audit names the gap and fixes nothing."""
    from cadencia.report.service import audit_variables

    path = build_sample_template(tmp_path / "modelo.pptx")
    with open(path, "rb") as fh:
        uploaded = ok(client.post("/api/reports/templates", files={"file": ("modelo.pptx", fh)}))
    detail = ok(client.get(f"/api/reports/templates/{uploaded['id']}"))
    slide_id, spec = next(iter(detail["suggestion"]["slides"].items()))
    target = spec["bindings"][0]["target"]
    spec["bindings"] += [
        {"target": target, "mode": "text", "value": "Sobre: {{texto.definicao}}"},
        {"target": target, "mode": "text", "value": "{{texto.definicaoo}}"},
    ]
    detail["suggestion"].pop("notes")
    ok(client.put(f"/api/reports/templates/{uploaded['id']}/mapping", json=detail["suggestion"]))
    ok(client.post("/api/reports/variables", json={"name": "definicao"}))

    with sessions() as session:
        report = audit_variables(session, settings, uploaded["id"])
    result = report["results"][0]
    assert report["ok"] is False and report["breaking"] == [result["template"]]
    # The typo is unregistered; the registered one is merely unanswered. Different fixes.
    assert [u["name"] for u in result["unregistered"]] == ["texto.definicaoo"]
    assert [u["name"] for u in result["unfilled"]] == ["texto.definicao"]
    assert result["unfilled"][0]["boards"] and slide_id in result["unfilled"][0]["used_in"][0]
    # Nothing was registered on the template's behalf.
    assert [v["name"] for v in ok(client.get("/api/reports/variables"))["items"]] == ["definicao"]

    # A default on the registration answers the second one, and only the second one.
    variable = ok(client.get("/api/reports/variables"))["items"][0]
    ok(client.patch(f"/api/reports/variables/{variable['id']}", json={"default": "Sem definição."}))
    with sessions() as session:
        report = audit_variables(session, settings, uploaded["id"])
    assert report["results"][0]["unfilled"] == []
    assert [u["name"] for u in report["results"][0]["unregistered"]] == ["texto.definicaoo"]

    # And generation refuses the deck, naming the variable and the template.
    run = ok(client.post("/api/reports/generate", json={"template_id": uploaded["id"], "render": False}))
    errors = [e for e in run["validation"]["errors"] if e["kind"] == "variavel_nao_registrada"]
    assert errors and "{{texto.definicaoo}}" in errors[0]["message"]
    assert run["validation"]["ok"] is False


def test_the_catalogue_is_browsed_from_the_variable_manager(client):
    """RF-56: one screen answers "what can I put on a slide", with the kinds carried through."""
    definicao, _aviso = define_project_variables(client)
    view = ok(client.get("/api/reports/variables"))
    catalog = {f["key"]: f for f in view["catalog"]}
    # The Jira fields are all there, still typed — a table is rows, a chart a series, not text.
    assert catalog["riscos.itens"]["kind"] == "rows"
    assert catalog["grafico.vazao"]["kind"] == "series"
    assert catalog["board.status_geral"]["kind"] == "status"
    assert catalog["num.entregues"]["kind"] == "number"
    # Name, label, group and hint survive the move.
    assert catalog["espaco.nome"]["label"] and catalog["espaco.nome"]["group"]
    assert catalog["espaco.nome"]["origin"] == "jira"
    # The registered variables sit beside them, marked as such rather than merged in.
    assert catalog["texto.definicao"]["origin"] == "variavel"
    assert catalog["texto.definicao"]["kind"] == "text"
    # Same entries as the standalone endpoint: one catalogue, two doors onto it.
    fields = {f["key"] for f in ok(client.get("/api/reports/catalog"))["fields"]}
    assert fields == set(catalog)
    assert view["conflicts"] == []

    # A name registered twice is reported, never resolved by overwriting one of them.
    assert definicao["template_id"] is None
    bound = ok(
        client.post(
            "/api/reports/variables", json={"name": "definicao", "template_id": 99, "label": "Do modelo"}
        )
    )
    conflicts = ok(client.get("/api/reports/variables", params={"template_id": 99}))["conflicts"]
    assert [c["kind"] for c in conflicts] == ["sombreada"]
    assert conflicts[0]["name"] == "definicao" and conflicts[0]["winner_template_id"] == 99
    # Both rows still exist: nothing was overwritten to make the conflict go away.
    names = [(v["id"], v["template_id"]) for v in ok(client.get("/api/reports/variables"))["items"]]
    assert (definicao["id"], None) in names and (bound["id"], 99) in names


def test_a_variable_can_be_wired_to_a_jira_field(client):
    """RF-58: the connection lives on the registration, and only sensible fields are accepted."""
    wired = ok(
        client.post(
            "/api/reports/variables",
            json={"name": "time_do_board", "label": "Time", "source_field": "board.nome"},
        )
    )
    assert wired["source_field"] == "board.nome"

    # A board is needed for the field to say anything, and then the manager shows what the deck
    # will really carry — resolved through the generator's own `resolve`.
    view = ok(client.get("/api/reports/variables", params={"board_id": 1}))
    assert view["resolved"]["time_do_board"] == "Plataforma"
    assert "time_do_board" in view["from_source"]
    # With no board chosen there is nothing for the field to answer with, and it says so quietly.
    assert "time_do_board" not in ok(client.get("/api/reports/variables"))["resolved"]

    # A typed value is an override somebody made and still wins over the field.
    ok(
        client.put(
            f"/api/reports/variables/{wired['id']}/values",
            json={"scope": "board", "scope_key": "1", "value": "Escrito à mão"},
        )
    )
    view = ok(client.get("/api/reports/variables", params={"board_id": 1}))
    assert view["resolved"]["time_do_board"] == "Escrito à mão"
    assert "time_do_board" not in view["from_source"]

    # RF-59: a variable is typed by its source, so a rows field makes a rows variable. What used to
    # be refused here is now the point of the feature.
    for key, kind in (
        ("riscos.itens", "rows"),
        ("grafico.vazao", "series"),
        ("board.status_geral", "status"),
    ):
        typed = ok(
            client.post(
                "/api/reports/variables", json={"name": f"v_{key.split('.')[1]}", "source_field": key}
            )
        )
        assert typed["source_field"] == key
        catalog = {f["key"]: f for f in ok(client.get("/api/reports/variables"))["catalog"]}
        assert catalog[f"texto.{typed['name']}"]["kind"] == kind
        # …and a typed variable takes no hand-written value: there is no box to type a table into.
        refused = client.put(
            f"/api/reports/variables/{typed['id']}/values",
            json={"scope": "global", "value": "impossível"},
        )
        assert refused.status_code == 400 and kind in refused.json()["detail"]["message"]

    # A field that does not exist is still refused.
    missing = client.post("/api/reports/variables", json={"name": "x_nada", "source_field": "nao.existe"})
    assert missing.status_code == 400 and "catálogo" in missing.json()["detail"]["message"]
    # And never another variable: a variable reading a variable could be pointed at itself.
    cycle = client.post("/api/reports/variables", json={"name": "ciclo", "source_field": "texto.ciclo"})
    assert cycle.status_code == 400 and "outra variável" in cycle.json()["detail"]["message"]

    # Disconnecting needs its own flag, for the same reason clearing a default does.
    assert ok(client.patch(f"/api/reports/variables/{wired['id']}", json={"label": "Time"}))["source_field"]
    assert (
        ok(client.patch(f"/api/reports/variables/{wired['id']}", json={"clear_source": True}))["source_field"]
        is None
    )


def test_the_template_is_read_as_a_scene_like_a_generated_run(client, tmp_path):
    """RF-61: one renderer for the template and the deck, instead of a PNG with boxes on it."""
    path = build_sample_template(tmp_path / "modelo.pptx")
    with open(path, "rb") as fh:
        uploaded = ok(client.post("/api/reports/templates", files={"file": ("modelo.pptx", fh)}))
    scene = ok(client.get(f"/api/reports/templates/{uploaded['id']}/scene"))

    # The same shape as a run's scene: real EMU geometry, the theme, every shape of every slide.
    assert scene["slide_width"] > 0 and scene["slide_height"] > 0
    assert scene["theme"]["colors"] and scene["theme"]["fonts"]["minor"]
    assert scene["run_id"] is None and scene["template_id"] == uploaded["id"]
    assert len(scene["slides"]) == len(
        ok(client.get(f"/api/reports/templates/{uploaded['id']}"))["manifest"]["slides"]
    )
    first = scene["slides"][0]
    assert first["origin_slide_id"] and first["shapes"]
    assert all({"key", "frame", "z", "renderable"} <= set(shape) for shape in first["shapes"])

    # Binding a field marks the shape before anything has been generated, which is the one thing
    # the old clickable-hotspot view was genuinely good at.
    detail = ok(client.get(f"/api/reports/templates/{uploaded['id']}"))
    suggestion = detail["suggestion"]
    suggestion.pop("notes", None)
    ok(client.put(f"/api/reports/templates/{uploaded['id']}/mapping", json=suggestion))
    marked = ok(client.get(f"/api/reports/templates/{uploaded['id']}/scene"))
    assert any(shape["bound"] for slide in marked["slides"] for shape in slide["shapes"])

    assert client.get("/api/reports/templates/9999/scene").status_code == 404


def test_the_status_colours_are_configured_once_for_every_template(client, tmp_path):
    """RF-62: one palette in the manager; a template records only what its legend really says."""
    shared = ok(client.get("/api/reports/palette"))
    assert set(shared["statuses"]) == {"concluido", "em_dia", "risco", "atrasado"}
    assert shared["palette"]["risco"] == "FFC000"

    path = build_sample_template(tmp_path / "modelo.pptx")
    with open(path, "rb") as fh:
        uploaded = ok(client.post("/api/reports/templates", files={"file": ("modelo.pptx", fh)}))

    # The suggestion no longer writes a full copy of the palette into every mapping: it records
    # the statuses the legend actually names, and nothing else. `em_dia` is not in this legend.
    suggestion = ok(client.get(f"/api/reports/templates/{uploaded['id']}"))["suggestion"]
    assert set(suggestion["palette"]) < set(shared["palette"])
    assert "em_dia" not in suggestion["palette"]

    # A template that overrides nothing follows the shared setting…
    suggestion.pop("notes", None)
    suggestion["palette"] = {}
    suggestion["status_labels"] = {}
    ok(client.put(f"/api/reports/templates/{uploaded['id']}/mapping", json=suggestion))
    ok(client.put("/api/reports/palette", json={"palette": {"risco": "ABCDEF"}}))
    run = ok(client.post("/api/reports/generate", json={"template_id": uploaded["id"], "render": False}))
    assert run["status"] in {"ok", "invalid"}  # it generated; the colour came from the setting

    # …and one that does override keeps its own, untouched by the central change.
    suggestion["palette"] = {"risco": "123456"}
    ok(client.put(f"/api/reports/templates/{uploaded['id']}/mapping", json=suggestion))
    assert ok(client.get(f"/api/reports/templates/{uploaded['id']}"))["mapping"]["palette"] == {
        "risco": "123456"
    }

    # An unknown status cannot be smuggled into the shared configuration.
    ok(client.put("/api/reports/palette", json={"palette": {"inventado": "000000"}}))
    assert "inventado" not in ok(client.get("/api/reports/palette"))["palette"]


def test_an_accessible_palette_can_be_suggested_but_is_never_applied_on_its_own(client):
    """#5: Okabe-Ito, offered as a suggestion — saving it still goes through `PUT /palette`."""
    suggestion = ok(client.get("/api/reports/palette/accessible"))["palette"]
    assert set(suggestion) == {"concluido", "em_dia", "risco", "atrasado"}
    assert len(set(suggestion.values())) == 4  # every status gets its own colour
    # Fetching the suggestion changes nothing until it is actually saved.
    assert ok(client.get("/api/reports/palette"))["palette"]["risco"] == "FFC000"
    ok(client.put("/api/reports/palette", json={"palette": suggestion}))
    assert ok(client.get("/api/reports/palette"))["palette"] == suggestion


def test_the_editor_is_told_which_slides_are_not_ready(client, tmp_path):
    """RF-66: which slide, which variable, which board — before anything is generated."""
    path = build_sample_template(tmp_path / "modelo.pptx")
    with open(path, "rb") as fh:
        uploaded = ok(client.post("/api/reports/templates", files={"file": ("modelo.pptx", fh)}))
    detail = ok(client.get(f"/api/reports/templates/{uploaded['id']}"))
    slides = detail["manifest"]["slides"]
    area = next(a for a in slides[1]["areas"] if a.get("paragraphs"))

    # One slide asks for a variable nobody registered, and one for a registered but unfilled one.
    ok(client.post("/api/reports/variables", json={"name": "definicao", "label": "Definição"}))
    ok(
        client.put(
            f"/api/reports/templates/{uploaded['id']}/mapping",
            json={
                "slides": {
                    str(slides[1]["slide_id"]): {
                        "scope": "repeat",
                        "name": "Resumo do board",
                        "bindings": [
                            {
                                "target": area["id"],
                                "mode": "text",
                                "value": "{{texto.definicao}} {{texto.fantasma}}",
                            }
                        ],
                    }
                }
            },
        )
    )
    report = ok(client.get(f"/api/reports/templates/{uploaded['id']}/readiness"))
    assert report["ok"] is False
    entry = next(s for s in report["slides"] if s["slide_id"] == slides[1]["slide_id"])
    assert entry["name"] == "Resumo do board"  # the slide's given name, for the message
    kinds = {p["kind"]: p for p in entry["problems"]}
    assert kinds["nao_registrada"]["variable"] == "texto.fantasma"
    assert kinds["sem_valor"]["variable"] == "texto.definicao"
    assert kinds["sem_valor"]["boards"]  # named, so the message can say who is short of it

    # Giving it a value for every board clears that one, and only that one.
    variable = next(v for v in ok(client.get("/api/reports/variables"))["items"] if v["name"] == "definicao")
    ok(
        client.put(
            f"/api/reports/variables/{variable['id']}/values",
            json={"scope": "global", "value": "Uma frase."},
        )
    )
    again = ok(client.get(f"/api/reports/templates/{uploaded['id']}/readiness"))
    entry = next(s for s in again["slides"] if s["slide_id"] == slides[1]["slide_id"])
    assert [p["kind"] for p in entry["problems"]] == ["nao_registrada"]


def test_the_deck_can_be_downloaded_at_the_end_of_the_process(client, tmp_path):
    """RF-70: the file is the point. It comes back as a .pptx, and a refused one still comes back."""
    path = build_sample_template(tmp_path / "modelo.pptx")
    with open(path, "rb") as fh:
        uploaded = ok(client.post("/api/reports/templates", files={"file": ("modelo.pptx", fh)}))
    suggestion = ok(client.get(f"/api/reports/templates/{uploaded['id']}"))["suggestion"]
    suggestion.pop("notes", None)
    ok(client.put(f"/api/reports/templates/{uploaded['id']}/mapping", json=suggestion))
    run = ok(client.post("/api/reports/generate", json={"template_id": uploaded["id"], "render": False}))

    download = client.get(f"/api/reports/runs/{run['id']}/download")
    assert download.status_code == 200
    assert download.content[:2] == b"PK"  # a real zip, which is what a .pptx is
    assert "presentationml" in download.headers["content-type"]

    # A run the validation refused is still downloadable with ?force=true — the deck somebody needs
    # to open to understand what is wrong with it must not be the one they cannot have.
    slide_id, first = next(iter(suggestion["slides"].items()))
    ok(
        client.put(
            f"/api/reports/templates/{uploaded['id']}/mapping",
            json={
                **suggestion,
                "slides": {
                    **suggestion["slides"],
                    slide_id: {
                        **first,
                        "bindings": [
                            {
                                "target": first["bindings"][0]["target"],
                                "mode": "text",
                                "value": "{{texto.fantasma}}",
                            }
                        ],
                    },
                },
            },
        )
    )
    bad = ok(client.post("/api/reports/generate", json={"template_id": uploaded["id"], "render": False}))
    assert bad["status"] == "invalid"
    assert client.get(f"/api/reports/runs/{bad['id']}/download").status_code == 409
    forced = client.get(f"/api/reports/runs/{bad['id']}/download", params={"force": "true"})
    assert forced.status_code == 200 and forced.content[:2] == b"PK"
